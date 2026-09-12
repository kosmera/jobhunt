"""Rendu DOCX d'un CV ATS.

Mise en page volontairement austère : une colonne, polices standard, pas de
tableau ni d'image — exactement ce que les analyseurs de candidatures lisent
sans se tromper.
"""

from __future__ import annotations

import io
from typing import cast

from docx import Document
from docx.shared import Pt
from docx.styles.style import ParagraphStyle

from jobhunt_ai.agents.schemas import ATSResume

#: Intitulés de sections par langue du CV.
SECTION_TITLES = {
    "fr": {
        "summary": "Profil",
        "skills": "Compétences",
        "experience": "Expérience professionnelle",
        "education": "Formation",
        "certifications": "Certifications",
        "languages": "Langues",
    },
    "en": {
        "summary": "Summary",
        "skills": "Skills",
        "experience": "Professional Experience",
        "education": "Education",
        "certifications": "Certifications",
        "languages": "Languages",
    },
    "nl": {
        "summary": "Profiel",
        "skills": "Vaardigheden",
        "experience": "Werkervaring",
        "education": "Opleiding",
        "certifications": "Certificaten",
        "languages": "Talen",
    },
}


def render_docx(resume: ATSResume, language: str = "fr") -> bytes:
    titles = SECTION_TITLES.get(language, SECTION_TITLES["fr"])
    document = Document()

    # ``styles[...]`` est typé BaseStyle ; « Normal » est un style de
    # paragraphe, seul porteur d'une police.
    style = cast(ParagraphStyle, document.styles["Normal"])
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    document.add_heading(resume.full_name, level=0)
    if resume.headline:
        document.add_paragraph(resume.headline)
    if resume.contact_line:
        document.add_paragraph(resume.contact_line)
    for link in resume.links:
        document.add_paragraph(link)

    if resume.summary:
        document.add_heading(titles["summary"], level=1)
        document.add_paragraph(resume.summary)

    if resume.skill_groups:
        document.add_heading(titles["skills"], level=1)
        for group in resume.skill_groups:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(f"{group.category} : ")
            run.bold = True
            paragraph.add_run(", ".join(group.items))

    if resume.experiences:
        document.add_heading(titles["experience"], level=1)
        for experience in resume.experiences:
            heading = document.add_paragraph()
            run = heading.add_run(experience.title)
            run.bold = True
            meta = " — ".join(
                part
                for part in (experience.company, experience.location, experience.period)
                if part
            )
            if meta:
                heading.add_run(f" — {meta}")
            for bullet in experience.bullets:
                document.add_paragraph(bullet, style="List Bullet")

    for title_key, lines in (
        ("education", resume.education),
        ("certifications", resume.certifications),
        ("languages", resume.languages),
    ):
        if lines:
            document.add_heading(titles[title_key], level=1)
            for line in lines:
                document.add_paragraph(line, style="List Bullet")

    for section in resume.extra_sections:
        document.add_heading(section.title, level=1)
        for line in section.lines:
            document.add_paragraph(line, style="List Bullet")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
