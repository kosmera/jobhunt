"""Tests des quatre agents, modèle simulé."""

from __future__ import annotations

import contextlib
import io
import json
import shutil
import tempfile
from datetime import timedelta
from unittest import mock
from urllib.parse import quote_plus

import docx
import rls
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import TestCase, override_settings
from django.utils import timezone

from accounts.services import profile_for
from jobhunt_ai.tests.fakes import make_premium_user as make_user
from tracker.models import (
    ActivityEvent,
    Application,
    Company,
    Document,
    DocumentKind,
    SkillGap,
    Status,
)

from jobhunt_ai.models import (
    LEAD_SUMMARY_PREFIX,
    AgentRun,
    CandidateProfile,
    GeneratedCV,
    MatchReport,
    OfferLead,
    RunKind,
    RunStatus,
)
from jobhunt_ai.services import runner
from jobhunt_ai.tests import fakes

MEDIA = tempfile.mkdtemp(prefix="jobhunt-ai-tests-")

#: Un CV comme il en arrive vraiment : avec le nom, l'e-mail et le numéro du
#: candidat en tête. Le cœur les masque avant que le copilote n'en voie une
#: ligne — c'est justement ce que les tests d'analyse vérifient.
CV_TEXT = (
    "Lionel Test — Ingénieur DevOps senior.\n"
    "lionel.test@example.be · 0470 12 34 56 · Nivelles\n"
    "Quinze ans d'infrastructure Linux, d'automatisation Ansible et de Python.\n"
    "Expérience : Acme (2019-), migration de 200 serveurs sans coupure.\n"
    "Formation : Master en informatique, UCL, 2008. Langues : français C2."
)


def default_owner() -> User:
    """Le compte connecté du test — ou un compte neuf pour les tests d'agents."""
    return User.objects.order_by("pk").first() or make_user()


def make_cv_document(**kwargs) -> Document:
    upload = SimpleUploadedFile("cv-lionel.txt", CV_TEXT.encode("utf-8"))
    defaults = {"kind": DocumentKind.CV, "label": "CV de base", "file": upload}
    defaults.update(kwargs)
    defaults.setdefault("owner", default_owner())
    return Document.objects.create(**defaults)


def ingest_cv_document(
    owner=None,
    *,
    body: str = CV_TEXT,
    filename: str = "cv-lionel.txt",
    label: str = "CV de base",
    make_primary: bool = False,
):
    """Le vrai chemin d'un CV, de bout en bout.

    Le cœur range les octets par le port de stockage, en relit le texte,
    l'anonymise, puis appelle l'analyseur que l'extension déclare sur son
    descripteur — lequel lance l'exécution. Renvoie ``(intake, run)`` ;
    ``run`` vaut ``None`` quand le CV n'était pas analysable.
    """
    from tracker import privacy, services

    from jobhunt_ai.hooks import CopilotCVAnalyzer

    owner = owner or default_owner()
    account = profile_for(owner)
    analyzer = CopilotCVAnalyzer(make_primary=make_primary)
    intake = services.ingest_cv(
        owner,
        None,
        Document(kind=DocumentKind.CV),
        upload=SimpleUploadedFile(filename, body.encode("utf-8")),
        label=label,
        known=privacy.known_identity(
            display_name=account.display_name,
            username=owner.get_username(),
            email=owner.email,
            location=account.location,
        ),
        analyzer=analyzer,
    )
    return intake, analyzer.run


def read_stored(document: Document) -> bytes:
    """Les octets d'un document, par le port de stockage — jamais par un
    chemin : ``Storage.path`` n'existe ni en mémoire ni chez Azure."""
    from tracker.adapters import storage

    with storage().open_file(document.file.name or "") as handle:
        return handle.read()


def make_application(**kwargs) -> Application:
    company = kwargs.pop("company", None)
    owner = kwargs.pop("owner", None) or (company.owner if company else default_owner())
    company = company or Company.objects.create(owner=owner, name=kwargs.pop("company_name", "Acme"))
    defaults = {
        "title": "Senior DevOps Engineer",
        "status": Status.TO_APPLY,
        "posting_raw": "Nous cherchons un profil Ansible / Linux / Terraform. " * 4,
    }
    defaults.update(kwargs)
    return Application.objects.create(owner=owner, company=company, **defaults)


def make_profile(**kwargs) -> CandidateProfile:
    defaults = {
        "label": "CV de base",
        "is_primary": True,
        "full_name": "Lionel Test",
        "headline": "Ingénieur DevOps senior",
        "summary": "Quinze ans d'infrastructure.",
        "skills": [{"name": "Ansible", "category": "DevOps"}],
    }
    defaults.update(kwargs)
    defaults.setdefault("owner", default_owner())
    return CandidateProfile.objects.create(**defaults)


@override_settings(MEDIA_ROOT=MEDIA)
class CVParserTests(TestCase):
    """L'analyse de CV, du téléversement au profil.

    Le copilote n'ouvre plus le fichier : c'est ``services.ingest_cv`` du
    cœur qui l'enregistre par le port de stockage, en relit le texte,
    l'anonymise et appelle l'analyseur déclaré sur le descripteur. Les tests
    passent donc par ce chemin plutôt que de fabriquer une exécution à la
    main.
    """

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA, ignore_errors=True)

    def test_parse_creates_profile_from_document(self):
        with fakes.fake_llm(), fakes.eager_runs():
            intake, run = ingest_cv_document(make_primary=True)
        assert run is not None
        self.assertEqual(run.status, RunStatus.SUCCEEDED, run.error)
        self.assertTrue(intake.analyzed)
        profile = CandidateProfile.objects.get()
        self.assertTrue(profile.is_primary)
        self.assertEqual(profile.source_document, intake.document)
        self.assertEqual(profile.label, "CV de base")
        self.assertEqual(len(profile.skills), 3)
        self.assertEqual(profile.language, "fr")
        self.assertIn("Ansible", profile.skill_names)
        self.assertEqual(run.profile, profile)
        self.assertEqual(run.result["skill_count"], 3)

    def test_first_profile_becomes_primary_even_unrequested(self):
        with fakes.fake_llm(), fakes.eager_runs():
            _, run = ingest_cv_document(make_primary=False)
        assert run is not None
        self.assertEqual(run.status, RunStatus.SUCCEEDED, run.error)
        self.assertTrue(CandidateProfile.objects.get().is_primary)

    def test_experience_skills_complete_the_profile_without_losing_existing_details(self):
        from jobhunt_ai.agents.schemas import ParsedProfile, ProfileSkill

        parsed = fakes.PARSED_PROFILE.model_copy(deep=True)
        names = ["Python", "Go", "C/C++", "JavaScript", "Ansible", "Jenkins", "OpenStack", "GitLab CI"]
        parsed.skills = [
            ProfileSkill(name=name, category="", level="", years="") for name in names
        ]
        parsed.skills[4] = ProfileSkill(name="Ansible", category="DevOps", level="expert", years="8")
        parsed.skills.extend([
            ProfileSkill(name="  ansible  ", category="", level="", years=""),
            ProfileSkill(name=" \n ", category="", level="", years=""),
        ])
        parsed.experiences[0].skills = [
            "ANSIBLE", " GitLab   CI ", "", " \n ", "Dart", "Kotlin", "OpenCV",
            "Flutter", "TestRail", "Linux", "  React   Native  ", "flutter",
        ]
        original = parsed.model_dump()
        expected_names = names + ["Dart", "Kotlin", "OpenCV", "Flutter", "TestRail", "Linux", "React Native"]
        body = CV_TEXT + "\nCompétences : " + ", ".join(expected_names)

        with fakes.fake_llm({ParsedProfile: parsed}), fakes.eager_runs():
            _, run = ingest_cv_document(body=body)

        assert run is not None
        self.assertEqual(run.status, RunStatus.SUCCEEDED, run.error)
        stored = CandidateProfile.objects.get()
        self.assertEqual(stored.skill_names, expected_names)
        self.assertEqual(run.result["skill_count"], 15)
        self.assertEqual(stored.skills[4], {
            "name": "Ansible", "category": "DevOps", "level": "expert", "years": "8",
        })
        self.assertEqual(stored.skills[-1], {
            "name": "React Native", "category": "", "level": "", "years": "",
        })
        self.assertEqual(parsed.model_dump(), original)

    def test_declared_languages_and_proficiency_are_preserved_for_matching(self):
        from jobhunt_ai.agents.qualifications import profile_context
        from jobhunt_ai.agents.schemas import ParsedProfile, ProfileLanguage

        expected = [
            {"name": "Français", "level": "Langue maternelle"},
            {"name": "Anglais", "level": "C1"},
            {"name": "Néerlandais", "level": "Notions"},
            {"name": "Allemand", "level": ""},
        ]
        parsed = fakes.PARSED_PROFILE.model_copy(deep=True)
        parsed.languages = [ProfileLanguage(**language) for language in expected]
        body = CV_TEXT.replace(
            "Langues : français C2.",
            "Langues : français — langue maternelle ; anglais — C1 ; "
            "néerlandais — notions ; allemand.",
        )

        with fakes.fake_llm({ParsedProfile: parsed}), fakes.eager_runs():
            _, run = ingest_cv_document(body=body)

        assert run is not None
        self.assertEqual(run.status, RunStatus.SUCCEEDED, run.error)
        stored = CandidateProfile.objects.get()
        self.assertEqual(stored.languages, expected)
        self.assertEqual(json.loads(profile_context(stored))["langues"], expected)

    def test_document_language_does_not_fill_in_unstated_language_skills(self):
        from jobhunt_ai.agents.qualifications import profile_context
        from jobhunt_ai.agents.schemas import ParsedProfile

        parsed = fakes.PARSED_PROFILE.model_copy(deep=True)
        parsed.languages = []
        parsed.detected_language = "fr"
        body = CV_TEXT.replace("Langues : français C2.", "")

        with fakes.fake_llm({ParsedProfile: parsed}), fakes.eager_runs():
            _, run = ingest_cv_document(body=body)

        assert run is not None
        self.assertEqual(run.status, RunStatus.SUCCEEDED, run.error)
        stored = CandidateProfile.objects.get()
        self.assertEqual(stored.language, "fr")
        self.assertEqual(stored.languages, [])
        self.assertNotIn("langues", json.loads(profile_context(stored)))

    def test_the_model_only_ever_sees_the_anonymised_text(self):
        """Le CV part au modèle sans e-mail, sans téléphone, sans nom."""
        owner = make_user("Lionel Test", email="lionel.test@example.be", location="Nivelles")
        with fakes.fake_llm(), fakes.eager_runs():
            _, run = ingest_cv_document(owner)
        assert run is not None
        sent = run.params["text"]
        for identifier in ("lionel.test@example.be", "0470 12 34 56", "Lionel Test"):
            self.assertNotIn(identifier, sent)
        for marker in ("[EMAIL]", "[TELEPHONE]", "[NOM]"):
            self.assertIn(marker, sent)
        # Ce qui est rangé sur le profil est ce même texte, pas l'original.
        self.assertEqual(CandidateProfile.objects.get().raw_text, sent)
        self.assertTrue(run.params["redactions"])

    def test_the_profile_identity_comes_from_the_account(self):
        """Le CV n'en porte plus que des marqueurs : nom, e-mail, téléphone et
        point de départ viennent du compte, jamais de ``[NOM]``/``[EMAIL]``."""
        owner = make_user(
            "Lionel Test", email="lionel.test@example.be", location="Nivelles",
            phone="+32 470 12 34 56",
        )
        with fakes.fake_llm(), fakes.eager_runs():
            ingest_cv_document(owner)
        profile = CandidateProfile.objects.get()
        self.assertEqual(profile.full_name, "Lionel Test")
        self.assertEqual(profile.email, "lionel.test@example.be")
        self.assertEqual(profile.phone, "+32 470 12 34 56")
        self.assertEqual(profile.location, "Nivelles")
        # Le titre, lui, reste celui que le modèle a lu dans le CV.
        self.assertEqual(profile.headline, "Ingénieur DevOps senior")

    def test_an_account_without_a_phone_leaves_the_field_empty(self):
        """Le téléphone est facultatif : rien n'est inventé pour le combler."""
        owner = make_user("Lionel Test", email="lionel.test@example.be")
        with fakes.fake_llm(), fakes.eager_runs():
            ingest_cv_document(owner)
        profile = CandidateProfile.objects.get()
        self.assertEqual(profile.phone, "")
        # …et le contexte remis au générateur ne porte pas de clé vide.
        from jobhunt_ai.agents.qualifications import profile_context

        self.assertNotIn("telephone", profile_context(profile))

    def test_a_cv_without_enough_text_is_stored_but_not_analysed(self):
        """Plus de repli « PDF envoyé tel quel » : le renvoyer rendrait au
        modèle exactement ce que l'anonymisation vient de masquer."""
        with fakes.fake_llm(), fakes.eager_runs():
            intake, run = ingest_cv_document(body="trop court")
        self.assertIsNone(run)
        self.assertFalse(intake.analyzed)
        self.assertEqual(AgentRun.objects.count(), 0)
        self.assertTrue(Document.objects.filter(pk=intake.document.pk).exists())

    def test_a_run_launched_without_text_fails_cleanly(self):
        """Filet pour une exécution créée autrement (l'admin, une relance)."""
        document = make_cv_document()
        with fakes.fake_llm(), fakes.eager_runs():
            run = runner.launch(
                RunKind.PARSE_CV, owner=document.owner, params={"document_id": document.pk}
            )
        self.assertEqual(run.status, RunStatus.FAILED)
        self.assertIn("texte exploitable", run.error)

    @override_settings(STORAGES=fakes.REMOTE_STORAGES)
    def test_the_flow_never_needs_a_filesystem_path(self):
        """Le fournisseur « mémoire » n'a pas de chemin, Azure non plus."""
        with fakes.fake_llm(), fakes.eager_runs():
            intake, run = ingest_cv_document(make_primary=True)
        assert run is not None
        self.assertEqual(run.status, RunStatus.SUCCEEDED, run.error)
        self.assertEqual(CandidateProfile.objects.count(), 1)
        with self.assertRaises(NotImplementedError):
            intake.document.file.path


class CVAnalyzerHookTests(TestCase):
    """Le descripteur branche bien le copilote sur le port du cœur."""

    def test_core_publishes_cv_events_without_a_direct_ai_analyzer(self):
        from tracker.adapters import cv_analyzer
        from tracker.events import CVEventPublisher, cv_ingested

        self.assertIsInstance(cv_analyzer(), CVEventPublisher)
        # Django keeps ((dispatch_uid, sender_id), receiver, is_async) entries.
        uids = [entry[0][0] for entry in cv_ingested.receivers]
        self.assertIn("jobhunt_ai.cv_ingested", uids)

    def test_the_analyzer_returns_at_once_and_leaves_a_pending_run(self):
        """Le cœur appelle depuis sa transaction : rien de long ici, et le
        worker ne voit le message qu'après le COMMIT."""
        from tracker.privacy import AnonymizedText

        from jobhunt_ai.hooks import CopilotCVAnalyzer

        owner = default_owner()
        document = make_cv_document(owner=owner)
        analyzer = CopilotCVAnalyzer(make_primary=True)
        with mock.patch.object(runner, "_dispatch") as dispatch:
            analyzer.analyze_cv(
                owner,
                document_id=document.pk,
                label="Mon CV",
                language="fr",
                text=AnonymizedText("Un CV anonymisé.", {"EMAIL": 1}),
            )
        run = analyzer.run
        assert run is not None
        self.assertEqual(run.status, RunStatus.PENDING)
        self.assertEqual(run.owner, owner)
        self.assertEqual(
            run.params,
            {
                "document_id": document.pk,
                "label": "Mon CV",
                "language": "fr",
                "make_primary": True,
                "text": "Un CV anonymisé.",
                "redactions": {"EMAIL": 1},
            },
        )
        # The task is durable, but no worker executes it inside the request.
        dispatch.assert_not_called()
        self.assertIsNotNone(run.task_id)


class MatcherTests(TestCase):
    def test_match_scores_and_reports(self):
        profile = make_profile()
        application = make_application()
        with fakes.fake_llm(), fakes.eager_runs():
            run = runner.launch(RunKind.MATCH, application=application)
        self.assertEqual(run.status, RunStatus.SUCCEEDED, run.error)

        application.refresh_from_db()
        self.assertEqual(application.score, 76)
        report = MatchReport.objects.get()
        self.assertEqual(report.application, application)
        self.assertEqual(report.profile, profile)
        self.assertEqual(report.missing_skills, ["Terraform"])
        # Les lacunes remontent dans la table SkillGap du cœur.
        gap = SkillGap.objects.get(name="Terraform")
        self.assertIn("homelab", gap.action_plan)
        # L'événement apparaît sur la ligne du temps.
        self.assertTrue(
            ActivityEvent.objects.filter(
                application=application, title__icontains="Compatibilité évaluée"
            ).exists()
        )

    def test_match_fills_only_empty_analysis_fields(self):
        make_profile()
        application = make_application(strengths="- Déjà rédigé à la main")
        with fakes.fake_llm(), fakes.eager_runs():
            runner.launch(RunKind.MATCH, application=application)
        application.refresh_from_db()
        self.assertEqual(application.strengths, "- Déjà rédigé à la main")
        self.assertEqual(application.weaknesses, "- Pas de Terraform en production")

    def test_match_replaces_the_pre_triage_summary(self):
        """Le résumé écrit à l'import justifie le pré-tri. Laissé en place, il
        expliquerait un score que l'évaluation vient de remplacer — c'est le
        « 88 % » qui restait affiché à côté d'un 65 %."""
        make_profile()
        application = make_application(
            summary=f"{LEAD_SUMMARY_PREFIX}Stack alignée, à 15 km."
        )
        with fakes.fake_llm(), fakes.eager_runs():
            runner.launch(RunKind.MATCH, application=application)
        application.refresh_from_db()
        self.assertEqual(application.score, 76)
        self.assertEqual(
            application.summary, "Bon recouvrement technique, mais Terraform manque."
        )

    def test_match_keeps_a_handwritten_summary(self):
        make_profile()
        application = make_application(summary="Résumé écrit à la main.")
        with fakes.fake_llm(), fakes.eager_runs():
            runner.launch(RunKind.MATCH, application=application)
        application.refresh_from_db()
        self.assertEqual(application.summary, "Résumé écrit à la main.")

    def test_match_existing_skill_gap_not_overwritten(self):
        make_profile()
        SkillGap.objects.create(
            owner=default_owner(), name="Terraform", demand_count=6, why_it_matters="Déjà documenté."
        )
        application = make_application()
        with fakes.fake_llm(), fakes.eager_runs():
            runner.launch(RunKind.MATCH, application=application)
        gap = SkillGap.objects.get(name="Terraform")
        self.assertEqual(gap.why_it_matters, "Déjà documenté.")
        self.assertEqual(gap.demand_count, 6)

    def test_match_uses_the_account_s_own_profile_and_gaps(self):
        owner = make_user("Lionel")
        other = make_user("Marie")
        mine = make_profile(owner=owner)
        # Marie's profile is the most recent one: an owner-blind ``primary()``
        # would pick it.
        theirs = make_profile(owner=other, label="Profil de Marie")
        SkillGap.objects.create(owner=other, name="Terraform", why_it_matters="Chez Marie.")
        self.assertEqual(CandidateProfile.primary(owner), mine)
        self.assertEqual(CandidateProfile.primary(other), theirs)
        application = make_application(owner=owner)
        with fakes.fake_llm(), fakes.eager_runs():
            run = runner.launch(RunKind.MATCH, application=application)
        self.assertEqual(run.status, RunStatus.SUCCEEDED, run.error)
        self.assertEqual(run.profile, mine)
        self.assertEqual(SkillGap.objects.filter(name="Terraform").count(), 2)
        self.assertEqual(SkillGap.objects.get(owner=other, name="Terraform").why_it_matters, "Chez Marie.")

    def test_match_without_profile_fails(self):
        application = make_application()
        with fakes.fake_llm(), fakes.eager_runs():
            run = runner.launch(RunKind.MATCH, application=application)
        self.assertEqual(run.status, RunStatus.FAILED)
        self.assertIn("profil", run.error.lower())

    def test_match_without_posting_text_fails(self):
        make_profile()
        application = make_application(posting_raw="", summary="", url="")
        with fakes.fake_llm(), fakes.eager_runs():
            run = runner.launch(RunKind.MATCH, application=application)
        self.assertEqual(run.status, RunStatus.FAILED)
        self.assertIn("annonce", run.error)


@override_settings(MEDIA_ROOT=MEDIA)
class GeneratorTests(TestCase):
    def test_generates_ats_docx_attached_to_application(self):
        profile = make_profile()
        application = make_application()
        with fakes.fake_llm(), fakes.eager_runs():
            run = runner.launch(
                RunKind.GENERATE_CV, params={"language": "fr"}, application=application
            )
        self.assertEqual(run.status, RunStatus.SUCCEEDED, run.error)

        generated = GeneratedCV.objects.get()
        self.assertEqual(generated.application, application)
        self.assertEqual(generated.profile, profile)
        document = generated.document
        self.assertIsNotNone(document)
        self.assertEqual(document.application, application)
        self.assertEqual(document.kind, DocumentKind.CV)
        self.assertIn("CV IA", document.label)

        # Relu par le port de stockage : ``file.path`` n'existe pas partout.
        self.assertEqual(document.size_bytes, len(read_stored(document)))
        rendered = docx.Document(io.BytesIO(read_stored(document)))
        text = "\n".join(p.text for p in rendered.paragraphs)
        self.assertIn("Lionel Test", text)
        self.assertIn("Migration de 200 serveurs sans coupure", text)
        self.assertIn("Compétences", text)

    @override_settings(STORAGES=fakes.REMOTE_STORAGES)
    def test_generates_without_a_filesystem_path(self):
        """Les octets partent par le port, hors transaction : rien n'exige un
        disque, et le fournisseur « mémoire » n'a pas de chemin."""
        make_profile()
        application = make_application()
        with fakes.fake_llm(), fakes.eager_runs():
            run = runner.launch(
                RunKind.GENERATE_CV, params={"language": "fr"}, application=application
            )
        self.assertEqual(run.status, RunStatus.SUCCEEDED, run.error)
        document = GeneratedCV.objects.get().document
        assert document is not None
        with self.assertRaises(NotImplementedError):
            document.file.path
        self.assertTrue(read_stored(document).startswith(b"PK"))

    @override_settings(STORAGES=fakes.REMOTE_STORAGES)
    def test_a_failed_row_write_leaves_no_orphan_file(self):
        """Le fichier est écrit avant les lignes : si elles ne passent pas,
        il ne doit rester aucun blob que plus rien ne désigne. Le magasin de
        la doublure est neuf, donc « vide » se lit sans ambiguïté."""
        from tracker.adapters import storage

        make_profile()
        application = make_application()
        with fakes.fake_llm(), fakes.eager_runs(), mock.patch.object(
            GeneratedCV.objects, "create", side_effect=RuntimeError("panne")
        ):
            run = runner.launch(
                RunKind.GENERATE_CV, params={"language": "fr"}, application=application
            )
        self.assertEqual(run.status, RunStatus.FAILED)
        self.assertIn("panne", run.error)
        files = storage()
        assert isinstance(files, fakes.PathlessStorage)
        self.assertEqual(files.blobs, {})
        self.assertEqual(GeneratedCV.objects.count(), 0)


class ScoutTests(TestCase):
    def test_scout_creates_scored_and_deduplicated_leads(self):
        profile = make_profile()
        # L'offre Acme est déjà suivie : elle doit être dédupliquée par URL.
        make_application(url="https://example.org/jobs/acme-devops")

        with fakes.fake_llm(), fakes.eager_runs(), mock.patch(
            "jobhunt_ai.scraping.sources.fetch_page",
            return_value="DevOps Engineer chez Widgets SA…",
        ):
            run = runner.launch(
                RunKind.SCOUT, owner=profile.owner,
                params={"location": "Nivelles", "radius_km": 40},
            )
        self.assertEqual(run.status, RunStatus.SUCCEEDED, run.error)
        lead = OfferLead.objects.get()
        self.assertEqual(lead.title, "DevOps Engineer")
        self.assertEqual(lead.score, 81)
        self.assertEqual(lead.score_confidence, "moyenne")
        self.assertEqual(lead.score_blockers, ["Terraform exigé"])
        self.assertEqual(lead.owner, profile.owner)
        self.assertEqual(run.result["created"], 1)

    def test_scout_only_deduplicates_against_the_same_account(self):
        """Une offre suivie par un autre compte n'est pas une raison de la taire."""
        profile = make_profile()
        other = make_user("Marie")
        make_application(owner=other, url="https://example.org/jobs/acme-devops", company_name="Acme")
        with fakes.fake_llm(), fakes.eager_runs(), mock.patch(
            "jobhunt_ai.scraping.sources.fetch_page", return_value="des offres…"
        ):
            run = runner.launch(RunKind.SCOUT, owner=profile.owner, params={"keywords": "devops"})
        self.assertEqual(run.status, RunStatus.SUCCEEDED, run.error)
        self.assertEqual(OfferLead.objects.filter(owner=profile.owner).count(), 2)

    def test_scout_defaults_to_the_account_s_home_base(self):
        profile = make_profile()
        from accounts.services import preferences_for, profile_for

        account = profile_for(profile.owner)
        account.location = "Braine-l'Alleud"
        account.save()
        preferences = preferences_for(profile.owner)
        preferences.search_radius_km = 25
        preferences.save()
        with fakes.fake_llm() as calls, fakes.eager_runs(), mock.patch(
            "jobhunt_ai.scraping.sources.fetch_page", return_value="des offres…"
        ):
            run = runner.launch(RunKind.SCOUT, owner=profile.owner, params={"keywords": "devops"})
        self.assertEqual(run.status, RunStatus.SUCCEEDED, run.error)
        from jobhunt_ai.agents.schemas import LeadScores

        scoring = next(call for call in calls if call["schema"] is LeadScores)
        self.assertIn("Braine-l'Alleud, rayon 25 km", scoring["content"])

    def test_scout_scores_against_the_qualification_rubric(self):
        profile = make_profile()
        with fakes.fake_llm() as calls, fakes.eager_runs(), mock.patch(
            "jobhunt_ai.scraping.sources.fetch_page", return_value="des offres…"
        ):
            run = runner.launch(RunKind.SCOUT, owner=profile.owner, params={"keywords": "devops"})
        self.assertEqual(run.status, RunStatus.SUCCEEDED, run.error)
        from jobhunt_ai.agents.schemas import LeadScores

        scoring = next(call for call in calls if call["schema"] is LeadScores)
        self.assertIn("<qualifications>", scoring["content"])
        self.assertIn("Maîtrisé : Ansible, Linux", scoring["content"])

    def test_scout_uses_keywords_without_llm_queries(self):
        profile = make_profile()
        with fakes.fake_llm() as calls, fakes.eager_runs(), mock.patch(
            "jobhunt_ai.scraping.sources.fetch_page", return_value="peu importe"
        ):
            run = runner.launch(
                RunKind.SCOUT, owner=profile.owner, params={"keywords": "sre bruxelles"}
            )
        self.assertEqual(run.status, RunStatus.SUCCEEDED, run.error)
        from jobhunt_ai.agents.schemas import ScoutQueries

        self.assertFalse(any(call["schema"] is ScoutQueries for call in calls))

    def test_scout_all_sources_down_fails_with_detail(self):
        profile = make_profile()
        with fakes.fake_llm(), fakes.eager_runs(), mock.patch(
            "jobhunt_ai.scraping.sources.fetch_page",
            side_effect=OSError("connexion refusée"),
        ):
            run = runner.launch(RunKind.SCOUT, owner=profile.owner, params={"keywords": "devops"})
        self.assertEqual(run.status, RunStatus.FAILED)
        self.assertGreater(run.result["targets_failed"], 0)
        self.assertTrue(run.result["warnings"])


class QualificationTests(TestCase):
    """La grille de qualifications et l'échelle partagée par les deux noteurs."""

    def test_the_two_scorers_share_one_scale(self):
        """Sans texte commun, chaque invite redéfinit le score à sa façon et la
        même offre passe de 88 % au pré-tri à 65 % à l'évaluation."""
        from jobhunt_ai.agents import matcher, scout
        from jobhunt_ai.agents.qualifications import SCORE_SCALE

        self.assertIn(SCORE_SCALE, matcher.SYSTEM_PROMPT)
        self.assertIn(SCORE_SCALE, scout.SCORE_PROMPT)

    def test_rubric_is_derived_once_and_reused(self):
        from jobhunt_ai.agents import qualifications
        from jobhunt_ai.agents.schemas import QualificationRubric

        profile = make_profile()
        with fakes.fake_llm() as calls:
            first = qualifications.rubric_for(profile)
            second = qualifications.rubric_for(profile)
        self.assertEqual(first, second)
        self.assertEqual(first["core_skills"], ["Ansible", "Linux"])
        self.assertEqual(
            sum(1 for call in calls if call["schema"] is QualificationRubric), 1
        )

        # Rangée sur le profil : une veille ultérieure ne la repaiera pas.
        reloaded = CandidateProfile.objects.get(pk=profile.pk)
        self.assertEqual(reloaded.qualifications["seniority"], "senior, 15 ans")
        with fakes.fake_llm() as calls:
            qualifications.rubric_for(reloaded)
        self.assertEqual(calls, [])

    def test_rubric_text_drops_the_empty_lines(self):
        from jobhunt_ai.agents import qualifications

        text = qualifications.rubric_text(
            {"seniority": "senior", "core_skills": [], "deal_breakers": ["néerlandais"]}
        )
        self.assertIn("Séniorité : senior", text)
        self.assertIn("Rédhibitoire : néerlandais", text)
        self.assertNotIn("Maîtrisé", text)


class BrightDataTests(TestCase):
    """La lecture des sources passe par Bright Data dès qu'un jeton existe."""

    def test_endpoint_carries_token_and_keeps_other_parameters(self):
        from jobhunt_ai.scraping import brightdata

        with mock.patch("jobhunt_ai.conf.BRIGHTDATA_API_TOKEN", "jeton-secret"), \
                mock.patch(
                    "jobhunt_ai.conf.BRIGHTDATA_MCP_URL",
                    "https://mcp.brightdata.com/mcp?groups=browser",
                ):
            endpoint = brightdata._endpoint()
        self.assertIn("groups=browser", endpoint)
        self.assertIn("token=jeton-secret", endpoint)

    def test_token_is_redacted_from_messages(self):
        """Le jeton est dans l'URL : une erreur réseau la recrache souvent, et
        ``AgentRun.error`` est affiché dans l'interface."""
        from jobhunt_ai.scraping import brightdata

        with mock.patch("jobhunt_ai.conf.BRIGHTDATA_API_TOKEN", "jeton-secret"):
            message = brightdata._redact(
                "connect error for https://mcp.brightdata.com/mcp?token=jeton-secret"
            )
        self.assertNotIn("jeton-secret", message)
        self.assertIn("***", message)

    def test_exception_group_is_unwrapped_to_the_real_cause(self):
        """Le client MCP tourne sous anyio : sans dépliage, l'utilisateur ne
        lit que « unhandled errors in a TaskGroup »."""
        from jobhunt_ai.scraping import brightdata

        group = ExceptionGroup(
            "unhandled errors in a TaskGroup",
            [RuntimeError("Client error '401 Invalid API Token'")],
        )
        self.assertEqual(
            brightdata._describe(group), "Client error '401 Invalid API Token'"
        )

    def test_missing_token_is_reported_not_silently_unauthenticated(self):
        from jobhunt_ai.scraping import brightdata

        with mock.patch("jobhunt_ai.conf.BRIGHTDATA_API_TOKEN", ""):
            self.assertFalse(brightdata.is_configured())
            with self.assertRaises(brightdata.BrightDataError):
                brightdata._endpoint()

    def test_tool_error_becomes_brightdata_error(self):
        from jobhunt_ai.scraping import brightdata

        result = mock.Mock(
            isError=True,
            content=[mock.Mock(type="text", text="quota dépassé")],
        )
        with mock.patch("jobhunt_ai.conf.BRIGHTDATA_API_TOKEN", "jeton"), mock.patch.object(
            brightdata, "_run_sync", return_value=None
        ):
            # _run_sync est court-circuité : on valide la lecture du résultat.
            self.assertEqual(brightdata._result_text(result), "quota dépassé")

    def test_fetch_page_text_prefers_brightdata(self):
        from jobhunt_ai.scraping import sources

        with mock.patch("jobhunt_ai.conf.BRIGHTDATA_API_TOKEN", "jeton"), mock.patch(
            "jobhunt_ai.scraping.brightdata.scrape_as_markdown",
            return_value="# Offres\n- [DevOps](https://example.org/x)",
        ) as scrape, mock.patch(
            "jobhunt_ai.scraping.sources._fetch_direct"
        ) as direct:
            text = sources.fetch_page_text("https://www.ictjob.be/fr/x")

        scrape.assert_called_once_with("https://www.ictjob.be/fr/x")
        direct.assert_not_called()
        self.assertIn("DevOps", text)

    def test_fetch_page_text_falls_back_to_direct_read(self):
        from jobhunt_ai.scraping import brightdata, sources

        with mock.patch("jobhunt_ai.conf.BRIGHTDATA_API_TOKEN", "jeton"), mock.patch(
            "jobhunt_ai.conf.BRIGHTDATA_FALLBACK", True
        ), mock.patch(
            "jobhunt_ai.scraping.brightdata.scrape_as_markdown",
            side_effect=brightdata.BrightDataError("quota dépassé"),
        ), mock.patch(
            "jobhunt_ai.scraping.sources._fetch_direct", return_value="page nue"
        ):
            self.assertEqual(sources.fetch_page_text("https://example.org"), "page nue")

    def test_fallback_can_be_disabled_to_surface_the_error(self):
        from jobhunt_ai.scraping import brightdata, sources

        with mock.patch("jobhunt_ai.conf.BRIGHTDATA_API_TOKEN", "jeton"), mock.patch(
            "jobhunt_ai.conf.BRIGHTDATA_FALLBACK", False
        ), mock.patch(
            "jobhunt_ai.scraping.brightdata.scrape_as_markdown",
            side_effect=brightdata.BrightDataError("quota dépassé"),
        ), mock.patch(
            "jobhunt_ai.scraping.sources._fetch_direct"
        ) as direct, self.assertRaises(brightdata.BrightDataError):
            sources.fetch_page_text("https://example.org")
        direct.assert_not_called()

    def test_without_token_reading_stays_direct(self):
        from jobhunt_ai.scraping import sources

        with mock.patch("jobhunt_ai.conf.BRIGHTDATA_API_TOKEN", ""), mock.patch(
            "jobhunt_ai.scraping.brightdata.scrape_as_markdown"
        ) as scrape, mock.patch(
            "jobhunt_ai.scraping.sources._fetch_direct", return_value="page nue"
        ):
            self.assertEqual(sources.fetch_page_text("https://example.org"), "page nue")
        scrape.assert_not_called()

    def test_page_text_is_truncated_to_the_limit(self):
        from jobhunt_ai.scraping import sources

        with mock.patch("jobhunt_ai.conf.BRIGHTDATA_API_TOKEN", "jeton"), mock.patch(
            "jobhunt_ai.scraping.brightdata.scrape_as_markdown", return_value="x" * 5000
        ):
            self.assertEqual(len(sources.fetch_page_text("https://example.org", 100)), 100)


class ScoutSourceShapeTests(TestCase):
    """Une source est soit une page à lire, soit une recherche à lancer."""

    SEARCH_SOURCE = (
        '[{"name": "Google", "search": "{query} emploi Brabant wallon", '
        '"engine": "google", "geo_location": "be"}]'
    )

    def test_search_source_is_routed_to_the_engine(self):
        import json

        from jobhunt_ai.scraping import sources

        with override_settings(
            JOBHUNT_AI_SCOUT_SOURCES=json.loads(self.SEARCH_SOURCE)
        ):
            pages = list(sources.iter_search_pages(["devops"]))

        self.assertEqual(len(pages), 1)
        page = pages[0]
        self.assertEqual(page["mode"], "search")
        self.assertEqual(page["search"], "devops emploi Brabant wallon")
        self.assertEqual(page["geo_location"], "be")
        self.assertEqual(page["url"], "")

        with mock.patch("jobhunt_ai.conf.BRIGHTDATA_API_TOKEN", "jeton"), mock.patch(
            "jobhunt_ai.scraping.brightdata.search_engine", return_value="résultats"
        ) as search:
            self.assertEqual(sources.fetch_page(page), "résultats")
        search.assert_called_once_with(
            "devops emploi Brabant wallon", engine="google", geo_location="be"
        )

    def test_linkedin_and_indeed_are_default_sources(self):
        from jobhunt_ai import conf

        names = [source["name"] for source in conf.scout_sources()]
        self.assertIn("LinkedIn", names)
        self.assertIn("Indeed", names)

    def test_page_budget_covers_every_query_times_source(self):
        """Le plafond tronque la boucle requête × source : trop bas, la
        deuxième requête n'est jamais lue."""
        from jobhunt_ai import conf

        self.assertGreaterEqual(
            conf.SCOUT_MAX_PAGES,
            conf.SCOUT_MAX_QUERIES * len(conf.scout_sources()),
        )

    def test_location_and_radius_fill_the_templates(self):
        from jobhunt_ai.scraping import sources

        pages = {
            page["name"]: page["url"]
            for page in sources.iter_search_pages(["sre"], "Wavre, Belgique", 25)
        }
        # Encodés : la zone part dans une chaîne de requête.
        self.assertIn("location=Wavre%2C+Belgique", pages["LinkedIn"])
        self.assertIn("l=Wavre%2C+Belgique", pages["Indeed"])
        self.assertIn("radius=25", pages["Indeed"])
        # Un gabarit qui ignore ces emplacements reste valide.
        self.assertNotIn("Wavre", pages["ICTjob"])

    def test_location_falls_back_to_the_configured_default(self):
        from jobhunt_ai import conf
        from jobhunt_ai.scraping import sources

        pages = {p["name"]: p["url"] for p in sources.iter_search_pages(["sre"])}
        self.assertIn(quote_plus(conf.DEFAULT_LOCATION), pages["LinkedIn"])
        self.assertIn(f"radius={conf.DEFAULT_RADIUS_KM}", pages["Indeed"])

    def test_unknown_placeholder_names_the_faulty_source(self):
        import json

        from jobhunt_ai.scraping import sources

        bad = '[{"name": "Cassée", "url": "https://x.test/?q={query}&z={inconnu}"}]'
        with override_settings(JOBHUNT_AI_SCOUT_SOURCES=json.loads(bad)):
            with self.assertRaisesMessage(RuntimeError, "Cassée"):
                list(sources.iter_search_pages(["devops"]))

    def test_indeed_relative_links_resolve_against_the_search_page(self):
        """Indeed rend des liens relatifs (``/rc/clk?jk=…``), contrairement à
        LinkedIn : sans base ils seraient perdus."""
        from jobhunt_ai.scraping.sources import absolutize

        base = "https://be.indeed.com/emplois?q=devops&l=Nivelles"
        self.assertEqual(
            absolutize("/rc/clk?jk=abc123", base),
            "https://be.indeed.com/rc/clk?jk=abc123",
        )
        self.assertEqual(
            absolutize("https://be.linkedin.com/jobs/view/x-123", base),
            "https://be.linkedin.com/jobs/view/x-123",
        )

    def test_search_source_without_brightdata_says_why(self):
        from jobhunt_ai.scraping import sources

        page = {"mode": "search", "search": "devops", "engine": "google"}
        with mock.patch(
            "jobhunt_ai.conf.BRIGHTDATA_API_TOKEN", ""
        ), self.assertRaisesMessage(RuntimeError, "BRIGHTDATA_API_TOKEN"):
            sources.fetch_page(page)

    def test_url_source_keeps_escaping_the_query(self):
        from jobhunt_ai.scraping import sources

        pages = list(sources.iter_search_pages(["ingénieur devops"]))
        self.assertTrue(pages)
        self.assertEqual(pages[0]["mode"], "scrape")
        self.assertIn("ing%C3%A9nieur+devops", pages[0]["url"])

    def test_relative_url_from_a_search_result_is_dropped(self):
        """Une recherche n'a pas d'URL de page : sans base, un lien relatif ne
        peut pas être résolu et ne doit pas être conservé tel quel."""
        from jobhunt_ai.scraping.sources import absolutize

        self.assertEqual(absolutize("/fr/detail/1", ""), "")
        self.assertEqual(
            absolutize("https://www.ictjob.be/fr/detail/1", ""),
            "https://www.ictjob.be/fr/detail/1",
        )

    def test_scout_runs_end_to_end_on_a_search_source(self):
        import json

        profile = make_profile()
        with override_settings(
            JOBHUNT_AI_SCOUT_SOURCES=json.loads(self.SEARCH_SOURCE)
        ), mock.patch("jobhunt_ai.conf.BRIGHTDATA_API_TOKEN", "jeton"), mock.patch(
            "jobhunt_ai.scraping.brightdata.search_engine",
            return_value="1. DevOps Engineer — Widgets SA https://example.org/jobs/devops",
        ), fakes.fake_llm(), fakes.eager_runs():
            run = runner.launch(RunKind.SCOUT, owner=profile.owner, params={"keywords": "devops"})

        self.assertEqual(run.status, RunStatus.SUCCEEDED, run.error)
        lead = OfferLead.objects.get(title="DevOps Engineer")
        self.assertEqual(lead.source_name, "Google")
        self.assertEqual(lead.url, "https://example.org/jobs/devops")


class OwnerFollowTests(TestCase):
    def test_runs_follow_a_reassigned_application(self):
        application = make_application()
        run = AgentRun.objects.create(
            owner=application.owner, kind=RunKind.MATCH, status=RunStatus.SUCCEEDED,
            application=application,
        )
        other = make_user("Marie")
        application = Application.objects.get(pk=application.pk)
        application.owner = other
        application.company = Company.objects.create(owner=other, name="Acme bis")
        application.save()
        run.refresh_from_db()
        self.assertEqual(run.owner, other)


class RunnerTests(TestCase):
    def test_run_lifecycle_timestamps(self):
        make_profile()
        application = make_application()
        with fakes.fake_llm(), fakes.eager_runs():
            run = runner.launch(RunKind.MATCH, application=application)
        self.assertTrue(run.is_finished)
        self.assertIsNotNone(run.started_at)
        self.assertIsNotNone(run.finished_at)
        self.assertIsNotNone(run.duration_seconds)

    def test_unknown_kind_fails(self):
        run = AgentRun.objects.create(owner=make_user(), kind="unknown")
        runner.execute(run.pk, run.owner_id)
        run.refresh_from_db()
        self.assertEqual(run.status, RunStatus.FAILED)


class SweepOrphansTests(TestCase):
    """Deadlines belong to jobs and survive web-process restarts."""

    def test_an_expired_run_is_closed(self):
        user = make_user()
        run = AgentRun.objects.create(
            owner=user, kind=RunKind.MATCH, status=RunStatus.RUNNING,
            deadline_at=timezone.now() - timedelta(seconds=1),
        )
        runner.sweep_orphans(user)
        run.refresh_from_db()
        self.assertEqual(run.status, RunStatus.FAILED)
        self.assertIn("worker", run.error)
        self.assertIsNotNone(run.finished_at)

    def test_a_live_run_is_left_alone_even_if_created_long_ago(self):
        user = make_user()
        run = AgentRun.objects.create(
            owner=user, kind=RunKind.MATCH, status=RunStatus.RUNNING,
            deadline_at=timezone.now() + timedelta(minutes=10),
        )
        AgentRun.objects.filter(pk=run.pk).update(created_at=timezone.now() - timedelta(days=2))
        runner.sweep_orphans(user)
        run.refresh_from_db()
        self.assertEqual(run.status, RunStatus.RUNNING)

    def test_only_the_accounts_own_runs_are_touched(self):
        mine, theirs = make_user(), make_user()
        runs = {
            owner: AgentRun.objects.create(
                owner=owner, kind=RunKind.MATCH, status=RunStatus.RUNNING,
                deadline_at=timezone.now() - timedelta(seconds=1),
            )
            for owner in (mine, theirs)
        }
        runner.sweep_orphans(mine)
        self.assertEqual(AgentRun.objects.get(pk=runs[mine].pk).status, RunStatus.FAILED)
        self.assertEqual(AgentRun.objects.get(pk=runs[theirs].pk).status, RunStatus.RUNNING)


class ReviewRegressionTests(TestCase):
    """Verrous posés après la revue adversariale du 2026-08-31."""

    def test_generator_without_posting_text_fails(self):
        make_profile()
        application = make_application(posting_raw="", summary="", url="")
        with fakes.fake_llm(), fakes.eager_runs():
            run = runner.launch(
                RunKind.GENERATE_CV, params={"language": "fr"}, application=application
            )
        self.assertEqual(run.status, RunStatus.FAILED)
        self.assertIn("annonce", run.error)

    def test_scraped_javascript_url_is_neutralized(self):
        from jobhunt_ai.scraping.sources import absolutize, safe_http_url

        base = "https://www.ictjob.be/fr/rechercher-emplois-it?keywords=devops"
        self.assertEqual(absolutize("javascript:alert(1)", base), "")
        self.assertEqual(absolutize("JaVaScRiPt:alert(1)", base), "")
        self.assertEqual(absolutize("data:text/html,<script>1</script>", base), "")
        self.assertEqual(safe_http_url("ftp://example.org/x"), "")
        self.assertEqual(
            absolutize("/fr/detail/1", base), "https://www.ictjob.be/fr/detail/1"
        )

    def test_skill_gap_demand_counts_exact_names_only(self):
        from jobhunt_ai.services.applications import sync_skill_gaps
        from jobhunt_ai.agents.schemas import MatchVerdict, SkillGapItem

        # « Go » ne doit pas compter les rapports où manque « Django ».
        first = make_application(company_name="Alpha")
        second = make_application(company_name="Beta")
        MatchReport.objects.create(application=first, score=50, missing_skills=["Django"])
        MatchReport.objects.create(application=second, score=50, missing_skills=["Go"])
        verdict = MatchVerdict(
            score=50, summary="s", strengths="f", weaknesses="w", strategy="st",
            matched_skills=[], partial_skills=[], missing_skills=["Go"],
            gaps=[SkillGapItem(name="Go", severity="importante",
                               why_it_matters="x", action_plan="y")],
        )
        sync_skill_gaps(verdict, first.owner_id)
        self.assertEqual(SkillGap.objects.get(name="Go").demand_count, 1)

    def test_fetched_offer_text_is_persisted_once(self):
        from jobhunt_ai.agents.matcher import resolve_offer_text

        application = make_application(
            posting_raw="", summary="", url="https://example.org/offre"
        )
        page_text = "Un poste DevOps très détaillé. " * 10
        with mock.patch(
            "jobhunt_ai.agents.matcher.fetch_offer_text", return_value=page_text
        ) as fetch:
            resolve_offer_text(application)
            application.refresh_from_db()
            self.assertEqual(application.posting_raw, page_text)
            resolve_offer_text(application)
        fetch.assert_called_once()


class ThreadBindingTests(TestCase):
    """Chaque requête SQL d'une exécution est émise liée au compte.

    Le thread d'un agent ne descend d'aucune requête HTTP : Python le démarre
    avec un contexte vide, donc rien ne le relie au compte et, sur
    PostgreSQL, les politiques d'isolation ne lui montreraient aucune ligne —
    l'agent échouerait sans bruit plutôt qu'avec une erreur. Le corps du test
    est délibérément hors de tout ``as_user`` (comme le thread), et chaque
    requête qui s'échappe d'un bloc lié est collectée.

    Le contrôle vaut sur SQLite comme sur PostgreSQL : ``rls.as_user`` lie la
    variable de contexte sur les deux, et n'ouvre la transaction qui annonce
    le compte au serveur que sur PostgreSQL.
    """

    MEDIA = tempfile.mkdtemp(prefix="jobhunt-ai-binding-")

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(cls.MEDIA, ignore_errors=True)

    @contextlib.contextmanager
    def stray_queries(self, owner):
        stray: list[str] = []

        def wrapper(execute, sql, params, many, context):
            if rls.current_user_id() != owner.pk:
                stray.append(sql)
            return execute(sql, params, many, context)

        with connection.execute_wrapper(wrapper):
            yield stray

    def assert_stays_bound(self, run):
        with self.stray_queries(run.owner) as stray, fakes.eager_runs():
            runner.execute(run.pk, run.owner_id)
        run.refresh_from_db()
        self.assertEqual(run.status, RunStatus.SUCCEEDED, run.error)
        self.assertEqual(stray, [], f"{len(stray)} requête(s) hors du compte lié")

    def pending(self, kind, owner, **kwargs) -> AgentRun:
        """Une exécution en attente, comme ``launch`` la crée avant le thread."""
        return AgentRun.objects.create(owner=owner, kind=kind, **kwargs)

    def test_cv_parser_stays_bound(self):
        with override_settings(MEDIA_ROOT=self.MEDIA):
            document = make_cv_document()
            run = self.pending(
                RunKind.PARSE_CV,
                document.owner,
                # Le texte anonymisé voyage avec l'exécution, comme
                # ``hooks.CopilotCVAnalyzer`` le pose.
                params={"document_id": document.pk, "make_primary": True, "text": CV_TEXT},
            )
            with fakes.fake_llm():
                self.assert_stays_bound(run)
        self.assertEqual(CandidateProfile.objects.count(), 1)

    def test_matcher_stays_bound(self):
        profile = make_profile()
        application = make_application(owner=profile.owner)
        run = self.pending(RunKind.MATCH, profile.owner, application=application)
        with fakes.fake_llm():
            self.assert_stays_bound(run)
        self.assertEqual(MatchReport.objects.count(), 1)

    def test_generator_stays_bound(self):
        with override_settings(MEDIA_ROOT=self.MEDIA):
            profile = make_profile()
            application = make_application(owner=profile.owner)
            run = self.pending(
                RunKind.GENERATE_CV,
                profile.owner,
                application=application,
                params={"language": "fr"},
            )
            with fakes.fake_llm():
                self.assert_stays_bound(run)
        self.assertEqual(GeneratedCV.objects.count(), 1)

    def test_scout_stays_bound(self):
        profile = make_profile()
        run = self.pending(
            RunKind.SCOUT, profile.owner, params={"keywords": "devops"}
        )
        with fakes.fake_llm(), mock.patch(
            "jobhunt_ai.scraping.sources.fetch_page", return_value="des offres…"
        ):
            self.assert_stays_bound(run)
        self.assertTrue(OfferLead.objects.exists())

    def test_the_rubric_derivation_stays_bound(self):
        """La grille est dérivée à la première veille : lecture, appel au
        modèle, puis écriture sur le profil — chacune dans son bloc."""
        profile = make_profile(qualifications={})
        run = self.pending(
            RunKind.SCOUT, profile.owner, params={"keywords": "devops"}
        )
        with fakes.fake_llm(), mock.patch(
            "jobhunt_ai.scraping.sources.fetch_page", return_value="des offres…"
        ):
            self.assert_stays_bound(run)
        profile.refresh_from_db()
        self.assertTrue(profile.qualifications)
