"""Import the original spreadsheet and the Offres/ folder tree into the database.

Idempotent: run it twice and nothing is duplicated. Existing rows are matched on
(company, title) and updated, never cloned. Files are copied into MEDIA_ROOT so
the original folders stay exactly as they are.
"""

from __future__ import annotations

import datetime as dt
import re
import unicodedata
from pathlib import Path

from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.files import File
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils.text import slugify

from tracker.models import (
    Application,
    Company,
    Document,
    DocumentKind,
    EventKind,
    Language,
    Platform,
    Sector,
    SkillGap,
    Status,
    WorkMode,
)

SECTOR_MAP = {
    "privé": Sector.PRIVATE,
    "prive": Sector.PRIVATE,
    "parapublic": Sector.PARAPUBLIC,
    "public": Sector.PUBLIC,
}

LANGUAGE_MAP = {"fr": Language.FR, "en": Language.EN, "nl": Language.NL}

STATUS_MAP = {
    "à postuler": Status.TO_APPLY,
    "a postuler": Status.TO_APPLY,
    "à examiner": Status.BACKLOG,
    "a examiner": Status.BACKLOG,
    "candidature envoyée": Status.SENT,
    "envoyée": Status.SENT,
    "entretien": Status.INTERVIEW,
    "refusée": Status.REJECTED,
    "refus": Status.REJECTED,
    "sans réponse": Status.GHOSTED,
    "abandonnée": Status.WITHDRAWN,
}

#: Section headings used by the generated notes.md files.
NOTE_SECTIONS = {
    "ce qui joue pour toi": "strengths",
    "ce qui joue contre toi": "weaknesses",
    "comment aborder cette candidature": "strategy",
}

DOCUMENT_EXTENSIONS = {".docx", ".doc", ".pdf", ".odt", ".rtf", ".txt", ".md"}


def is_blank(value) -> bool:
    """Treat 'unknown' enum members as absent so they never clobber real data."""
    return value in (None, "", WorkMode.UNKNOWN, Sector.UNKNOWN)


def norm(value) -> str:
    return str(value or "").strip()


def squash(value: str) -> str:
    """Bare letters and digits, accents folded — for tolerant name matching."""
    return re.sub(r"[^a-z0-9]", "", fold(value))


def fold(value: str) -> str:
    """Lower-case and strip accents, for tolerant lookups."""
    text = unicodedata.normalize("NFKD", str(value or "").strip().lower())
    return "".join(c for c in text if not unicodedata.combining(c))


def parse_distance(value) -> int | None:
    match = re.search(r"(\d+)", norm(value))
    return int(match.group(1)) if match else None


def parse_score(value) -> int | None:
    if value in (None, ""):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if number <= 1:
        number *= 100
    return max(0, min(100, round(number)))


def parse_date(value) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = norm(value)
    for pattern in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return dt.datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


ON_SITE_HINTS = ("sur site", "sur place", "on site", "on-site", "j/sem", "jours sur")
REMOTE_HINTS = ("remote", "teletravail", "telework", "distanciel")
NEGATED_REMOTE = re.compile(
    r"\b(?:pas d[e']|aucune?|sans|no|non)\s+(?:\w+\s+){0,2}"
    r"(?:remote|teletravail|telework|distanciel)"
)


def guess_work_mode(*sources: str) -> str:
    """Read the work arrangement out of a location line.

    "5 jours sur site, pas de télétravail" mentions remote work only to rule it
    out, so a plain keyword match would get it exactly backwards.
    """
    haystack = fold(" ".join(s for s in sources if s))
    on_site = any(hint in haystack for hint in ON_SITE_HINTS)
    remote = any(hint in haystack for hint in REMOTE_HINTS)
    denied = remote and bool(NEGATED_REMOTE.search(haystack))

    if denied:
        return WorkMode.ON_SITE
    if "hybride" in haystack or "hybrid" in haystack:
        return WorkMode.HYBRID
    if remote and on_site:
        return WorkMode.HYBRID
    if remote:
        return WorkMode.REMOTE
    if on_site:
        return WorkMode.ON_SITE
    return WorkMode.UNKNOWN


def guess_language(filename: str) -> str:
    stem = fold(Path(filename).stem)
    if stem.endswith("_en") or "_en_" in stem or stem.endswith("-en"):
        return Language.EN
    if stem.endswith("_nl") or "_nl_" in stem:
        return Language.NL
    if stem.endswith("_fr") or "_fr_" in stem or "sante" in stem:
        return Language.FR
    return ""


def parse_notes(text: str) -> dict[str, str]:
    """Split a generated notes.md into its three analysis sections."""
    sections: dict[str, list[str]] = {}
    current: str | None = None
    for line in text.replace("\r\n", "\n").split("\n"):
        heading = re.match(r"^#{2,4}\s+(.*)$", line.strip())
        if heading:
            current = NOTE_SECTIONS.get(fold(heading.group(1)))
            if current:
                sections[current] = []
            continue
        if current and current in sections:
            sections[current].append(line)
    return {key: "\n".join(lines).strip() for key, lines in sections.items()}


class Command(BaseCommand):
    help = "Reprend le tableur 00_Suivi_candidatures.xlsx et le dossier Offres/ dans la base."

    def add_arguments(self, parser):
        parser.add_argument(
            "--workbook",
            default=str(settings.LEGACY_WORKBOOK),
            help="Chemin du classeur Excel à importer.",
        )
        parser.add_argument(
            "--root",
            default=str(settings.LEGACY_ROOT),
            help="Dossier contenant Offres/ et CV_base/.",
        )
        parser.add_argument(
            "--skip-files",
            action="store_true",
            help="N'importer que les données, sans copier les documents.",
        )
        parser.add_argument(
            "--user",
            default="",
            help="Identifiant du profil qui reçoit les données (facultatif s'il n'y en a qu'un).",
        )

    def resolve_owner(self, username: str):
        User = get_user_model()
        if username:
            user = User.objects.filter(username=username).first()
            if user is None:
                raise CommandError(f"Aucun profil nommé « {username} ».")
            return user
        users = list(User.objects.order_by("pk")[:3])
        if len(users) == 1:
            return users[0]
        if not users:
            raise CommandError(
                "Aucun profil : ouvre l'application une première fois, puis relance l'import."
            )
        names = ", ".join(u.get_username() for u in User.objects.order_by("username"))
        raise CommandError(f"Plusieurs profils : précise --user parmi {names}.")

    def handle(self, *args, **options):
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover - dependency is declared
            raise CommandError("openpyxl est requis : pip install openpyxl") from exc

        self.owner = self.resolve_owner(options["user"])
        workbook_path = Path(options["workbook"])
        root = Path(options["root"])
        if not workbook_path.exists():
            raise CommandError(f"Classeur introuvable : {workbook_path}")

        workbook = openpyxl.load_workbook(workbook_path, data_only=True)
        self.counts = {
            "companies": 0, "applications": 0, "updated": 0,
            "platforms": 0, "gaps": 0, "documents": 0,
        }

        with transaction.atomic():
            platforms = self.import_platforms(workbook)
            self.import_applications(workbook, platforms)
            self.import_discarded(workbook)
            self.import_gaps(workbook)
            if not options["skip_files"]:
                self.import_offer_folders(root)
                self.import_base_cvs(root)

        self.stdout.write("")
        self.stdout.write(self.style.SUCCESS("Import terminé."))
        for label, key in [
            ("sociétés", "companies"),
            ("candidatures créées", "applications"),
            ("candidatures mises à jour", "updated"),
            ("plateformes", "platforms"),
            ("lacunes", "gaps"),
            ("documents", "documents"),
        ]:
            self.stdout.write(f"  {self.counts[key]:>4}  {label}")

    # -- sheets ------------------------------------------------------------

    def sheet(self, workbook, *names):
        for name in names:
            for title in workbook.sheetnames:
                if fold(title) == fold(name):
                    return workbook[title]
        return None

    def rows(self, sheet, header_row: int = 1):
        """Yield dicts keyed by the folded header label."""
        headers = [
            fold(cell.value) for cell in next(sheet.iter_rows(min_row=header_row, max_row=header_row))
        ]
        for row in sheet.iter_rows(min_row=header_row + 1):
            values = [cell.value for cell in row]
            if not any(norm(v) for v in values):
                continue
            yield dict(zip(headers, values))

    def get_company(self, name: str, sector: str = "", location: str = "") -> Company:
        name = norm(name)
        company = Company.objects.for_user(self.owner).filter(name__iexact=name).first()
        if company is None:
            company = Company.objects.create(
                owner=self.owner,
                name=name,
                sector=SECTOR_MAP.get(fold(sector), Sector.UNKNOWN),
                location=norm(location),
            )
            self.counts["companies"] += 1
        else:
            changed = []
            if sector and company.sector == Sector.UNKNOWN:
                company.sector = SECTOR_MAP.get(fold(sector), Sector.UNKNOWN)
                changed.append("sector")
            if location and not company.location:
                company.location = norm(location)
                changed.append("location")
            if changed:
                company.save(update_fields=changed)
        return company

    def import_platforms(self, workbook) -> dict[str, Platform]:
        sheet = self.sheet(workbook, "Plateformes explorées", "Plateformes")
        registry: dict[str, Platform] = {}
        if sheet is None:
            return registry

        for row in self.rows(sheet):
            name = norm(row.get("plateforme"))
            if not name:
                continue
            is_lead = fold(name).startswith("pistes non")
            platform, created = Platform.objects.update_or_create(
                owner=self.owner,
                name=name,
                defaults={
                    "searched_for": norm(row.get("ce que j'y ai cherche"))
                    or norm(row.get("ce que j'y ai cherché")),
                    "outcome": norm(row.get("resultat")) or norm(row.get("résultat")),
                    "is_lead": is_lead,
                },
            )
            if created:
                self.counts["platforms"] += 1
            registry[fold(name)] = platform
        return registry

    def match_platform(self, label: str, registry: dict[str, Platform]) -> Platform | None:
        """The spreadsheet writes 'LinkedIn / Indeed' — take the first channel."""
        first = fold(label).split("/")[0].strip()
        if not first:
            return None
        if first in registry:
            return registry[first]
        for key, platform in registry.items():
            if key.startswith(first) or first.startswith(key):
                return platform
        platform = Platform.objects.for_user(self.owner).filter(name__iexact=first).first()
        if platform is None:
            platform = Platform.objects.create(owner=self.owner, name=first.title())
            self.counts["platforms"] += 1
            registry[first] = platform
        return platform

    def import_applications(self, workbook, platforms):
        sheet = self.sheet(workbook, "Candidatures")
        if sheet is None:
            raise CommandError("Onglet « Candidatures » introuvable.")

        for row in self.rows(sheet):
            company_name = norm(row.get("societe")) or norm(row.get("société"))
            title = norm(row.get("intitule du poste")) or norm(row.get("intitulé du poste"))
            # The summary block at the bottom of the sheet has no company column.
            if not company_name or not title:
                continue

            sector = norm(row.get("secteur"))
            location = norm(row.get("lieu"))
            company = self.get_company(company_name, sector, location)

            status_label = fold(row.get("statut"))
            defaults = {
                "location": location,
                "distance_km": parse_distance(
                    row.get("distance nivelles") or row.get("distance")
                ),
                "score": parse_score(row.get("score")),
                "url": norm(row.get("lien de l'offre")),
                "cv_language": LANGUAGE_MAP.get(
                    fold(row.get("langue du cv")), Language.FR
                ),
                "status": STATUS_MAP.get(status_label, Status.TO_APPLY),
                "summary": norm(row.get("notes")),
                "applied_on": parse_date(row.get("date de candidature")),
                "follow_up_on": parse_date(row.get("relance prevue"))
                or parse_date(row.get("relance prévue")),
                "source_platform": self.match_platform(
                    norm(row.get("plateforme")), platforms
                ),
                "work_mode": guess_work_mode(location),
            }
            # The sheet's "CV personnalisé" column points at a file; the folder
            # pass sets legacy_folder to the directory, which is the useful one.
            self.upsert_application(company, title, defaults)

    def import_discarded(self, workbook):
        sheet = self.sheet(workbook, "Offres écartées", "Offres ecartees")
        if sheet is None:
            return
        for row in self.rows(sheet):
            company_name = norm(row.get("societe")) or norm(row.get("société"))
            title = norm(row.get("poste"))
            if not company_name or not title:
                continue
            location = norm(row.get("lieu"))
            company = self.get_company(company_name, norm(row.get("secteur")), location)
            reason = (
                norm(row.get("pourquoi je ne l'ai pas retenue"))
                or norm(row.get("raison"))
            )
            self.upsert_application(
                company,
                title,
                {
                    "location": location,
                    "status": Status.DISCARDED,
                    "discard_reason": reason,
                    "work_mode": guess_work_mode(location),
                },
            )

    def import_gaps(self, workbook):
        sheet = self.sheet(workbook, "Lacunes à combler", "Lacunes a combler", "Lacunes")
        if sheet is None:
            return
        for position, row in enumerate(self.rows(sheet)):
            name = norm(row.get("competence manquante")) or norm(
                row.get("compétence manquante")
            )
            if not name:
                continue
            demand_label = norm(row.get("combien d'offres la demandent"))
            match = re.search(r"(\d+)", demand_label)
            _, created = SkillGap.objects.update_or_create(
                owner=self.owner,
                name=name,
                defaults={
                    "demand_count": int(match.group(1)) if match else 0,
                    "demand_label": demand_label,
                    "why_it_matters": norm(row.get("pourquoi ca compte"))
                    or norm(row.get("pourquoi ça compte")),
                    "action_plan": norm(row.get("piste")),
                    "position": position,
                },
            )
            if created:
                self.counts["gaps"] += 1

    def upsert_application(self, company: Company, title: str, defaults: dict) -> Application:
        application = Application.objects.for_user(self.owner).filter(
            company=company, title__iexact=title
        ).first()
        if application is None:
            application = Application(owner=self.owner, company=company, title=title)
            for field, value in defaults.items():
                if not is_blank(value):
                    setattr(application, field, value)
            application.save()
            application.log(
                EventKind.NOTE,
                "Reprise du tableur",
                "Ligne importée depuis 00_Suivi_candidatures.xlsx.",
            )
            self.counts["applications"] += 1
            return application

        changed = []
        for field, value in defaults.items():
            if is_blank(value):
                continue
            if getattr(application, field) != value:
                setattr(application, field, value)
                changed.append(field)
        if changed:
            application.save()
            self.counts["updated"] += 1
        return application

    # -- files -------------------------------------------------------------

    def import_offer_folders(self, root: Path):
        offers_dir = root / "Offres"
        if not offers_dir.is_dir():
            self.stdout.write(self.style.WARNING(f"Dossier Offres/ absent sous {root}."))
            return

        for folder in sorted(p for p in offers_dir.iterdir() if p.is_dir()):
            application = self.match_folder(folder)
            if application is None:
                self.stdout.write(
                    self.style.WARNING(f"  · aucun rapprochement pour {folder.name}")
                )
                continue

            application.legacy_folder = str(folder.relative_to(root))

            notes_file = folder / "notes.md"
            if notes_file.exists():
                text = notes_file.read_text(encoding="utf-8", errors="replace")
                application.notes_raw = text
                for field, value in parse_notes(text).items():
                    if value:
                        setattr(application, field, value)
                discovered = re.search(r"Relevé le\*{0,2}\s*:?\s*(\d{1,2}\s+\w+\s+\d{4})", text)
                if discovered and not application.discovered_on:
                    application.discovered_on = self.parse_french_date(discovered.group(1))
                # Only the "Lieu" line describes the arrangement; the rest of the
                # note discusses it and would poison a whole-document match.
                place = re.search(r"\*{0,2}Lieu\*{0,2}\s*:?\s*(.+)", text)
                if place:
                    mode = guess_work_mode(place.group(1), application.location)
                    if mode != WorkMode.UNKNOWN:
                        application.work_mode = mode

            posting_file = folder / "annonce_originale.md"
            if posting_file.exists():
                application.posting_raw = posting_file.read_text(
                    encoding="utf-8", errors="replace"
                )

            application.save()
            self.attach_documents(application, folder)

    def match_folder(self, folder: Path) -> Application | None:
        """Folders are named NN_Company_Job-Title; match on the company name.

        Folder tokens drop the spaces the spreadsheet keeps ("EgovSelect" for
        "Egov Select"), so both sides are reduced to bare letters and digits
        before comparing.
        """
        parts = folder.name.split("_", 2)
        if len(parts) < 2:
            return None
        company_token = squash(parts[1])
        title_words = set(slugify(parts[2].replace("-", " ")).split("-")) if len(parts) > 2 else set()

        candidates = []
        for application in Application.objects.for_user(self.owner).select_related("company"):
            company_key = squash(application.company.name)
            if not company_key or not company_token:
                continue
            head = min(len(company_key), len(company_token), 10)
            if company_key[:head] == company_token[:head]:
                candidates.append(application)

        if not candidates:
            return None
        if len(candidates) == 1:
            return candidates[0]

        # Several offers from the same company: pick the closest job title.
        def overlap(application: Application) -> int:
            return len(set(slugify(application.title).split("-")) & title_words)

        best = max(candidates, key=overlap)
        return best if overlap(best) else candidates[0]

    def attach_documents(self, application: Application, folder: Path):
        files = [
            path
            for path in sorted(folder.iterdir())
            if path.is_file()
            and path.suffix.lower() in DOCUMENT_EXTENSIONS
            and path.name not in {"notes.md", "annonce_originale.md"}
        ]
        stems = {path.stem for path in files}
        has_primary = application.documents.filter(
            kind=DocumentKind.CV, is_primary=True
        ).exists()

        for path in files:
            # Skip Finder-style duplicates ("… EN 2.docx") when the original is there.
            duplicate = re.match(r"^(.*) \d+$", path.stem)
            if duplicate and duplicate.group(1) in stems:
                continue
            is_cv = "cv" in fold(path.stem)
            self.store_document(
                path,
                application=application,
                kind=DocumentKind.CV if is_cv else DocumentKind.OTHER,
                is_primary=is_cv and not has_primary,
            )
            if is_cv:
                has_primary = True

    def import_base_cvs(self, root: Path):
        base_dir = root / "CV_base"
        if not base_dir.is_dir():
            return
        for path in sorted(base_dir.iterdir()):
            if path.is_file() and path.suffix.lower() in DOCUMENT_EXTENSIONS:
                self.store_document(path, application=None, kind=DocumentKind.CV)

    def store_document(self, path: Path, *, application, kind, is_primary=False):
        source = str(path)
        if Document.objects.for_user(self.owner).filter(source_path=source).exists():
            return
        label = path.name
        with path.open("rb") as handle:
            document = Document(
                owner=self.owner,
                application=application,
                kind=kind,
                label=label,
                language=guess_language(path.name),
                is_primary=is_primary,
                source_path=source,
            )
            document.file.save(path.name, File(handle), save=False)
            document.save()
        self.counts["documents"] += 1

    @staticmethod
    def parse_french_date(text: str) -> dt.date | None:
        months = {
            "janvier": 1, "fevrier": 2, "mars": 3, "avril": 4, "mai": 5, "juin": 6,
            "juillet": 7, "aout": 8, "septembre": 9, "octobre": 10,
            "novembre": 11, "decembre": 12,
        }
        match = re.match(r"(\d{1,2})\s+(\w+)\s+(\d{4})", fold(text))
        if not match:
            return None
        month = months.get(match.group(2))
        if not month:
            return None
        try:
            return dt.date(int(match.group(3)), month, int(match.group(1)))
        except ValueError:
            return None
