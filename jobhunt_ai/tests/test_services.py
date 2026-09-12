"""Tests des services : extraction de documents, rendu ATS, veille."""

from __future__ import annotations

import tempfile
from unittest import mock

import docx
from django.test import TestCase

from jobhunt_ai.services import ats, documents
from jobhunt_ai.scraping import sources
from jobhunt_ai.tests import fakes
from jobhunt_ai.tests.test_agents import make_application


class DocumentFormatTests(TestCase):
    """L'extension ne lit plus un document : elle ne fait que refuser au
    téléversement ce que l'extracteur du cœur ne saurait pas relire."""

    def test_accepted_formats_follow_the_core_extractor(self):
        from tracker.adapters import document_text

        self.assertEqual(documents.ALLOWED_EXTENSIONS, document_text.SUPPORTED_SUFFIXES)
        self.assertTrue(issubclass(documents.UnsupportedFormat, ValueError))
        for name in ("cv.pdf", "CV.DOCX", "cv.txt", "notes.md"):
            with self.subTest(name=name):
                self.assertIn(documents.check_extension(name), documents.ALLOWED_EXTENSIONS)

    def test_unsupported_extension_raises_and_lists_the_accepted_ones(self):
        with self.assertRaises(documents.UnsupportedFormat) as raised:
            documents.check_extension("cv.rtf")
        self.assertIn("DOCX", str(raised.exception))
        with self.assertRaises(documents.UnsupportedFormat):
            documents.check_extension("cv")

    def test_the_readable_threshold_is_the_core_s(self):
        from tracker import services

        self.assertEqual(documents.MIN_TEXT_LENGTH, services.MIN_TEXT_LENGTH)


class ATSRenderTests(TestCase):
    def test_renders_all_sections_in_french(self):
        payload = ats.render_docx(fakes.ATS_RESUME, language="fr")
        with tempfile.NamedTemporaryFile(suffix=".docx") as handle:
            handle.write(payload)
            handle.flush()
            rendered = docx.Document(handle.name)
        text = "\n".join(p.text for p in rendered.paragraphs)
        for expected in (
            "Lionel Test",
            "Profil",
            "Compétences",
            "Expérience professionnelle",
            "Formation",
            "Certifications",
            "Langues",
            "Migration de 200 serveurs sans coupure",
        ):
            self.assertIn(expected, text)

    def test_renders_english_titles(self):
        payload = ats.render_docx(fakes.ATS_RESUME, language="en")
        with tempfile.NamedTemporaryFile(suffix=".docx") as handle:
            handle.write(payload)
            handle.flush()
            rendered = docx.Document(handle.name)
        text = "\n".join(p.text for p in rendered.paragraphs)
        self.assertIn("Professional Experience", text)
        self.assertNotIn("Expérience professionnelle", text)


class ScrapingTests(TestCase):
    def test_iter_search_pages_crosses_sources_and_queries(self):
        with mock.patch(
            "jobhunt_ai.conf.scout_sources",
            return_value=[{"name": "Test", "url": "https://t.example/s?q={query}"}],
        ):
            pages = list(sources.iter_search_pages(["devops engineer", "sre"]))
        self.assertEqual(len(pages), 2)
        self.assertEqual(pages[0]["url"], "https://t.example/s?q=devops+engineer")

    def test_absolutize(self):
        base = "https://www.ictjob.be/fr/rechercher-emplois-it?keywords=devops"
        self.assertEqual(
            sources.absolutize("/fr/detail/123", base),
            "https://www.ictjob.be/fr/detail/123",
        )
        self.assertEqual(sources.absolutize("", base), "")
        self.assertEqual(
            sources.absolutize("https://autre.example/x", base), "https://autre.example/x"
        )

    def test_deduplicate_against_tracked_applications(self):
        application = make_application(url="https://example.org/jobs/1", title="DevOps Engineer")
        offers = [
            {"title": "DevOps Engineer", "company_name": "Acme", "url": "https://example.org/jobs/1"},
            {"title": "devops engineer", "company_name": "ACME", "url": ""},
            {"title": "SRE", "company_name": "Widgets", "url": "https://example.org/jobs/2"},
            {"title": "SRE", "company_name": "Widgets", "url": "https://example.org/jobs/2"},
            {"title": "", "company_name": "Sans titre"},
        ]
        kept = sources.deduplicate(offers, application.owner_id)
        self.assertEqual(len(kept), 1)
        self.assertEqual(kept[0]["title"], "SRE")
