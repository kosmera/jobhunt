"""Tests for the row-level security layer.

The policy tests need PostgreSQL (``JOBHUNT_DATABASE_URL=postgres://…``, see
the README) and are skipped elsewhere; the context, registry, check and
middleware tests run on any engine.
"""

from __future__ import annotations

import io
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest import mock, skipUnless

from django.apps import apps
from django.conf import settings
from django.contrib.auth.models import AnonymousUser, Group, User
from django.contrib.sessions.backends.db import SessionStore
from django.contrib.sessions.models import Session
from django.core.exceptions import ImproperlyConfigured
from django.core.management import call_command
from django.db import DatabaseError, connection, transaction
from django.http import HttpResponse
from django.test import (
    Client,
    RequestFactory,
    SimpleTestCase,
    TestCase,
    TransactionTestCase,
    override_settings,
)
from django.urls import reverse
from openpyxl import Workbook

from accounts.models import Preferences, Profile, SearchProfile
from accounts.testing import make_user
from rls import checks, context, registry, sql, verify
from rls.middleware import RowLevelSecurityMiddleware
from rls.operations import EnableRowLevelSecurity
from rls.testing import AppRoleClient, AppRoleTestCase, app_role
from tracker.models import ActivityEvent, Application, Company, Contact

POSTGRES = connection.vendor == "postgresql"


def foreign(problem: str) -> bool:
    """A problem about an installed extension that has not adopted rls yet."""
    return "jobhunt_ai" in problem


def make_application(owner, company_name: str, title: str = "Poste") -> Application:
    company, _ = Company.objects.get_or_create(owner=owner, name=company_name)
    application = Application.objects.create(owner=owner, company=company, title=title)
    ActivityEvent.objects.create(application=application, title=f"Note {company_name}")
    Contact.objects.create(application=application, name=f"Contact {company_name}")
    return application


# ---------------------------------------------------------------------------
# Engine-independent: context, registry, SQL, checks, middleware
# ---------------------------------------------------------------------------


class ContextTests(TestCase):
    def test_unscoped_by_default(self):
        self.assertIsNone(context.current_user_id())
        self.assertFalse(context.is_scoped())

    def test_bound_sets_and_resets(self):
        with context.bound(7) as user_id:
            self.assertEqual(user_id, 7)
            self.assertEqual(context.current_user_id(), 7)
            self.assertTrue(context.is_scoped())
        self.assertIsNone(context.current_user_id())
        self.assertFalse(context.is_scoped())

    def test_user_id_of_accepts_users_ints_and_nobody(self):
        user = make_user("Alice")
        self.assertEqual(context.user_id_of(user), user.pk)
        self.assertEqual(context.user_id_of(3), 3)
        self.assertIsNone(context.user_id_of(None))
        self.assertIsNone(context.user_id_of(User()))

    def test_rebind_only_inside_a_scope(self):
        context.rebind(5)
        self.assertIsNone(context.current_user_id())
        with context.bound(None):
            context.rebind(5)
            self.assertEqual(context.current_user_id(), 5)
        self.assertIsNone(context.current_user_id())

    def test_as_user_binds_and_restores(self):
        with context.as_user(1):
            self.assertEqual(context.current_user_id(), 1)
            with context.as_user(2):
                self.assertEqual(context.current_user_id(), 2)
            self.assertEqual(context.current_user_id(), 1)
        self.assertIsNone(context.current_user_id())

    def test_rebind_is_refused_in_a_nested_scope(self):
        with context.as_user(None):
            context.rebind(1)
            self.assertEqual(context.current_user_id(), 1)
            with context.as_user(2):
                with self.assertRaises(ImproperlyConfigured):
                    context.rebind(3)
                self.assertEqual(context.current_user_id(), 2)
            self.assertEqual(context.current_user_id(), 1)

    def test_on_commit_keeps_the_account(self):
        seen = []
        with self.captureOnCommitCallbacks(execute=True):
            with context.as_user(3):
                context.on_commit(lambda: seen.append(context.current_user_id()))
        self.assertEqual(seen, [3])


class RegistryTests(SimpleTestCase):
    def test_rules_cover_the_project(self):
        labels = {rule.label.lower() for rule in registry.rules()}
        for label in (
            "auth.user",
            "accounts.profile",
            "accounts.preferences",
            "accounts.searchprofile",
            "tracker.application",
            "tracker.document",
            "tracker.activityevent",
            "tracker.contact",
        ):
            self.assertIn(label, labels)

    def test_rule_needs_owner_or_via(self):
        with self.assertRaises(ValueError):
            registry.Rule(label="x.Y")
        with self.assertRaises(ValueError):
            registry.Rule(label="x.Y", owner="a", via="b")

    def test_rule_for_model(self):
        rule = registry.rule_for(Application)
        assert rule is not None
        self.assertEqual(rule.owner, "owner")
        self.assertEqual(rule.table, "tracker_application")
        via = registry.rule_for(Contact)
        assert via is not None
        self.assertEqual(via.via, "application")


class ReportSideTests(SimpleTestCase):
    """Which side a connection is on, before and after the first migration."""

    def role(self, **overrides):
        base = dict(
            name="r", superuser=False, bypass_rls=False, create_role=False, create_db=False,
            app_role="jobhunt_app", app_role_exists=False, can_set_app_role=False,
            session_value=None, preset=False, is_app_role=False, owns_database=False,
        )
        return verify.RoleReport(**{**base, **overrides})

    def empty_tables(self):
        return [verify.TableReport(rule.label, rule.table) for rule in registry.rules()]

    def test_the_admin_of_an_empty_database_is_maintenance_side(self):
        # The Azure admin: no superuser, no table yet — its first migrate must not be refused.
        self.assertTrue(verify.Report(self.role(owns_database=True), self.empty_tables()).maintenance_side)
        self.assertTrue(verify.Report(self.role(create_role=True), self.empty_tables()).maintenance_side)

    def test_the_runtime_role_stays_runtime_side_on_an_empty_database(self):
        self.assertFalse(verify.Report(self.role(), self.empty_tables()).maintenance_side)

    def test_once_tables_exist_only_ownership_counts(self):
        tables = self.empty_tables()
        tables[0].exists = True
        self.assertFalse(verify.Report(self.role(owns_database=True), tables).maintenance_side)
        self.assertTrue(verify.Report(self.role(owns_database=True), tables, owned=["x"]).maintenance_side)


class SqlTests(SimpleTestCase):
    def test_owner_policy(self):
        statements = sql.enable_statements("t", sql.owner_predicate("owner_id"))
        self.assertEqual(statements[0], 'ALTER TABLE "t" ENABLE ROW LEVEL SECURITY')
        self.assertIn("FOR ALL TO PUBLIC", statements[2])
        self.assertIn(
            "\"owner_id\" = NULLIF(current_setting('app.current_user_id', true), '')::bigint",
            statements[2],
        )
        self.assertIn("WITH CHECK", statements[2])

    def test_via_policy_lists_the_parent_rows(self):
        predicate = sql.via_predicate("application_id", "tracker_application", "id", "owner_id")
        self.assertTrue(predicate.startswith('"application_id" = ANY (ARRAY(SELECT "id" FROM "tracker_application"'))
        self.assertIn('WHERE "owner_id" = NULLIF(', predicate)

    def test_unbound_visible_policy(self):
        predicate = sql.unbound_visible_predicate("id")
        self.assertTrue(predicate.startswith("(NULLIF("))
        self.assertIn("IS NULL OR", predicate)

    def test_unbound_visible_table_gets_one_policy_per_command(self):
        statements = sql.enable_unbound_visible_statements("auth_user", "id")
        commands = [s.split(" FOR ")[1].split(" TO ")[0] for s in statements if s.startswith("CREATE POLICY")]
        self.assertEqual(commands, ["SELECT", "INSERT", "UPDATE", "DELETE"])
        delete = [s for s in statements if " FOR DELETE " in s][0]
        self.assertNotIn("IS NULL OR", delete)
        self.assertIn('USING ("id" = NULLIF(', delete)
        self.assertEqual(sum(1 for s in statements if s.startswith("DROP POLICY")), 4)

    def test_bookkeeping_tables_lose_write_privileges(self):
        block = [s for s in sql.application_role_statements("jobhunt_app") if "REVOKE" in s][0]
        for table in sql.BOOKKEEPING_TABLES:
            self.assertIn(f'REVOKE INSERT, UPDATE, DELETE, TRUNCATE ON TABLE "{table}" FROM jobhunt_app', block)
        self.assertIn("to_regclass", block)

    @skipUnless(apps.is_installed("django_q"), "the queue tables come with the copilot")
    def test_queue_tables_are_allowed_by_name(self):
        # The names must match the models: the allow-list works without them.
        queue = {apps.get_model("django_q", name)._meta.db_table for name in ("OrmQ", "Schedule", "Task")}
        self.assertEqual(set(sql.QUEUE_TABLES), queue)
        self.assertTrue(queue <= set(sql.UNPROTECTED_ALLOWED))

    def test_role_name_is_validated(self):
        self.assertEqual(sql.role_name("jobhunt_app"), "jobhunt_app")
        for bad in ("Jobhunt", "app;drop", "app role", ""):
            with self.subTest(bad=bad), self.assertRaises(ValueError):
                sql.role_name(bad)

    def test_role_statements_create_grant_and_allow_set_role(self):
        statements = sql.application_role_statements("jobhunt_app")
        self.assertIn("CREATE ROLE jobhunt_app NOLOGIN NOSUPERUSER", statements[0])
        self.assertTrue(any("ALTER DEFAULT PRIVILEGES" in s and "ON TABLES" in s for s in statements))
        # The SET ROLE grant is best effort: the owner may not administer the role.
        self.assertIn("'GRANT jobhunt_app TO ' || quote_ident(current_user) || ' WITH SET TRUE, INHERIT FALSE'", statements[-1])
        self.assertIn("EXCEPTION WHEN insufficient_privilege", statements[-1])
        # Executed by the schema editor without parameters: no placeholder may appear.
        self.assertNotIn("%", "".join(statements))


class OperationTests(SimpleTestCase):
    def test_predicate_resolves_columns_from_the_model_state(self):
        from django.apps import apps

        table, predicate = EnableRowLevelSecurity("tracker.Application", owner="owner").predicate(apps)
        self.assertEqual(table, "tracker_application")
        self.assertIn('"owner_id" =', predicate)
        table, predicate = EnableRowLevelSecurity("tracker.Contact", via="application").predicate(apps)
        self.assertEqual(table, "tracker_contact")
        self.assertIn('FROM "tracker_application" WHERE "owner_id" =', predicate)
        statements = EnableRowLevelSecurity("auth.User", owner="id", unbound_visible=True).statements(apps)
        self.assertTrue(any('ON "auth_user" AS PERMISSIVE FOR DELETE' in s for s in statements))

    def test_deconstruct_round_trips(self):
        operation = EnableRowLevelSecurity("tracker.Contact", via="application", parent_owner="owner")
        name, args, kwargs = operation.deconstruct()
        self.assertEqual((name, args), ("EnableRowLevelSecurity", ["tracker.Contact"]))
        self.assertEqual(kwargs, {"unbound_visible": False, "via": "application", "parent_owner": "owner"})


class CheckTests(SimpleTestCase):
    def test_configuration_is_valid(self):
        self.assertEqual(checks.check_configuration(None), [])

    def test_every_model_referencing_the_account_is_covered(self):
        # An out-of-tree extension may be installed without a rule yet: only
        # the project's own apps are held to the check here.
        ours = ("auth.", "admin.", "sessions.", "contenttypes.", "accounts.", "tracker.")
        uncovered = [w.msg for w in checks.check_coverage(None) if w.msg.startswith(ours)]
        self.assertEqual(uncovered, [])

    def test_autocommit_off_is_refused_with_the_rls_engine(self):
        import warnings

        databases = {"default": {"ENGINE": checks.ENGINE, "NAME": "x", "AUTOCOMMIT": False}}
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            with override_settings(DATABASES=databases):
                ids = [p.id for p in checks.check_configuration(None)]
        self.assertIn("rls.E003", ids)

    def test_missing_or_misplaced_middleware(self):
        base = [m for m in settings.MIDDLEWARE if "rls." not in m]
        with override_settings(MIDDLEWARE=base):
            ids = [p.id for p in checks.check_configuration(None)]
        self.assertEqual(ids, ["rls.E001"])
        swapped = list(base)
        swapped.insert(swapped.index("accounts.middleware.AccountsMiddleware") + 1, checks.MIDDLEWARE)
        with override_settings(MIDDLEWARE=swapped):
            ids = [p.id for p in checks.check_configuration(None)]
        self.assertEqual(ids, ["rls.E002"])

    def test_rls_must_precede_auth_in_installed_apps(self):
        apps_list = list(settings.INSTALLED_APPS)
        apps_list.remove("rls")
        apps_list.append("rls")
        fake = SimpleNamespace(
            MIDDLEWARE=settings.MIDDLEWARE, DATABASES=settings.DATABASES, INSTALLED_APPS=apps_list
        )
        with mock.patch.object(checks, "settings", fake):
            ids = [p.id for p in checks.check_configuration(None)]
        self.assertEqual(ids, ["rls.E005"])

    def test_coverage_follows_relations_to_protected_models(self):
        rule = registry._rules.pop("tracker.contact")
        try:
            warnings = [w.msg for w in checks.check_coverage(None) if "tracker.Contact" in w.msg]
        finally:
            registry._rules["tracker.contact"] = rule
        self.assertEqual(len(warnings), 1)
        self.assertIn("référence tracker.Application", warnings[0])


@override_settings(RLS_ENFORCE=False)
class MiddlewareTests(TestCase):
    """The middleware driven by hand, on whatever role the test connection has."""

    def setUp(self):
        self.factory = RequestFactory()
        self.user = make_user("Alice")

    def run_middleware(self, request, get_response):
        if not hasattr(request, "user"):
            request.user = AnonymousUser()
        return RowLevelSecurityMiddleware(get_response)(request)

    def test_requires_authentication_middleware(self):
        request = self.factory.get("/")
        with self.assertRaises(ImproperlyConfigured):
            RowLevelSecurityMiddleware(lambda r: HttpResponse())(request)

    def test_binds_the_signed_in_account_for_the_view(self):
        seen = {}

        def view(request):
            seen["user_id"] = context.current_user_id()
            seen["scoped"] = context.is_scoped()
            return HttpResponse("ok")

        request = self.factory.get("/")
        request.user = self.user
        self.run_middleware(request, view)
        self.assertEqual(seen, {"user_id": self.user.pk, "scoped": True})
        self.assertFalse(context.is_scoped())

    def test_anonymous_request_is_scoped_to_nobody(self):
        seen = {}

        def view(request):
            seen["user_id"] = context.current_user_id()
            seen["scoped"] = context.is_scoped()
            return HttpResponse("ok")

        self.run_middleware(self.factory.get("/"), view)
        self.assertEqual(seen, {"user_id": None, "scoped": True})

    def test_signed_in_responses_are_never_cached_and_fragments_vary(self):
        request = self.factory.get("/", HTTP_HX_REQUEST="true")
        request.user = self.user
        response = self.run_middleware(request, lambda r: HttpResponse("fragment"))
        self.assertIn("HX-Request", response.headers["Vary"])
        self.assertEqual(response.headers["Cache-Control"], "private, no-store")
        page = self.factory.get("/")
        page.user = self.user
        response = self.run_middleware(page, lambda r: HttpResponse("page"))
        self.assertEqual(response.headers["Cache-Control"], "private, no-store")
        self.assertNotIn("Vary", response.headers)
        anonymous = self.run_middleware(self.factory.get("/"), lambda r: HttpResponse("public"))
        self.assertNotIn("Cache-Control", anonymous.headers)

    def test_a_view_exception_marks_the_transaction_for_rollback(self):
        request = self.factory.get("/")
        request.user = self.user
        middleware = RowLevelSecurityMiddleware(lambda r: HttpResponse())
        self.assertTrue(connection.in_atomic_block)  # the test case's own transaction
        # Not this middleware's transaction: left alone.
        self.assertIsNone(middleware.process_exception(request, RuntimeError("x")))
        self.assertFalse(transaction.get_rollback())
        setattr(request, "rls_transaction", True)
        self.assertIsNone(middleware.process_exception(request, RuntimeError("x")))
        self.assertTrue(transaction.get_rollback())
        transaction.set_rollback(False)

    def test_a_view_cache_header_is_kept(self):
        request = self.factory.get("/")
        request.user = self.user

        def view(request):
            response = HttpResponse("x")
            response.headers["Cache-Control"] = "max-age=60"
            return response

        self.assertEqual(self.run_middleware(request, view).headers["Cache-Control"], "max-age=60")


class CsrfFailureTests(TestCase):
    def test_htmx_request_with_a_stale_token_is_told_to_reload(self):
        user = make_user("Alice")
        client = self.client_class(enforce_csrf_checks=True)
        client.force_login(user)
        response = client.post(reverse("accounts:settings"), {"display_name": "A"}, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.headers["HX-Refresh"], "true")
        page = client.post(reverse("accounts:settings"), {"display_name": "A"})
        self.assertEqual(page.status_code, 403)
        self.assertNotIn("HX-Refresh", page.headers)
        self.assertContains(page, "CSRF", status_code=403)  # Django's own failure page


# ---------------------------------------------------------------------------
# PostgreSQL only: the policies themselves
# ---------------------------------------------------------------------------


@skipUnless(POSTGRES, "row-level security only exists on PostgreSQL")
class PolicyTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.alice = make_user("Alice", username="alice")
        cls.bob = make_user("Bob", username="bob")
        cls.alice_app = make_application(cls.alice, "Acme", "Ingénieure")
        cls.bob_app = make_application(cls.bob, "Globex", "Ingénieur")
        SearchProfile.objects.create(user=cls.alice, job_titles=["Ingénieure"])

    def test_every_registered_table_carries_the_policy(self):
        report = verify.inspect(connection)
        for table in report.tables:
            with self.subTest(table=table.table):
                self.assertTrue(table.exists)
                self.assertTrue(table.rls_enabled)
                assert table.policy is not None
                self.assertIn("current_setting('app.current_user_id'", table.policy)
        self.assertTrue(report.maintenance_side)  # the test connection is the superuser
        self.assertFalse(report.role.is_app_role)
        with app_role():
            self.assertTrue(verify.inspect(connection).role.is_app_role)

    def test_database_check_warns_for_the_owner_and_errors_for_the_application_role(self):
        # The superuser is the maintenance side: nothing is an Error, and the
        # only warnings concern an installed extension without policies.
        owner_side = checks.check_database(None, databases=["default"])
        self.assertEqual({p.id for p in owner_side} - {"rls.W002"}, set())
        self.assertEqual([p.msg for p in owner_side if not foreign(p.msg)], [])
        with app_role():
            app_side = checks.check_database(None, databases=["default"])
        self.assertEqual({p.id for p in app_side} - {"rls.E004"}, set())
        self.assertEqual([p.msg for p in app_side if not foreign(p.msg)], [])

    def test_rls_grant_is_idempotent_and_refused_to_the_application_role(self):
        out = io.StringIO()
        call_command("rls_grant", stdout=out)
        self.assertIn("Droits posés", out.getvalue())
        with app_role():
            with self.assertRaises(Exception) as caught:
                call_command("rls_grant", stdout=io.StringIO())
        self.assertIn("propriétaire", str(caught.exception))

    def test_superuser_connection_bypasses_the_policies(self):
        with context.as_user(self.alice):
            self.assertEqual(Application.objects.count(), 2)

    def test_nobody_bound_sees_nothing_but_the_accounts(self):
        with app_role(), context.as_user(None):
            self.assertEqual(Application.objects.count(), 0)
            self.assertEqual(Company.objects.count(), 0)
            self.assertEqual(Contact.objects.count(), 0)
            self.assertEqual(ActivityEvent.objects.count(), 0)
            self.assertEqual(Profile.objects.count(), 0)
            self.assertEqual(Preferences.objects.count(), 0)
            self.assertEqual(SearchProfile.objects.count(), 0)
            # Sign-in must be able to find the account.
            self.assertEqual(User.objects.count(), 2)

    def test_bound_account_sees_its_rows_only_even_without_a_filter(self):
        with app_role(), context.as_user(self.alice):
            self.assertEqual(list(Application.objects.all()), [self.alice_app])
            self.assertEqual(list(Company.objects.values_list("name", flat=True)), ["Acme"])
            self.assertEqual(Contact.objects.get().name, "Contact Acme")
            self.assertEqual(ActivityEvent.objects.get().title, "Note Acme")
            self.assertEqual(list(User.objects.all()), [self.alice])
            self.assertEqual(Profile.objects.get().user, self.alice)
            self.assertEqual(SearchProfile.objects.get().user, self.alice)
            with self.assertRaises(Application.DoesNotExist):
                Application.objects.get(pk=self.bob_app.pk)

    def test_nobody_may_delete_an_account_unbound(self):
        with app_role(), context.as_user(None):
            deleted, _ = User.objects.filter(pk=self.bob.pk).delete()
            self.assertEqual(deleted, 0)
            self.assertEqual(User.objects.filter(pk=self.bob.pk).update(first_name="x"), 1)
        self.assertTrue(User.objects.filter(pk=self.bob.pk).exists())

    def test_an_anonymous_transaction_overrides_a_session_value(self):
        # A session-level SET (startup options, a stray psql session through
        # PgBouncer) must not turn an anonymous request into that account.
        connection.connection.execute(f"SET {context.GUC} = %s", (str(self.alice.pk),)).close()
        self.addCleanup(lambda: connection.connection.execute(f"RESET {context.GUC}").close())
        with app_role():
            with context.as_user(None):
                self.assertEqual(Application.objects.count(), 0)
            seen = {}

            def view(request):
                seen["n"] = Application.objects.count()
                return HttpResponse()

            request = RequestFactory().get("/")
            request.user = AnonymousUser()
            RowLevelSecurityMiddleware(view)(request)
            self.assertEqual(seen["n"], 0)

    def test_bookkeeping_tables_are_read_only_for_the_application_role(self):
        with app_role():
            self.assertGreater(Group.objects.count() + 1, 0)  # readable
            with self.assertRaises(DatabaseError), transaction.atomic():
                Group.objects.create(name="Intrus")

    def test_readable_tables_without_a_policy_are_reported(self):
        with connection.cursor() as cursor:
            cursor.execute("CREATE TABLE rls_probe_unprotected (id integer)")
            cursor.execute(f"GRANT SELECT ON rls_probe_unprotected TO {settings.RLS_APP_ROLE}")
        self.addCleanup(lambda: connection.cursor().execute("DROP TABLE IF EXISTS rls_probe_unprotected"))
        report = verify.inspect(connection)
        self.assertIn("rls_probe_unprotected", report.unprotected)
        self.assertTrue(any("rls_probe_unprotected" in p for p in verify.problems(connection, runtime=False)))
        self.assertNotIn("django_session", report.unprotected)
        self.assertEqual(report.writable_bookkeeping, [])

        probe = SimpleNamespace(_meta=SimpleNamespace(
            label="rls.Probe", db_table="rls_probe_unprotected",
        ))
        with mock.patch("rls.verify.apps.get_models", return_value=[probe]), mock.patch.object(
            registry, "_exempt", registry._exempt | {"rls.probe"},
        ):
            self.assertNotIn("rls_probe_unprotected", verify.inspect(connection).unprotected)

    @skipUnless(apps.is_installed("django_q"), "the queue tables come with the copilot")
    def test_queue_tables_stay_allowed_once_the_copilot_is_off(self):
        # ``COPILOT_ENABLED=0`` on a database migrated with the copilot: the
        # django_q tables remain, readable by the application role, and no
        # model carries ``rls.exempt`` any more. The guard must not turn
        # that into ``ImproperlyConfigured`` on the first request.
        with mock.patch.object(registry, "_exempt", set()):
            report = verify.inspect(connection)
            with app_role():  # what the middleware sees on the first request
                problems = verify.problems(connection, runtime=True)
        for table in sql.QUEUE_TABLES:
            with self.subTest(table=table):
                self.assertNotIn(table, report.unprotected)
        self.assertEqual(problems, [])

    def test_rows_for_another_account_are_refused(self):
        with app_role(), context.as_user(self.alice):
            with self.assertRaises(DatabaseError), transaction.atomic():
                Company.objects.create(owner=self.bob, name="Initech")
            with self.assertRaises(DatabaseError), transaction.atomic():
                Contact.objects.create(application=self.bob_app, name="Intrus")
            with self.assertRaises(DatabaseError), transaction.atomic():
                Application.objects.filter(pk=self.alice_app.pk).update(owner=self.bob)
        self.assertEqual(Company.objects.filter(name="Initech").count(), 0)

    def test_updates_and_deletes_cannot_reach_another_account(self):
        with app_role(), context.as_user(self.alice):
            touched = Application.objects.filter(pk=self.bob_app.pk).update(title="Piraté")
            self.assertEqual(touched, 0)
            deleted, _ = Application.objects.filter(pk=self.bob_app.pk).delete()
            self.assertEqual(deleted, 0)
        self.bob_app.refresh_from_db()
        self.assertEqual(self.bob_app.title, "Ingénieur")

    def test_nested_as_user_restores_the_previous_account(self):
        with app_role(), context.as_user(self.alice):
            with context.as_user(self.bob):
                self.assertEqual(list(Application.objects.all()), [self.bob_app])
            self.assertEqual(list(Application.objects.all()), [self.alice_app])
            with context.as_user(None):
                self.assertEqual(Application.objects.count(), 0)
            self.assertEqual(list(Application.objects.all()), [self.alice_app])

    def test_a_rolled_back_block_restores_the_previous_account(self):
        with app_role(), context.as_user(self.alice):
            with self.assertRaises(RuntimeError):
                with context.as_user(self.bob):
                    self.assertEqual(Application.objects.count(), 1)
                    raise RuntimeError("abandon")
            self.assertEqual(list(Application.objects.all()), [self.alice_app])

    def test_rebind_switches_the_open_transaction(self):
        with app_role(), context.as_user(None):
            self.assertEqual(Application.objects.count(), 0)
            context.rebind(self.bob)
            self.assertEqual(list(Application.objects.all()), [self.bob_app])

    def test_status_command_reports_the_maintenance_side(self):
        out = io.StringIO()
        try:
            call_command("rls_status", "--probe", stdout=out)
        except SystemExit:
            # An installed extension without policies is the one accepted reason.
            pass
        text = out.getvalue()
        self.assertIn("maintenance", text)
        self.assertIn("tracker_application", text)
        self.assertIn("SET ROLE possible", text)
        self.assertNotIn("FUITE", text)
        problems = [line for line in text.splitlines() if line.startswith("✗")]
        self.assertEqual([p for p in problems if not foreign(p)], [])


@skipUnless(POSTGRES, "row-level security only exists on PostgreSQL")
@override_settings(AUTH_MODE="local")
class RequestTests(TestCase):
    """Whole requests as the application role (the runner's test client)."""

    @classmethod
    def setUpTestData(cls):
        cls.alice = make_user("Alice", username="alice")
        cls.bob = make_user("Bob", username="bob")
        cls.alice_app = make_application(cls.alice, "Acme", "Ingénieure")
        cls.bob_app = make_application(cls.bob, "Globex", "Ingénieur")

    def setUp(self):
        self.assertIsInstance(self.client, AppRoleClient)
        self.client.force_login(self.alice)

    def test_pages_show_the_signed_in_account_only(self):
        response = self.client.get(reverse("tracker:application_list"))
        self.assertContains(response, "Acme")
        self.assertNotContains(response, "Globex")

    def test_another_accounts_row_is_a_404(self):
        response = self.client.get(reverse("tracker:application_detail", args=[self.bob_app.pk]))
        self.assertEqual(response.status_code, 404)

    def test_writes_land_on_the_signed_in_account(self):
        response = self.client.post(
            reverse("tracker:quick_create"),
            {"company_name": "Initech", "title": "SRE", "status": "backlog", "cv_language": "fr"},
            HTTP_HX_REQUEST="true",
        )
        self.assertLess(response.status_code, 400)
        created = Application.objects.get(title="SRE")
        self.assertEqual(created.owner, self.alice)
        self.assertEqual(created.company.owner, self.alice)

    def test_five_hundred_rolls_the_request_back(self):
        def failing_view(request):
            Company.objects.create(owner=self.alice, name="Éphémère")
            return HttpResponse(status=500)

        request = RequestFactory().get("/")
        request.user = self.alice
        with app_role():
            RowLevelSecurityMiddleware(failing_view)(request)
        self.assertFalse(Company.objects.filter(name="Éphémère").exists())

    def test_deleting_the_account_removes_everything_it_owns(self):
        response = self.client.post(reverse("accounts:settings_section", args=["supprimer"]), {"confirm": "Alice"})
        self.assertEqual(response.status_code, 302)
        self.assertFalse(User.objects.filter(pk=self.alice.pk).exists())
        self.assertFalse(Application.objects.filter(pk=self.alice_app.pk).exists())
        self.assertTrue(Application.objects.filter(pk=self.bob_app.pk).exists())

    def test_local_chooser_reads_every_profile(self):
        self.client.logout()
        response = self.client.get(reverse("accounts:login"))
        self.assertContains(response, "Alice")
        self.assertContains(response, "Bob")
        response = self.client.post(reverse("accounts:login"), {"user": self.bob.pk})
        self.assertRedirects(response, reverse("tracker:dashboard"))
        page = self.client.get(reverse("tracker:application_list"))
        self.assertContains(page, "Globex")
        self.assertNotContains(page, "Acme")

    @override_settings(AUTH_MODE="accounts", SIGNUP_OPEN=True)
    def test_signup_creates_the_profile_rows(self):
        self.client.logout()
        response = self.client.post(
            reverse("accounts:signup"),
            {
                "display_name": "Carol",
                "email": "carol@example.com",
                "password1": "un-mot-de-passe-solide",
                "password2": "un-mot-de-passe-solide",
            },
        )
        self.assertEqual(response.status_code, 302, response.content[:300])
        carol = User.objects.get(email="carol@example.com")
        self.assertTrue(Profile.objects.filter(user=carol).exists())
        self.assertTrue(Preferences.objects.filter(user=carol).exists())

    def test_runtime_guard_refuses_a_role_that_bypasses_the_policies(self):
        RowLevelSecurityMiddleware.reset_verification()
        self.addCleanup(RowLevelSecurityMiddleware.reset_verification)
        with override_settings(RLS_ENFORCE=True):
            superuser_client = Client()
            superuser_client.force_login(self.alice)
            with self.assertRaises(ImproperlyConfigured):
                superuser_client.get(reverse("tracker:application_list"))
        # The application role passes — except for the tables of an installed
        # extension that has not registered them, which production refuses.
        with app_role():
            problems = [p for p in verify.problems(connection) if not foreign(p)]
        self.assertEqual(problems, [])

    def test_a_rolled_back_sign_in_keeps_a_session(self):
        request = RequestFactory().get("/")
        request.user = self.alice
        request.session = SessionStore()
        request.session.create()
        key_before = request.session.session_key

        def signs_in_then_fails(request):
            request.session.cycle_key()
            return HttpResponse(status=500)

        with app_role():
            RowLevelSecurityMiddleware(signs_in_then_fails)(request)
        key_after = request.session.session_key
        self.assertNotEqual(key_after, key_before)
        self.assertTrue(Session.objects.filter(session_key=key_after).exists())


@skipUnless(POSTGRES, "row-level security only exists on PostgreSQL")
class CommandTests(AppRoleTestCase, TestCase):
    """Management commands run with the runtime credentials."""

    @classmethod
    def setUpTestData(cls):
        cls.alice = make_user("Alice", username="alice")

    def test_legacy_import_writes_on_the_owners_behalf(self):
        # CI has no private legacy files; exercise the real importer with a
        # minimal workbook while still enforcing the runtime role's policies.
        with TemporaryDirectory(prefix="jobhunt-rls-import-") as directory:
            workbook_path = Path(directory) / "applications.xlsx"
            workbook = Workbook()
            sheet = workbook.create_sheet("Candidatures")
            sheet.append(["Société", "Intitulé du poste"])
            sheet.append(["Example", "Engineer"])
            workbook.save(workbook_path)
            workbook.close()
            call_command(
                "import_legacy",
                workbook=str(workbook_path),
                root=directory,
                skip_files=True,
                user="alice",
                stdout=io.StringIO(),
            )
        self.assertEqual(Application.objects.count(), 0)  # unbound: nothing visible
        with context.as_user(self.alice):
            self.assertEqual(Application.objects.get().title, "Engineer")


@skipUnless(POSTGRES, "row-level security only exists on PostgreSQL")
class BackendHookTests(TransactionTestCase):
    """Outside any enclosing transaction the engine's hook is the only announcer."""

    def test_an_outermost_transaction_announces_the_account(self):
        alice = make_user("Alice", username="alice")
        bob = make_user("Bob", username="bob")
        make_application(alice, "Acme")
        make_application(bob, "Globex")
        self.assertFalse(connection.in_atomic_block)
        with app_role():
            with context.as_user(alice):
                self.assertFalse(connection.get_autocommit())
                self.assertEqual(list(Application.objects.values_list("company__name", flat=True)), ["Acme"])
            with context.as_user(None):
                self.assertEqual(Application.objects.count(), 0)
            # Bound but outside any transaction: nothing was announced.
            with context.bound(alice):
                self.assertEqual(Application.objects.count(), 0)
