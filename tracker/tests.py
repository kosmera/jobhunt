"""Tests for the tracker.

Coverage is deliberately weighted towards rendering: most of the risk in an
HTMX app is a template that only breaks once a branch finally has data in it.
"""

from __future__ import annotations

import ast
import base64
import datetime as dt
import io
import logging
import re
import shutil
import tempfile
import time
from pathlib import Path
from types import SimpleNamespace
from typing import Any
from unittest import mock, skipUnless
from urllib.parse import parse_qs, quote, urlsplit

from django.conf import settings
from django.contrib.auth import get_user_model
from django.contrib.auth.models import User
from django.core import signing
from django.core.exceptions import ImproperlyConfigured, SuspiciousFileOperation
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.db import connection, transaction
from django.db.migrations.executor import MigrationExecutor
from django.test import (
    RequestFactory,
    SimpleTestCase,
    TestCase,
    TransactionTestCase,
    override_settings,
)
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone

import tracker
from accounts.services import LOCAL_USERNAME, preferences_for, profile_for
from accounts.testing import OwnedTestCase, make_user
from jobhunt import plugins as plugin_registry
from jobhunt.storage import parse_connection_string
from tracker import checks, domain, links, privacy, queries, services
from tracker.adapters import cv_analyzer, document_text, persistence, storage
from tracker.adapters.azure_storage import AzureStorageAdapter
from tracker.adapters.django_orm import DjangoPersistence
from tracker.adapters.memory import MemoryPersistence
from tracker.adapters.file_storage import LocalStorageAdapter, MemoryStorageAdapter
from tracker.management.commands import storage_status
from tracker.models import (
    DOCUMENT_NAME_MAX_LENGTH,
    PIPELINE_STATUSES,
    ActivityEvent,
    Application,
    Company,
    Contact,
    Document,
    DocumentKind,
    EventKind,
    GapStatus,
    Language,
    Platform,
    Sector,
    SkillGap,
    Status,
    WorkMode,
)
from tracker.ports import MissingFile, NotFound, StorageError, StoragePort

MEDIA = tempfile.mkdtemp(prefix="jobhunt-tests-")

#: A base64 account key, as ``generate_blob_sas`` decodes it before signing.
AZURE_KEY = base64.b64encode(b"k" * 32).decode()


def make_docx(*paragraphs: str, table: tuple[str, ...] = ()) -> bytes:
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=1, cols=len(table))
        for cell, value in zip(grid.rows[0].cells, table):
            cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pdf() -> bytes:
    """A one-page PDF without a text layer (a scan, as far as extraction goes)."""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


class FakeAnalyzer:
    """A ``CVAnalyzer`` that records what the core hands it."""

    def __init__(self):
        self.calls: list[tuple[Any, dict]] = []

    def analyze_cv(self, owner, **kwargs):
        self.calls.append((owner, kwargs))


# --- A stand-in for the Azure SDK's clients: dicts, the SDK's own exceptions ---


class FakeDownloader:
    def __init__(self, data: bytes):
        self._buffer = io.BytesIO(data)
        self.size = len(data)

    def read(self, size: int = -1) -> bytes:
        return self._buffer.read(size)

    def readall(self) -> bytes:
        return self._buffer.read()


class FakeBlob:
    def __init__(self, container, name: str):
        self.container, self.name = container, name

    @property
    def url(self) -> str:
        return f"https://jobhunt.blob.core.windows.net/{self.container.name}/{quote(self.name, safe='~/')}"

    def _data(self) -> bytes:
        from azure.core.exceptions import ResourceNotFoundError

        try:
            return self.container.blobs[self.name]
        except KeyError:
            raise ResourceNotFoundError(f"Pas de blob {self.name}")

    def exists(self) -> bool:
        return self.name in self.container.blobs

    def download_blob(self) -> FakeDownloader:
        return FakeDownloader(self._data())

    def delete_blob(self) -> None:
        self._data()
        del self.container.blobs[self.name]

    def get_blob_properties(self):
        return SimpleNamespace(
            size=len(self._data()), last_modified=dt.datetime.now(dt.timezone.utc)
        )


class BrokenBlob:
    """Every call fails the way an unreachable account does."""

    url = "https://jobhunt.blob.core.windows.net/documents/x"

    def _down(self):
        from azure.core.exceptions import ServiceRequestError

        raise ServiceRequestError("connexion refusée")

    exists = download_blob = delete_blob = get_blob_properties = _down


class FakeContainer:
    def __init__(self, name: str = "documents"):
        self.name = name
        self.blobs: dict[str, bytes] = {}
        self.uploads: list[tuple] = []
        self.fail_next_upload = False
        self.race_on_next_upload = False
        self.broken = False
        self.missing = False

    def get_blob_client(self, name: str):
        return BrokenBlob() if self.broken else FakeBlob(self, name)

    def list_blobs(self, name_starts_with=None):
        for name in sorted(self.blobs):
            if not name_starts_with or name.startswith(name_starts_with):
                yield SimpleNamespace(name=name)

    def get_container_properties(self):
        from azure.core.exceptions import ResourceNotFoundError, ServiceRequestError

        if self.broken:
            raise ServiceRequestError("connexion refusée")
        if self.missing:
            # The SDK sets ``error_code`` from the response; the stub types it
            # read-only, hence the targeted ignore rather than a fake response.
            error = ResourceNotFoundError("pas de conteneur")
            error.error_code = "ContainerNotFound"  # pyright: ignore[reportAttributeAccessIssue]
            raise error
        return SimpleNamespace(name=self.name)

    def upload_blob(self, name, data, length=None, overwrite=False, content_settings=None, **kwargs):
        from azure.core.exceptions import ResourceExistsError

        if self.fail_next_upload:
            self.fail_next_upload = False
            raise ResourceExistsError("existe déjà")
        if self.race_on_next_upload:
            # Another writer got there first: the blob is really there now.
            self.race_on_next_upload = False
            self.blobs[name] = b"quelqu'un d'autre"
            raise ResourceExistsError("existe déjà")
        if name in self.blobs and not overwrite:
            raise ResourceExistsError("existe déjà")
        payload = data.read() if hasattr(data, "read") else bytes(data)
        content_type = getattr(content_settings, "content_type", None)
        self.uploads.append((name, length, len(payload), content_type))
        self.blobs[name] = payload


class FakeTokenCredential:
    def get_token(self, *scopes, **kwargs):  # pragma: no cover — shape only
        raise AssertionError("jamais appelé : la clé de délégation vient du service")


class FakeService:
    def __init__(self, *, account_key: str | None = None, token: bool = False):
        self.account_name = "jobhunt"
        if account_key:
            self.credential: Any = SimpleNamespace(account_name="jobhunt", account_key=account_key)
        elif token:
            self.credential = FakeTokenCredential()
        else:
            self.credential = "sv=2024-05-04&sig=abc"  # a SAS-only client
        self.container = FakeContainer()
        self.delegation_calls: list[tuple[dt.datetime, dt.datetime]] = []

    def get_container_client(self, name: str):
        assert name == self.container.name, name
        return self.container

    def get_user_delegation_key(self, start, expiry):
        from azure.storage.blob import UserDelegationKey

        self.delegation_calls.append((start, expiry))
        key = UserDelegationKey()
        key.signed_oid = "00000000-0000-0000-0000-000000000001"
        key.signed_tid = "00000000-0000-0000-0000-000000000002"
        key.signed_start = start.strftime("%Y-%m-%dT%H:%M:%SZ")
        key.signed_expiry = expiry.strftime("%Y-%m-%dT%H:%M:%SZ")
        key.signed_service = "b"
        key.signed_version = "2020-02-10"
        key.value = AZURE_KEY
        return key


def azure_adapter(**options) -> tuple[AzureStorageAdapter, FakeService]:
    service = FakeService(account_key=AZURE_KEY)
    options.setdefault("container", "documents")
    return AzureStorageAdapter(client=service, **options), service


def azure_settings(service: FakeService, **options) -> dict:
    """``STORAGES`` pointing the default storage at a fake Azure account."""
    return {
        "default": {
            "BACKEND": "tracker.adapters.azure_storage.AzureStorageAdapter",
            "OPTIONS": {"client": service, "container": service.container.name, **options},
        },
        "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
    }


#: Enough text for ``services.MIN_TEXT_LENGTH``: a CV the extractor can read.
CV_BODY = (
    "Ingénieur DevOps senior, douze ans d'expérience sur des plateformes Linux. "
    "Automatisation avec Ansible et Terraform, conteneurs Docker et Kubernetes, "
    "intégration continue GitLab CI, supervision Prometheus et Grafana. "
    "Bases PostgreSQL et Redis, réseaux et pare-feux, scripts Python et Bash. "
    "Habitué au travail en équipe réduite et à la reprise de systèmes existants."
)


MEMORY_STORAGES = {
    "default": {"BACKEND": "tracker.adapters.file_storage.MemoryStorageAdapter", "OPTIONS": {"link_ttl": 60}},
    "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
}


def default_owner():
    """The test's signed-in profile — or a fresh one for model-level tests."""
    return get_user_model().objects.order_by("pk").first() or make_user()


def make_company(**kwargs) -> Company:
    kwargs.setdefault("owner", default_owner())
    kwargs.setdefault("sector", Sector.PRIVATE)
    return Company.objects.create(**kwargs)


def make_platform(**kwargs) -> Platform:
    kwargs.setdefault("owner", default_owner())
    return Platform.objects.create(**kwargs)


def make_gap(**kwargs) -> SkillGap:
    kwargs.setdefault("owner", default_owner())
    return SkillGap.objects.create(**kwargs)


def make_document(**kwargs) -> Document:
    application = kwargs.get("application")
    kwargs.setdefault("owner", application.owner if application else default_owner())
    return Document.objects.create(**kwargs)


def make_application(**kwargs) -> Application:
    company = kwargs.pop("company", None)
    owner = kwargs.pop("owner", None) or (company.owner if company else default_owner())
    company = company or make_company(owner=owner, name=kwargs.pop("company_name", "Acme"))
    defaults = {
        "title": "Senior DevOps Engineer",
        "location": "Bruxelles",
        "distance_km": 30,
        "score": 72,
        "cv_language": Language.FR,
        "status": Status.TO_APPLY,
        "summary": "Un résumé.",
        "strengths": "- Linux\n- Ansible",
        "weaknesses": "- Pas de Terraform",
        "strategy": "Joue la carte réseau.",
    }
    defaults.update(kwargs)
    return Application.objects.create(owner=owner, company=company, **defaults)


@override_settings(MEDIA_ROOT=MEDIA)
class PageRenderTests(OwnedTestCase):
    """Every page must render with data, and again when the data is empty."""

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        platform = make_platform(
            name="LinkedIn", searched_for="DevOps", outcome="Source la plus productive."
        )
        make_platform(name="Pistes non explorées", outcome="uvcw.be", is_lead=True)
        make_gap(
            name="Terraform", demand_count=6, demand_label="6 offres sur 10",
            why_it_matters="Standard de fait.", action_plan="Reprovisionne le homelab.",
        )
        cls.ready = make_application(company_name="HMS Networks", score=88, distance_km=0)
        cls.ready.source_platform = platform
        cls.ready.save()

        cls.sent = make_application(
            company_name="Haulogy",
            status=Status.SENT,
            applied_on=timezone.localdate() - dt.timedelta(days=30),
            follow_up_on=timezone.localdate() - dt.timedelta(days=3),
        )
        cls.interview = make_application(company_name="Collibra", status=Status.INTERVIEW)
        cls.closed = make_application(
            company_name="Odoo", status=Status.REJECTED, closed_on=timezone.localdate()
        )
        cls.discarded = make_application(
            company_name="Smals",
            status=Status.DISCARDED,
            discard_reason="Poste Windows, et néerlandais exigé.",
        )
        cls.backlog = make_application(company_name="iMio", status=Status.BACKLOG, score=None)

        Contact.objects.create(
            application=cls.sent, name="Marie Dupont", role="Recruteuse",
            email="m@example.com", phone="+32 2 000 00 00",
        )
        ActivityEvent.objects.create(
            application=cls.sent, kind=EventKind.APPLIED, title="Candidature envoyée"
        )
        ActivityEvent.objects.create(
            application=cls.interview,
            kind=EventKind.INTERVIEW,
            title="Entretien technique",
            happened_on=timezone.localdate() + dt.timedelta(days=4),
        )
        make_document(
            application=cls.ready, kind=DocumentKind.CV, label="CV.docx",
            language=Language.EN, is_primary=True,
            file=SimpleUploadedFile("CV.docx", b"x" * 40),
        )
        make_document(
            kind=DocumentKind.CV, label="CV_base.docx", language=Language.FR,
            file=SimpleUploadedFile("CV_base.docx", b"y" * 40),
        )

    def test_every_page_renders(self):
        for name, args in [
            ("tracker:dashboard", []),
            ("tracker:pipeline", []),
            ("tracker:application_list", []),
            ("tracker:document_library", []),
            ("tracker:insights", []),
            ("tracker:application_create", []),
            ("tracker:application_detail", [self.ready.pk]),
            ("tracker:application_update", [self.ready.pk]),
            ("accounts:settings", []),
        ]:
            with self.subTest(page=name):
                response = self.client.get(reverse(name, args=args))
                self.assertEqual(response.status_code, 200)

    def test_detail_renders_for_every_status(self):
        for status, _ in Status.choices:
            with self.subTest(status=status):
                self.ready.status = status
                self.ready.save(update_fields=["status"])
                response = self.client.get(self.ready.get_absolute_url())
                self.assertEqual(response.status_code, 200)

    def test_every_sprite_icon_is_a_symbol_with_a_viewbox(self):
        """<g> cannot carry a viewBox, so a 24x24 icon would be clipped to 16px."""
        import re

        html = self.client.get(reverse("tracker:dashboard")).content.decode()
        # Le cœur et chaque extension apportent leur propre sprite.
        sprites = re.findall(r"(<svg width=\"0\".*?</svg>)", html, re.S)
        self.assertTrue(sprites, "aucun sprite trouvé dans la page")

        definitions = re.findall(r'<(\w+) id="(i-[a-z-]+)"([^>]*)>', "".join(sprites))
        self.assertTrue(definitions, "aucune icône trouvée dans le sprite")
        for tag, name, attrs in definitions:
            with self.subTest(icon=name):
                self.assertEqual(tag, "symbol", f"{name} est un <{tag}>, pas un <symbol>")
                self.assertIn("viewBox", attrs, f"{name} n'a pas de viewBox")

        # Every icon referenced anywhere must exist in the sprite.
        referenced = set(re.findall(r'<use href="#(i-[a-z-]+)"', html))
        defined = {name for _, name, _ in definitions}
        self.assertFalse(referenced - defined, f"icônes manquantes : {referenced - defined}")

    def test_no_svg_relies_on_the_hidden_attribute(self):
        """`hidden` is an HTMLElement property — on an <svg> it does nothing."""
        import re

        for name in ["tracker:dashboard", "tracker:application_list"]:
            html = self.client.get(reverse(name)).content.decode()
            for tag in re.findall(r"<svg[^>]*>", html):
                with self.subTest(page=name, tag=tag[:60]):
                    self.assertNotRegex(
                        tag,
                        r"(?<!aria-)\bhidden\b(?!-)",
                        "un <svg> masqué par l'attribut hidden reste visible",
                    )

    def test_the_favicon_data_uri_carries_usable_colours(self):
        """`urlencode` escapes "#" itself — pre-escaping it yields %2523."""
        import re
        import urllib.parse

        html = self.client.get(reverse("tracker:dashboard")).content.decode()
        match = re.search(r'<link rel="icon" href="([^"]+)"', html)
        assert match is not None
        href = match.group(1)
        svg = urllib.parse.unquote(href)
        colours = re.findall(r'(?:fill|stroke)="([^"]+)"', svg)
        self.assertTrue(colours)
        for colour in colours:
            self.assertTrue(
                colour == "none" or colour.startswith("#"),
                f"couleur illisible par le moteur SVG : {colour}",
            )

    def test_dashboard_surfaces_what_needs_attention(self):
        response = self.client.get(reverse("tracker:dashboard"))
        self.assertContains(response, "À traiter maintenant")
        self.assertIn(self.sent, response.context["attention"])
        self.assertIn(self.ready, response.context["to_apply"])
        # Import chatter about written-off offers must not crowd the feed.
        self.assertNotIn(self.discarded, [e.application for e in response.context["recent"]])

    def test_pages_render_with_an_empty_database(self):
        Document.objects.all().delete()
        ActivityEvent.objects.all().delete()
        Contact.objects.all().delete()
        Application.objects.all().delete()
        for name in [
            "tracker:dashboard", "tracker:pipeline", "tracker:application_list",
            "tracker:document_library", "tracker:insights",
        ]:
            with self.subTest(page=name):
                self.assertEqual(self.client.get(reverse(name)).status_code, 200)

    def test_css_numbers_never_use_a_decimal_comma(self):
        """A French locale would render 75.0 as "75,0" and break the CSS."""
        for name in ["tracker:insights", "tracker:dashboard"]:
            with self.subTest(page=name):
                html = self.client.get(reverse(name)).content.decode()
                for fragment in ["left: ", "width: ", "bottom: "]:
                    for chunk in html.split(fragment)[1:]:
                        value = chunk.split(";")[0].split('"')[0].strip()
                        self.assertNotIn(",", value, f"{fragment}{value}")

    def test_insights_plots_only_fully_measured_offers(self):
        response = self.client.get(reverse("tracker:insights"))
        plotted = [point["application"] for point in response.context["points"]]
        self.assertIn(self.ready, plotted)
        self.assertNotIn(self.backlog, plotted)  # no score
        self.assertNotIn(self.discarded, plotted)


@override_settings(MEDIA_ROOT=MEDIA)
class FilterTests(OwnedTestCase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.open = make_application(company_name="HMS Networks", score=88)
        cls.public = make_application(
            company=make_company(name="Egov Select", sector=Sector.PUBLIC),
            status=Status.SENT, cv_language=Language.FR, score=68,
        )
        cls.english = make_application(company_name="Collibra", cv_language=Language.EN, score=72)
        cls.discarded = make_application(
            company_name="Smals", status=Status.DISCARDED,
            discard_reason="Stack Windows et néerlandais exigé.",
        )

    def get_list(self, **params):
        return self.client.get(reverse("tracker:application_list"), params)

    def test_default_scope_hides_discarded(self):
        results = self.get_list().context["applications"]
        self.assertNotIn(self.discarded, results)
        self.assertIn(self.open, results)

    def test_discarded_scope(self):
        results = self.get_list(scope="discarded").context["applications"]
        self.assertEqual(list(results), [self.discarded])

    def test_search_reaches_into_the_analysis(self):
        self.open.weaknesses = "- Aucune expérience de production manufacturière"
        self.open.save()
        results = self.get_list(q="manufacturière").context["applications"]
        self.assertEqual(list(results), [self.open])

    def test_search_reaches_discard_reasons(self):
        results = self.get_list(q="néerlandais", scope="all").context["applications"]
        self.assertEqual(list(results), [self.discarded])

    def test_sector_and_language_filters(self):
        self.assertEqual(
            list(self.get_list(sector=Sector.PUBLIC).context["applications"]), [self.public]
        )
        self.assertEqual(
            list(self.get_list(language=Language.EN).context["applications"]), [self.english]
        )

    def test_the_search_box_carries_the_hook_its_trigger_looks_for(self):
        html = self.get_list().content.decode()
        self.assertIn("data-search-input", html)
        self.assertIn("keyup from:[data-search-input]", html)

    def test_htmx_request_returns_only_the_table(self):
        response = self.get_list(**{"HTTP_HX_REQUEST": "true"}) if False else self.client.get(
            reverse("tracker:application_list"), HTTP_HX_REQUEST="true"
        )
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "<html")
        self.assertContains(response, 'id="applications"')

    def test_sort_options_all_execute(self):
        for sort in ["-score", "pipeline", "-applied_on", "follow_up_on", "company__name", "-updated_at"]:
            with self.subTest(sort=sort):
                response = self.get_list(sort=sort)
                self.assertEqual(len(list(response.context["applications"])), 3)


@override_settings(MEDIA_ROOT=MEDIA)
class StatusTransitionTests(TestCase):
    def setUp(self):
        self.application = make_application()
        self.preferences = preferences_for(self.application.owner)

    def test_moving_to_sent_stamps_the_date_and_schedules_a_follow_up(self):
        self.application.apply_status(Status.SENT)
        self.application.refresh_from_db()
        self.assertEqual(self.application.applied_on, timezone.localdate())
        self.assertEqual(
            self.application.follow_up_on,
            timezone.localdate() + dt.timedelta(days=self.preferences.follow_up_days),
        )
        event = self.application.events.first()
        assert event is not None
        self.assertEqual(event.kind, EventKind.APPLIED)

    def test_the_follow_up_delay_is_the_owner_s(self):
        self.preferences.follow_up_days = 3
        self.preferences.save()
        application = Application.objects.with_related().get(pk=self.application.pk)
        application.apply_status(Status.SENT)
        self.assertEqual(application.follow_up_on, timezone.localdate() + dt.timedelta(days=3))

    def test_closing_clears_the_follow_up_and_stamps_the_close_date(self):
        self.application.apply_status(Status.SENT)
        self.application.apply_status(Status.REJECTED)
        self.application.refresh_from_db()
        self.assertIsNone(self.application.follow_up_on)
        self.assertEqual(self.application.closed_on, timezone.localdate())

    def test_reopening_a_closed_application_clears_the_close_date(self):
        self.application.apply_status(Status.REJECTED)
        self.application.apply_status(Status.SENT)
        self.application.refresh_from_db()
        self.assertIsNone(self.application.closed_on)

    def test_a_no_op_transition_logs_nothing(self):
        self.assertFalse(self.application.apply_status(Status.TO_APPLY))
        self.assertEqual(self.application.events.count(), 0)

    def test_an_existing_applied_date_is_preserved(self):
        earlier = timezone.localdate() - dt.timedelta(days=12)
        self.application.applied_on = earlier
        self.application.apply_status(Status.SENT)
        self.application.refresh_from_db()
        self.assertEqual(self.application.applied_on, earlier)

    def test_staleness_and_follow_up_states(self):
        today = timezone.localdate()
        self.application.status = Status.SENT
        self.application.applied_on = today - dt.timedelta(days=self.preferences.stale_after_days + 1)
        self.application.follow_up_on = today - dt.timedelta(days=2)
        self.application.save()
        self.assertTrue(self.application.is_stale)
        self.assertEqual(self.application.follow_up_state, "overdue")
        self.assertTrue(self.application.needs_attention)

        self.application.follow_up_on = today
        self.assertEqual(self.application.follow_up_state, "today")
        self.application.follow_up_on = today + dt.timedelta(days=2)
        self.assertEqual(self.application.follow_up_state, "soon")
        self.application.follow_up_on = today + dt.timedelta(days=30)
        self.assertEqual(self.application.follow_up_state, "later")

    def test_score_bands(self):
        for score, band in [(88, "high"), (80, "high"), (72, "mid"), (65, "mid"),
                            (55, "low"), (None, "none")]:
            self.application.score = score
            self.assertEqual(self.application.score_band, band)

    def test_staleness_threshold_is_per_profile(self):
        today = timezone.localdate()
        self.application.status = Status.SENT
        self.application.applied_on = today - dt.timedelta(days=6)
        self.application.save()
        self.assertFalse(self.application.is_stale)  # default 14 days
        self.preferences.stale_after_days = 5
        self.preferences.save()
        reloaded = Application.objects.with_related().get(pk=self.application.pk)
        self.assertTrue(reloaded.is_stale)
        self.assertIn(reloaded, Application.objects.for_user(reloaded.owner).stale(days=5))
        self.assertNotIn(reloaded, Application.objects.for_user(reloaded.owner).stale(days=14))
        self.assertIn(reloaded, Application.objects.needs_attention(days=5))

    def test_application_and_company_must_share_an_owner(self):
        other = make_user("Marie")
        with self.assertRaises(ValueError):
            Application.objects.create(
                owner=other, company=self.application.company, title="Intrus"
            )

    def test_attached_document_inherits_the_application_owner(self):
        document = Document.objects.create(
            application=self.application, kind=DocumentKind.CV, label="CV",
            file=SimpleUploadedFile("cv.docx", b"x"),
        )
        self.assertEqual(document.owner, self.application.owner)
        with self.assertRaises(ValueError):
            Document.objects.create(
                owner=make_user("Marie"), application=self.application, kind=DocumentKind.CV,
                label="CV", file=SimpleUploadedFile("cv2.docx", b"x"),
            )


@override_settings(MEDIA_ROOT=MEDIA)
class HtmxEndpointTests(OwnedTestCase):
    def setUp(self):
        super().setUp()
        self.application = make_application()
        self.gap = make_gap(name="Terraform", demand_count=6)

    def post(self, route, args=(), **data):
        return self.client.post(reverse(route, args=args), data, HTTP_HX_REQUEST="true")

    def test_set_status_from_each_surface(self):
        for source in ["board", "row", "list", "detail"]:
            with self.subTest(source=source):
                response = self.post(
                    "tracker:set_status", [self.application.pk],
                    status=Status.SENT, source=source,
                )
                self.assertEqual(response.status_code, 200)
                self.assertIn("HX-Trigger", response.headers)
                if source == "list":
                    self.assertContains(response, f'id="app-row-{self.application.pk}"')
                self.application.status = Status.TO_APPLY
                self.application.save(update_fields=["status"])

    def test_set_status_rejects_an_unknown_value(self):
        response = self.post("tracker:set_status", [self.application.pk], status="nope")
        self.assertEqual(response.status_code, 400)

    def test_advance_walks_the_pipeline_and_stops_at_the_end(self):
        self.post("tracker:advance_status", [self.application.pk], source="board")
        self.application.refresh_from_db()
        self.assertEqual(self.application.status, Status.SENT)

        self.application.status = Status.ACCEPTED
        self.application.save(update_fields=["status"])
        response = self.post("tracker:advance_status", [self.application.pk], source="detail")
        self.assertEqual(response.status_code, 200)
        self.application.refresh_from_db()
        self.assertEqual(self.application.status, Status.ACCEPTED)

    def test_follow_up_presets_and_explicit_dates(self):
        self.application.status = Status.SENT
        self.application.save(update_fields=["status"])

        self.post("tracker:set_follow_up", [self.application.pk], in_days="7")
        self.application.refresh_from_db()
        self.assertEqual(self.application.follow_up_on, timezone.localdate() + dt.timedelta(days=7))

        self.post("tracker:set_follow_up", [self.application.pk], follow_up_on="2026-12-01")
        self.application.refresh_from_db()
        self.assertEqual(self.application.follow_up_on, dt.date(2026, 12, 1))

        self.post("tracker:set_follow_up", [self.application.pk], follow_up_on="")
        self.application.refresh_from_db()
        self.assertIsNone(self.application.follow_up_on)

        self.assertEqual(
            self.post("tracker:set_follow_up", [self.application.pk],
                      follow_up_on="pas une date").status_code,
            400,
        )

    def test_mark_followed_up_logs_and_reschedules(self):
        self.application.status = Status.SENT
        self.application.save(update_fields=["status"])
        preferences = preferences_for(self.user)
        preferences.follow_up_days = 4
        preferences.save()
        self.post("tracker:mark_followed_up", [self.application.pk])
        self.application.refresh_from_db()
        event = self.application.events.first()
        assert event is not None
        self.assertEqual(event.kind, EventKind.FOLLOW_UP)
        self.assertEqual(
            self.application.follow_up_on, timezone.localdate() + dt.timedelta(days=4)
        )

    def test_notes_round_trip(self):
        response = self.post("tracker:edit_notes", [self.application.pk],
                             personal_notes="Demander la part réelle de MDM.")
        self.assertEqual(response.status_code, 200)
        self.application.refresh_from_db()
        self.assertEqual(self.application.personal_notes, "Demander la part réelle de MDM.")

    def test_event_add_and_delete(self):
        response = self.post(
            "tracker:add_event", [self.application.pk],
            happened_on="2026-09-02", kind=EventKind.CALL,
            title="Appel du recruteur", detail="20 minutes.",
        )
        self.assertContains(response, "Appel du recruteur")
        event = self.application.events.get()
        response = self.post("tracker:delete_event", [event.pk])
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.application.events.count(), 0)

    def test_invalid_event_redisplays_the_panel_with_errors(self):
        response = self.post("tracker:add_event", [self.application.pk], title="", kind="note")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.application.events.count(), 0)

    def test_document_upload_and_delete(self):
        upload = SimpleUploadedFile("CV_Haulogy.docx", b"contenu")
        response = self.client.post(
            reverse("tracker:add_document", args=[self.application.pk]),
            {"file": upload, "kind": DocumentKind.CV, "label": "", "language": Language.FR,
             "is_primary": "on"},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 200)
        document = self.application.documents.get()
        self.assertEqual(document.label, "CV_Haulogy.docx")
        self.assertTrue(document.is_primary)
        self.assertEqual(self.application.primary_cv, document)
        name = document.file.name or ""
        self.assertEqual(name, f"documents/{self.user.pk}/{self.application.slug}/CV_Haulogy.docx")
        self.assertTrue(storage().file_exists(name))
        self.assertEqual(document.size_bytes, len(b"contenu"))

        with self.captureOnCommitCallbacks(execute=True):
            response = self.post("tracker:delete_document", [document.pk])
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.application.documents.count(), 0)
        # The bytes go once the rows agree, never before.
        self.assertFalse(storage().file_exists(name))

    def test_a_cv_upload_reaches_the_analyzer_anonymised_and_other_kinds_do_not(self):
        analyzer = FakeAnalyzer()
        with mock.patch("tracker.views.cv_analyzer", return_value=analyzer):
            self.client.post(
                reverse("tracker:add_document", args=[self.application.pk]),
                {"file": SimpleUploadedFile("cv.txt", f"Lionel, lionel@x.be, 0470 12 34 56\n{CV_BODY}".encode()),
                 "kind": DocumentKind.CV, "label": "Mon CV", "language": Language.FR},
                HTTP_HX_REQUEST="true",
            )
            self.client.post(
                reverse("tracker:add_document", args=[self.application.pk]),
                {"file": SimpleUploadedFile("annonce.txt", f"lionel@x.be\n{CV_BODY}".encode()),
                 "kind": DocumentKind.POSTING},
                HTTP_HX_REQUEST="true",
            )
        (owner, call), = analyzer.calls
        cv = self.application.documents.get(kind=DocumentKind.CV)
        self.assertEqual(owner, self.user)
        self.assertEqual(set(call), {"document_id", "label", "language", "text"})
        self.assertEqual((call["document_id"], call["label"], call["language"]), (cv.pk, "Mon CV", "fr"))
        self.assertTrue(call["text"].text.startswith("[NOM], [EMAIL], [TELEPHONE]"))
        # Only the text travels: no instance, no handle, no path.
        self.assertNotIn("document", call)

    def test_uploading_a_second_primary_demotes_the_first(self):
        for index in (1, 2):
            self.client.post(
                reverse("tracker:add_document", args=[self.application.pk]),
                {"file": SimpleUploadedFile(f"CV{index}.docx", b"x"),
                 "kind": DocumentKind.CV, "label": "", "is_primary": "on"},
                HTTP_HX_REQUEST="true",
            )
        primaries = self.application.documents.filter(is_primary=True)
        self.assertEqual(primaries.count(), 1)
        self.assertEqual(primaries.get().label, "CV2.docx")

    def test_library_upload_redirects(self):
        response = self.client.post(
            reverse("tracker:add_library_document"),
            {"file": SimpleUploadedFile("CV_base.docx", b"x"), "kind": DocumentKind.CV,
             "label": "CV générique"},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 204)
        self.assertEqual(response.headers["HX-Redirect"], reverse("tracker:document_library"))
        library = Document.objects.get(application__isnull=True)
        self.assertEqual(library.owner, self.user)

    def test_contact_add_and_delete(self):
        response = self.post("tracker:add_contact", [self.application.pk],
                             name="Marie Dupont", role="Recruteuse")
        self.assertContains(response, "Marie Dupont")
        contact = Contact.objects.get()
        self.post("tracker:delete_contact", [contact.pk])
        self.assertEqual(Contact.objects.count(), 0)

    def test_gap_status_toggle(self):
        response = self.post("tracker:set_gap_status", [self.gap.pk], status=GapStatus.DOING)
        self.assertEqual(response.status_code, 200)
        self.gap.refresh_from_db()
        self.assertEqual(self.gap.status, GapStatus.DOING)
        self.assertEqual(
            self.post("tracker:set_gap_status", [self.gap.pk], status="bogus").status_code, 400
        )

    def test_quick_create(self):
        response = self.client.post(
            reverse("tracker:quick_create"),
            {"company_name": "Vertuoza", "title": "Lead Platform Engineer",
             "url": "", "score": 60, "cv_language": Language.FR, "status": Status.BACKLOG},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(response.status_code, 204)
        created = Application.objects.get(title="Lead Platform Engineer")
        self.assertEqual(response.headers["HX-Redirect"], created.get_absolute_url())
        self.assertEqual(created.discovered_on, timezone.localdate())

    def test_quick_create_reuses_an_existing_company(self):
        self.client.post(
            reverse("tracker:quick_create"),
            {"company_name": "acme", "title": "Autre poste",
             "cv_language": Language.FR, "status": Status.BACKLOG},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(Company.objects.filter(name__iexact="acme").count(), 1)

    def test_quick_create_rejects_an_empty_form(self):
        response = self.client.post(reverse("tracker:quick_create"), {}, HTTP_HX_REQUEST="true")
        self.assertEqual(response.status_code, 422)

    def test_get_endpoints_render_their_forms(self):
        for name in ["tracker:quick_create", "tracker:stats_bar"]:
            self.assertEqual(self.client.get(reverse(name)).status_code, 200)
        for name in ["tracker:edit_notes", "tracker:add_event", "tracker:add_contact",
                     "tracker:add_document"]:
            with self.subTest(endpoint=name):
                response = self.client.get(reverse(name, args=[self.application.pk]))
                self.assertEqual(response.status_code, 200)

    def test_mutating_endpoints_refuse_get(self):
        for name in ["tracker:set_status", "tracker:advance_status", "tracker:set_follow_up"]:
            with self.subTest(endpoint=name):
                response = self.client.get(reverse(name, args=[self.application.pk]))
                self.assertEqual(response.status_code, 405)



@override_settings(MEDIA_ROOT=MEDIA)
class ApplicationFormTests(OwnedTestCase):
    def test_create_makes_the_company_when_it_is_new(self):
        response = self.client.post(
            reverse("tracker:application_create"),
            {"company_name": "Nouvelle SPRL", "company_sector": Sector.PRIVATE,
             "title": "Ingénieur système", "status": Status.BACKLOG,
             "cv_language": Language.FR, "work_mode": WorkMode.UNKNOWN},
        )
        self.assertEqual(response.status_code, 302)
        application = Application.objects.get(title="Ingénieur système")
        self.assertEqual(application.company.name, "Nouvelle SPRL")
        self.assertEqual(application.company.owner, self.user)
        self.assertEqual(application.owner, self.user)
        self.assertEqual(application.events.count(), 1)

    def test_the_same_company_name_can_exist_for_two_profiles(self):
        other = make_user("Marie")
        make_company(owner=other, name="Nouvelle SPRL")
        self.client.post(
            reverse("tracker:application_create"),
            {"company_name": "nouvelle sprl", "company_sector": Sector.PRIVATE,
             "title": "Poste", "status": Status.BACKLOG,
             "cv_language": Language.FR, "work_mode": WorkMode.UNKNOWN},
        )
        self.assertEqual(Company.objects.filter(name__iexact="nouvelle sprl").count(), 2)
        self.assertEqual(Application.objects.get().company.owner, self.user)

    def test_the_distance_help_names_the_home_base(self):
        response = self.client.get(reverse("tracker:application_create"))
        self.assertContains(response, "oiseau depuis Nivelles.")

    def test_a_follow_up_cannot_precede_the_application(self):
        response = self.client.post(
            reverse("tracker:application_create"),
            {"company_name": "Acme", "company_sector": Sector.PRIVATE, "title": "Poste",
             "status": Status.SENT, "cv_language": Language.FR, "work_mode": WorkMode.UNKNOWN,
             "applied_on": "2026-08-20", "follow_up_on": "2026-08-10"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("follow_up_on", response.context["form"].errors)

    def test_discarding_requires_a_reason(self):
        response = self.client.post(
            reverse("tracker:application_create"),
            {"company_name": "Acme", "company_sector": Sector.PRIVATE, "title": "Poste",
             "status": Status.DISCARDED, "cv_language": Language.FR,
             "work_mode": WorkMode.UNKNOWN},
        )
        self.assertIn("discard_reason", response.context["form"].errors)

    def test_score_is_bounded(self):
        response = self.client.post(
            reverse("tracker:application_create"),
            {"company_name": "Acme", "company_sector": Sector.PRIVATE, "title": "Poste",
             "status": Status.BACKLOG, "cv_language": Language.FR,
             "work_mode": WorkMode.UNKNOWN, "score": 140},
        )
        self.assertIn("score", response.context["form"].errors)

    def test_update_logs_a_status_change(self):
        application = make_application()
        self.client.post(
            reverse("tracker:application_update", args=[application.pk]),
            {"company_name": application.company.name, "company_sector": Sector.PRIVATE,
             "title": application.title, "status": Status.SENT,
             "cv_language": Language.FR, "work_mode": WorkMode.UNKNOWN},
        )
        application.refresh_from_db()
        self.assertEqual(application.status, Status.SENT)
        self.assertEqual(application.events.filter(kind=EventKind.STATUS).count(), 1)

    def test_delete(self):
        application = make_application()
        response = self.client.post(reverse("tracker:application_delete", args=[application.pk]))
        self.assertRedirects(response, reverse("tracker:application_list"))
        self.assertEqual(Application.objects.count(), 0)


@override_settings(MEDIA_ROOT=MEDIA)
class DocumentStorageTests(TestCase):
    def test_deleting_an_application_removes_its_files_from_disk(self):
        application = make_application()
        document = make_document(
            application=application, kind=DocumentKind.CV, label="CV.docx",
            file=SimpleUploadedFile("cascade.docx", b"contenu"),
        )
        path = Path(document.file.path)
        self.assertTrue(path.exists())
        with self.captureOnCommitCallbacks(execute=True):
            application.delete()
        self.assertFalse(path.exists())

    def test_deleting_a_document_removes_its_file_once_the_row_is_gone(self):
        document = make_document(
            kind=DocumentKind.CV, label="CV.docx",
            file=SimpleUploadedFile("single.docx", b"contenu"),
        )
        path = Path(document.file.path)
        with self.captureOnCommitCallbacks(execute=True) as callbacks:
            document.delete()
            # Still there while the deletion could still be rolled back.
            self.assertTrue(path.exists())
        self.assertEqual(len(callbacks), 1)
        self.assertFalse(path.exists())

    def test_a_rolled_back_deletion_keeps_the_file(self):
        document = make_document(
            kind=DocumentKind.CV, label="CV.docx",
            file=SimpleUploadedFile("rollback.docx", b"contenu"),
        )
        path, pk = Path(document.file.path), document.pk  # ``delete`` clears the pk
        with self.captureOnCommitCallbacks(execute=True) as callbacks:
            with transaction.atomic():
                document.delete()
                transaction.set_rollback(True)
        self.assertEqual(callbacks, [])
        self.assertTrue(path.exists())
        self.assertTrue(Document.objects.filter(pk=pk).exists())

    def test_the_size_is_recorded_once_and_read_from_the_row(self):
        document = make_document(
            kind=DocumentKind.CV, label="CV", file=SimpleUploadedFile("s.docx", b"12345"),
        )
        self.assertEqual(document.size_bytes, 5)
        document = Document.objects.get(pk=document.pk)
        with mock.patch.object(LocalStorageAdapter, "size", side_effect=AssertionError("appelé")):
            self.assertEqual(document.size_display, "5 o")
        # A row without the column (an extension wrote it): the storage, once.
        Document.objects.filter(pk=document.pk).update(size_bytes=None)
        document = Document.objects.get(pk=document.pk)
        self.assertEqual(document.size_display, "5 o")
        Path(document.file.path).unlink()
        self.assertEqual(document.size_display, "—")

    def test_a_provider_failure_on_delete_is_logged_not_raised(self):
        document = make_document(
            kind=DocumentKind.CV, label="CV", file=SimpleUploadedFile("d.docx", b"x"),
        )
        with mock.patch.object(LocalStorageAdapter, "delete", side_effect=StorageError("panne")):
            with self.assertLogs("tracker.models", "ERROR"):
                with self.captureOnCommitCallbacks(execute=True):
                    document.delete()
        self.assertFalse(Document.objects.filter(pk=document.pk).exists())


@override_settings(MEDIA_ROOT=MEDIA)
class DocumentDownloadTests(OwnedTestCase):
    def setUp(self):
        super().setUp()
        self.document = make_document(
            kind=DocumentKind.CV, label="CV base",
            file=SimpleUploadedFile("cv-base.docx", b"contenu"),
        )

    def test_owner_downloads_as_attachment(self):
        response = self.client.get(self.document.get_download_url())
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response.headers["Content-Disposition"])
        self.assertIn("cv-base", response.headers["Content-Disposition"])
        self.assertEqual(b"".join(getattr(response, "streaming_content")), b"contenu")

    def test_pages_link_the_download_view_not_a_file_path_nor_a_link(self):
        response = self.client.get(reverse("tracker:document_library"))
        self.assertContains(response, self.document.get_download_url())
        self.assertNotContains(response, "/media/")
        self.assertNotContains(response, "/fichiers/")

    def test_media_paths_are_not_served(self):
        """Also under DEBUG: the old ``static(MEDIA_URL)`` mount only existed
        there, and the test runner forces DEBUG off. ``file.url`` is no
        longer that path but a signed link (see PrivateFileViewTests)."""
        import importlib

        from django.urls import Resolver404, clear_url_caches, resolve

        import jobhunt.urls

        media_path = "/" + settings.MEDIA_URL.strip("/") + "/" + (self.document.file.name or "")
        self.assertTrue(self.document.file.url.startswith("/fichiers/"))
        self.assertEqual(self.client.get(media_path).status_code, 404)
        try:
            with override_settings(DEBUG=True):
                importlib.reload(jobhunt.urls)
                clear_url_caches()
                with self.assertRaises(Resolver404):
                    resolve(media_path)
                self.assertEqual(self.client.get(media_path).status_code, 404)
        finally:
            importlib.reload(jobhunt.urls)
            clear_url_caches()

    def test_download_streams_through_the_port_with_its_size(self):
        response = self.client.get(self.document.get_download_url())
        self.assertEqual(response.headers["Content-Length"], "7")
        self.assertEqual(response.headers["Content-Type"],
                         "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    def test_reassigning_an_application_carries_its_documents(self):
        other = make_user("Marie")
        application = make_application()
        attached = make_document(
            application=application, kind=DocumentKind.CV, label="CV",
            file=SimpleUploadedFile("attached.docx", b"x"),
        )
        application = Application.objects.get(pk=application.pk)
        application.owner = other
        application.company = make_company(owner=other, name="Acme bis")
        application.save()
        attached.refresh_from_db()
        self.assertEqual(attached.owner, other)

    def test_someone_else_s_document_is_a_404(self):
        self.client.force_login(make_user("Marie"))
        self.assertEqual(self.client.get(self.document.get_download_url()).status_code, 404)

    def test_missing_file_is_a_404(self):
        Path(self.document.file.path).unlink()
        self.assertEqual(self.client.get(self.document.get_download_url()).status_code, 404)


class SlugTests(TestCase):
    def test_slugs_stay_unique_for_identical_titles(self):
        company = make_company(name="HMS Networks")
        first = make_application(company=company, title="R&D Firmware Engineer")
        second = Application.objects.create(
            owner=company.owner, company=company, title="R&D Firmware Engineer",
            status=Status.BACKLOG,
        )
        self.assertNotEqual(first.slug, second.slug)
        self.assertTrue(second.slug.endswith("-2"))

    def test_company_slug_is_generated(self):
        company = make_company(name="Egov Select (Sûreté de l'État)")
        self.assertEqual(company.slug, "egov-select-surete-de-letat")


@override_settings(MEDIA_ROOT=MEDIA)
class IsolationTests(OwnedTestCase):
    """Signed in as Lionel, every URL that names one of Marie's rows is a 404
    and leaves the row exactly as it was."""

    def setUp(self):
        super().setUp()
        self.other = make_user("Marie")
        self.theirs = make_application(owner=self.other, company_name="Widgets", status=Status.TO_APPLY)
        self.event = ActivityEvent.objects.create(application=self.theirs, title="Note")
        self.contact = Contact.objects.create(application=self.theirs, name="Quelqu'un")
        self.document = make_document(
            application=self.theirs, kind=DocumentKind.CV, label="CV",
            file=SimpleUploadedFile("theirs.docx", b"x"),
        )
        self.gap = make_gap(owner=self.other, name="Cobol")
        make_application(status=Status.SENT)  # Lionel's own, for the list views

    def test_lists_never_show_another_profile(self):
        for name in ["tracker:dashboard", "tracker:pipeline", "tracker:application_list",
                     "tracker:document_library", "tracker:insights"]:
            with self.subTest(page=name):
                response = self.client.get(reverse(name))
                self.assertEqual(response.status_code, 200)
                self.assertNotContains(response, "Widgets")
                self.assertNotContains(response, "Cobol")
        self.assertEqual(self.client.get(reverse("tracker:dashboard")).context["nav_counters"]["tracked"], 1)

    def test_every_row_endpoint_is_a_404_for_another_profile(self):
        application = self.theirs.pk
        attempts = [
            ("get", "tracker:application_detail", [application], {}),
            ("get", "tracker:application_update", [application], {}),
            ("post", "tracker:application_update", [application], {"title": "Pirate"}),
            ("post", "tracker:application_delete", [application], {}),
            ("post", "tracker:set_status", [application], {"status": Status.SENT}),
            ("post", "tracker:advance_status", [application], {}),
            ("post", "tracker:set_follow_up", [application], {"in_days": "7"}),
            ("post", "tracker:mark_followed_up", [application], {}),
            ("post", "tracker:edit_notes", [application], {"personal_notes": "Pirate"}),
            ("get", "tracker:edit_notes", [application], {}),
            ("post", "tracker:add_event", [application], {"happened_on": "2026-09-01", "kind": "note", "title": "Pirate"}),
            ("post", "tracker:add_contact", [application], {"name": "Pirate"}),
            ("get", "tracker:add_document", [application], {}),
            ("post", "tracker:delete_event", [self.event.pk], {}),
            ("post", "tracker:delete_contact", [self.contact.pk], {}),
            ("post", "tracker:delete_document", [self.document.pk], {}),
            ("get", "tracker:document_download", [self.document.pk], {}),
            ("post", "tracker:set_gap_status", [self.gap.pk], {"status": GapStatus.DONE}),
        ]
        for method, name, args, data in attempts:
            with self.subTest(endpoint=name, method=method):
                call = getattr(self.client, method)
                response = call(reverse(name, args=args), data, HTTP_HX_REQUEST="true")
                self.assertEqual(response.status_code, 404)

        self.theirs.refresh_from_db()
        self.assertEqual(self.theirs.status, Status.TO_APPLY)
        self.assertEqual(self.theirs.title, "Senior DevOps Engineer")
        self.assertEqual(self.theirs.personal_notes, "")
        self.assertIsNone(self.theirs.follow_up_on)
        self.assertEqual(self.theirs.events.count(), 1)
        self.assertEqual(self.theirs.contacts.count(), 1)
        self.assertEqual(self.theirs.documents.count(), 1)
        self.gap.refresh_from_db()
        self.assertEqual(self.gap.status, GapStatus.TODO)

    def test_forms_only_offer_the_profile_s_own_reference_data(self):
        make_platform(owner=self.other, name="Leur plateforme")
        response = self.client.get(reverse("tracker:application_create"))
        self.assertNotContains(response, "Leur plateforme")
        self.assertNotContains(response, "Widgets")  # company datalist

    def test_a_company_from_another_profile_cannot_be_borrowed(self):
        """The form names companies, never their ids, so the worst case is a
        namesake created on Lionel's side."""
        self.client.post(
            reverse("tracker:application_create"),
            {"company_name": "Widgets", "company_sector": Sector.PRIVATE, "title": "Poste",
             "status": Status.BACKLOG, "cv_language": Language.FR, "work_mode": WorkMode.UNKNOWN},
        )
        mine = Application.objects.get(owner=self.user, title="Poste")
        self.assertNotEqual(mine.company, self.theirs.company)
        self.assertEqual(mine.company.owner, self.user)


class RichTextTests(TestCase):
    def render(self, value):
        from tracker.templatetags.tracker_extras import richtext

        return str(richtext(value))

    def test_bullets_bold_and_links(self):
        html = self.render("- **Python** exigé\n- Voir https://example.com")
        self.assertIn("<ul>", html)
        self.assertIn("<strong>Python</strong>", html)
        self.assertIn('href="https://example.com"', html)

    def test_html_is_escaped(self):
        html = self.render("<script>alert(1)</script>")
        self.assertNotIn("<script>", html)
        self.assertIn("&lt;script&gt;", html)

    def test_headings_and_empty_input(self):
        self.assertIn("<h3>", self.render("# Titre"))
        self.assertEqual(self.render(""), "")
        self.assertEqual(self.render(None), "")

    def test_french_pluralisation(self):
        from tracker.templatetags.tracker_extras import pluralize_fr

        self.assertEqual(pluralize_fr(0), "")
        self.assertEqual(pluralize_fr(1), "")
        self.assertEqual(pluralize_fr(2), "s")
        self.assertEqual(pluralize_fr(2, "est,sont"), "sont")
        self.assertEqual(pluralize_fr(1, "est,sont"), "est")


class WorkModeDetectionTests(TestCase):
    def detect(self, text):
        from tracker.management.commands.import_legacy import guess_work_mode

        return guess_work_mode(text)

    def test_negated_remote_reads_as_on_site(self):
        self.assertEqual(
            self.detect("Bruxelles — 5 jours sur site, pas de télétravail"), WorkMode.ON_SITE
        )

    def test_plain_remote(self):
        self.assertEqual(self.detect("Bruxelles / Remote Europe"), WorkMode.REMOTE)

    def test_on_site(self):
        self.assertEqual(self.detect("Nivelles — sur place"), WorkMode.ON_SITE)

    def test_unknown(self):
        self.assertEqual(self.detect("Bruxelles"), WorkMode.UNKNOWN)

    def test_hybrid(self):
        self.assertEqual(self.detect("Hybride, 2 jours à Nivelles"), WorkMode.HYBRID)


@override_settings(MEDIA_ROOT=MEDIA)
class ImportCommandTests(OwnedTestCase):
    """Runs against the real workbook and folders when they are present."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.workbook = Path(settings.LEGACY_WORKBOOK)

    def setUp(self):
        super().setUp()
        if not self.workbook.exists():
            self.skipTest("Le classeur d'origine n'est pas disponible.")

    def test_everything_imported_belongs_to_the_profile(self):
        call_command("import_legacy", "--skip-files", verbosity=0)
        for model in (Application, Company, Platform, SkillGap):
            with self.subTest(model=model.__name__):
                self.assertFalse(model.objects.exclude(owner=self.user).exists())

    def test_user_option_and_ambiguity(self):
        from django.core.management.base import CommandError

        make_user("Marie", username="marie")
        with self.assertRaises(CommandError):
            call_command("import_legacy", "--skip-files", verbosity=0)
        with self.assertRaises(CommandError):
            call_command("import_legacy", "--skip-files", user="personne", verbosity=0)
        call_command("import_legacy", "--skip-files", user="marie", verbosity=0)
        self.assertTrue(Application.objects.filter(owner__username="marie").exists())
        self.assertFalse(Application.objects.filter(owner=self.user).exists())

    def test_import_is_complete_and_idempotent(self):
        call_command("import_legacy", verbosity=0)
        first = {
            "applications": Application.objects.count(),
            "companies": Company.objects.count(),
            "documents": Document.objects.count(),
            "platforms": Platform.objects.count(),
            "gaps": SkillGap.objects.count(),
        }
        self.assertEqual(first["applications"], 24)
        self.assertEqual(Application.objects.discarded().count(), 14)
        self.assertEqual(first["gaps"], 7)

        call_command("import_legacy", verbosity=0)
        self.assertEqual(Application.objects.count(), first["applications"])
        self.assertEqual(Company.objects.count(), first["companies"])
        self.assertEqual(Document.objects.count(), first["documents"])

    def test_every_offer_folder_is_matched_and_parsed(self):
        call_command("import_legacy", verbosity=0)
        folders = [
            path for path in (Path(settings.LEGACY_ROOT) / "Offres").iterdir() if path.is_dir()
        ]
        self.assertEqual(
            Application.objects.exclude(legacy_folder="").count(), len(folders)
        )
        for application in Application.objects.exclude(legacy_folder=""):
            with self.subTest(folder=application.legacy_folder):
                self.assertTrue(application.strengths)
                self.assertTrue(application.weaknesses)
                self.assertTrue(application.strategy)
                self.assertTrue(application.posting_raw)
                self.assertEqual(
                    application.documents.filter(kind=DocumentKind.CV, is_primary=True).count(), 1
                )

    def test_skip_files_leaves_the_media_root_alone(self):
        call_command("import_legacy", "--skip-files", verbosity=0)
        self.assertEqual(Document.objects.count(), 0)
        self.assertTrue(Application.objects.exists())

    def test_import_survives_a_missing_workbook(self):
        from django.core.management.base import CommandError

        with self.assertRaises(CommandError):
            call_command("import_legacy", workbook="/nowhere/absent.xlsx", verbosity=0)


class PluginFrameworkTests(TestCase):
    """Le cœur doit fonctionner avec zéro, une ou une extension cassée."""

    def _with_entry_points(self, entries):
        """Fait tourner la découverte sur une liste de points d'entrée factices."""
        from unittest import mock

        from jobhunt import plugins

        patcher = mock.patch.object(plugins, "entry_points", return_value=entries)
        patcher.start()
        self.addCleanup(patcher.stop)
        plugins.get_plugins.cache_clear()
        self.addCleanup(plugins.get_plugins.cache_clear)
        return plugins

    def test_no_plugins_yields_empty_registry(self):
        plugins = self._with_entry_points([])
        request = RequestFactory().get("/")
        request.user = make_user("Lionel")
        self.assertEqual(plugins.get_plugins(), ())
        self.assertEqual(plugins.plugin_apps(), [])
        self.assertEqual(plugins.plugin_nav_items(), [])
        self.assertEqual(plugins.plugin_nav_badges(request), {})
        self.assertEqual(plugins.plugin_templates("icon_templates"), [])

    def test_broken_plugin_is_skipped(self):
        from unittest import mock

        broken = mock.Mock()
        broken.name = "casse"
        broken.load.side_effect = ImportError("paquet manquant")
        plugins = self._with_entry_points([broken])
        self.assertEqual(plugins.get_plugins(), ())

    def test_declared_plugin_is_exposed(self):
        from types import SimpleNamespace
        from unittest import mock

        descriptor = SimpleNamespace()
        descriptor.app = "exemple.apps.ExempleConfig"
        descriptor.nav_items = [("exemple:index", "Exemple", "gauge", None)]
        descriptor.nav_badges = ""
        descriptor.icon_templates = ["exemple/icons.html"]
        descriptor.application_panels = []
        entry = mock.Mock()
        entry.name = "exemple"
        entry.load.return_value = descriptor
        plugins = self._with_entry_points([entry])
        self.assertEqual(plugins.plugin_apps(), ["exemple.apps.ExempleConfig"])
        self.assertEqual(len(plugins.plugin_nav_items()), 1)
        self.assertEqual(plugins.plugin_templates("icon_templates"), ["exemple/icons.html"])

    def test_plugin_required_apps_are_loaded_once_before_the_plugins(self):
        from types import SimpleNamespace
        from unittest import mock

        entries = []
        for name in ("one", "two"):
            entry = mock.Mock()
            entry.name = name
            entry.load.return_value = SimpleNamespace(app=name, required_apps=("django_q",))
            entries.append(entry)
        plugins = self._with_entry_points(entries)
        self.assertEqual(plugins.plugin_apps(), ["django_q", "one", "two"])

    def test_navigation_pages_render_without_plugins(self):
        """Les pages du cœur ne dépendent d'aucune extension."""
        from unittest import mock

        self.client.force_login(make_user("Lionel"))

        for function in ("plugin_nav_items", "plugin_nav_badges"):
            patcher = mock.patch(
                f"tracker.context_processors.{function}",
                return_value=[] if function == "plugin_nav_items" else {},
            )
            patcher.start()
            self.addCleanup(patcher.stop)
        patcher = mock.patch("tracker.context_processors.plugin_templates", return_value=[])
        patcher.start()
        self.addCleanup(patcher.stop)

        response = self.client.get(reverse("tracker:dashboard"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["plugin_icon_templates"], [])


class AutoMigrateRunserverTests(TestCase):
    """``runserver`` writes and applies migrations in development, warns elsewhere."""

    PLAN = [(mock.Mock(app_label="tracker"), False), (mock.Mock(app_label="jobhunt_ai"), False)]

    def run_check(self, *, auto_migrate: bool = True, plan=(), changes=None):
        """Run ``check_migrations`` with the plan and the detected changes
        stubbed; returns the ``call_command`` mock and what was printed."""
        from io import StringIO

        from tracker.management.commands.runserver import Command

        self.enterContext(override_settings(AUTO_MIGRATE=auto_migrate))
        self.enterContext(
            mock.patch(
                "django.db.migrations.executor.MigrationExecutor.migration_plan",
                return_value=list(plan),
            )
        )
        if changes is not None:
            outcome: dict[str, Any] = (
                {"side_effect": changes} if isinstance(changes, Exception) else {"return_value": changes}
            )
            self.enterContext(
                mock.patch(
                    "django.db.migrations.autodetector.MigrationAutodetector.changes", **outcome
                )
            )
        call = self.enterContext(mock.patch("tracker.management.commands.runserver.call_command"))
        command = Command(stdout=StringIO())
        command.check_migrations()
        return call, command.stdout.getvalue()

    def test_runserver_resolves_to_the_override(self):
        from django.core.management import get_commands

        self.assertEqual(get_commands()["runserver"], "tracker")

    def test_pending_migrations_are_applied_in_development(self):
        call, output = self.run_check(plan=self.PLAN, changes={})
        call.assert_called_once()
        self.assertEqual(call.call_args.args[0], "migrate")
        self.assertFalse(call.call_args.kwargs["interactive"])
        self.assertIn("2 migration(s) en attente (jobhunt_ai, tracker)", output)

    def test_outside_development_django_only_warns(self):
        call, output = self.run_check(auto_migrate=False, plan=self.PLAN, changes={})
        call.assert_not_called()
        self.assertIn("unapplied migration", output)

    def test_nothing_to_do_stays_quiet(self):
        call, output = self.run_check(plan=[], changes={})
        call.assert_not_called()
        self.assertEqual(output, "")

    def test_model_changes_get_their_migration_before_migrate(self):
        call, output = self.run_check(plan=self.PLAN, changes={"tracker": [object()]})
        self.assertEqual(
            [c.args for c in call.call_args_list], [("makemigrations", "tracker"), ("migrate",)]
        )
        self.assertFalse(call.call_args_list[0].kwargs["interactive"])
        self.assertIn("Modèles modifiés sans migration (tracker)", output)

    def test_a_question_for_a_human_stops_the_generation(self):
        from tracker.management.commands.runserver import NeedsHuman

        call, output = self.run_check(
            plan=self.PLAN, changes=NeedsHuman("renommage possible de company.notes en remarks")
        )
        self.assertEqual([c.args[0] for c in call.call_args_list], ["migrate"])
        self.assertIn("renommage possible de company.notes en remarks", output)
        self.assertIn("makemigrations", output)

    def test_the_guard_refuses_what_django_would_ask_about(self):
        """A real change detector run on a doctored model state."""
        from django.apps import apps
        from django.db import models
        from django.db.migrations.autodetector import MigrationAutodetector
        from django.db.migrations.loader import MigrationLoader
        from django.db.migrations.state import ProjectState
        from django.utils import translation

        from tracker.management.commands.runserver import GuardedQuestioner, NeedsHuman

        loader = MigrationLoader(None, ignore_no_migrations=True)

        def detect(mutate):
            to_state = ProjectState.from_apps(apps)
            mutate(to_state.models["tracker", "company"])
            detector = MigrationAutodetector(loader.project_state(), to_state, GuardedQuestioner())
            with translation.override(None):
                return detector.changes(graph=loader.graph)

        def rename(model_state):
            model_state.fields["remarks"] = model_state.fields.pop("notes")

        def mandatory(model_state):
            model_state.fields["vat"] = models.CharField(max_length=20)

        def harmless(model_state):
            model_state.fields["vat"] = models.CharField(max_length=20, blank=True, default="")

        with self.assertRaises(NeedsHuman) as caught:
            detect(rename)
        self.assertIn("renommage possible de company.notes en remarks", str(caught.exception))
        with self.assertRaises(NeedsHuman) as caught:
            detect(mandatory)
        self.assertIn("company.vat est obligatoire", str(caught.exception))
        self.assertEqual(sorted(detect(harmless)), ["tracker"])

    def test_the_french_locale_does_not_fake_changes_in_contrib_apps(self):
        """Verbose names are compared with the English strings frozen in
        Django's migrations; the detector must run with translations off."""
        from django.utils import translation

        with translation.override("fr-be"):
            call, output = self.run_check(plan=[])
        call.assert_not_called()
        self.assertEqual(output, "")


class OwnershipMigrationTests(TransactionTestCase):
    """The ownership migrations against rows recorded before accounts existed.

    A ``TransactionTestCase``: SQLite refuses schema changes inside the
    transaction a ``TestCase`` wraps around each test.
    """

    BEFORE = [("tracker", "0001_initial")]

    def rewind(self):
        executor = MigrationExecutor(connection)
        executor.migrate(self.BEFORE)
        return executor.loader.project_state(self.BEFORE).apps

    def replay(self):
        call_command("migrate", verbosity=0)

    def tearDown(self):
        self.replay()

    def seed(self, apps):
        Company = apps.get_model("tracker", "Company")
        Application = apps.get_model("tracker", "Application")
        Document = apps.get_model("tracker", "Document")
        Platform = apps.get_model("tracker", "Platform")
        SkillGap = apps.get_model("tracker", "SkillGap")
        company = Company.objects.create(name="Acme", slug="acme")
        Application.objects.create(company=company, title="Poste", slug="acme-poste")
        Document.objects.create(kind="cv", label="CV", file="documents/bibliotheque/cv.docx")
        Platform.objects.create(name="LinkedIn")
        SkillGap.objects.create(name="Terraform")

    def assert_everything_belongs_to(self, user):
        for model in (Application, Company, Document, Platform, SkillGap):
            with self.subTest(model=model.__name__):
                self.assertEqual(model.objects.count(), 1)
                self.assertEqual(model.objects.get().owner, user)

    def test_orphans_are_parked_on_a_local_account_awaiting_onboarding(self):
        apps = self.rewind()
        self.seed(apps)
        self.replay()

        user = User.objects.get()
        self.assertEqual(user.username, LOCAL_USERNAME)
        self.assertFalse(user.has_usable_password())
        self.assertFalse(profile_for(user).is_onboarded)
        self.assertTrue(preferences_for(user).pk)
        self.assert_everything_belongs_to(user)

    def test_orphans_go_to_the_only_existing_account(self):
        apps = self.rewind()
        # ``auth`` is untouched by the rewind: the live model is the right one.
        admin = User.objects.create_user("admin", password="x")
        self.seed(apps)
        self.replay()

        self.assertEqual(User.objects.count(), 1)
        self.assert_everything_belongs_to(User.objects.get(pk=admin.pk))
        self.assertFalse(profile_for(User.objects.get()).is_onboarded)

    def test_an_empty_database_gets_no_account(self):
        self.rewind()
        self.replay()
        self.assertEqual(get_user_model().objects.count(), 0)


# ---------------------------------------------------------------------------
# The hexagon: ports, adapters, rules and use cases
# ---------------------------------------------------------------------------


@override_settings(MEDIA_ROOT=MEDIA)
class PersistenceContractTests(TestCase):
    """Every port behaves the same on the ORM adapter and on the in-memory
    one — owner scoping, ``NotFound``, ``None`` last in every sort."""

    ON = dt.date(2026, 9, 2)

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.lionel = make_user("Lionel", username="lionel")
        cls.marie = make_user("Marie", username="marie")

    def each_adapter(self, check):
        for name, store in (("django_orm", DjangoPersistence()), ("memory", MemoryPersistence())):
            with self.subTest(adapter=name):
                check(store)

    def application(self, store, owner, company_name, **fields):
        company = store.companies.get_or_create(owner, company_name, Sector.PRIVATE)
        defaults = {"title": "Poste", "status": Status.TO_APPLY, "score": 72}
        defaults.update(fields)
        return store.applications.add(Application(owner=owner, company=company, **defaults))

    def upload(self, store, application, name, kind=DocumentKind.CV, primary=False, **fields):
        return store.documents.add(
            Document(
                application=application, kind=kind, label=name, is_primary=primary,
                file=SimpleUploadedFile(name, b"x"), **fields,
            )
        )

    def test_get_is_scoped_on_the_owner(self):
        def check(store):
            mine = self.application(store, self.lionel, "Alpha")
            self.assertEqual(store.applications.get(self.lionel, mine.pk).pk, mine.pk)
            self.assertEqual(store.applications.detail(self.lionel, mine.pk).pk, mine.pk)
            for lookup in (store.applications.get, store.applications.detail):
                with self.assertRaises(NotFound):
                    lookup(self.marie, mine.pk)
                with self.assertRaises(NotFound):
                    lookup(self.lionel, mine.pk + 1000)

            loaded = store.applications.get(self.lionel, mine.pk)
            with self.assertNumQueries(0):
                self.assertEqual(loaded.company.name, "Alpha")
                self.assertIsNone(loaded.source_platform)
            if isinstance(store, DjangoPersistence):
                with self.assertNumQueries(0):
                    self.assertIsNotNone(loaded.owner.preferences)

            event = store.events.add(mine, EventKind.NOTE, "Note")
            self.assertEqual(store.events.get(self.lionel, event.pk).pk, event.pk)
            with self.assertRaises(NotFound):
                store.events.get(self.marie, event.pk)

            contact = store.contacts.add(Contact(application=mine, name="Marie Dupont"))
            self.assertEqual(store.contacts.get(self.lionel, contact.pk).pk, contact.pk)
            with self.assertRaises(NotFound):
                store.contacts.get(self.marie, contact.pk)

            document = self.upload(store, mine, "CV.docx")
            self.assertEqual(store.documents.get(self.lionel, document.pk).pk, document.pk)
            with self.assertRaises(NotFound):
                store.documents.get(self.marie, document.pk)

            gap = SkillGap(owner=self.lionel, name="Terraform")
            store.skill_gaps.save(gap)
            self.assertEqual(store.skill_gaps.get(self.lionel, gap.pk).pk, gap.pk)
            with self.assertRaises(NotFound):
                store.skill_gaps.get(self.marie, gap.pk)

        self.each_adapter(check)

    def test_add_assigns_pk_slug_and_timestamps(self):
        def check(store):
            first = self.application(store, self.lionel, "Alpha", title="R&D Engineer")
            self.assertIsNotNone(first.pk)
            self.assertEqual(first.slug, "alpha-rd-engineer")
            self.assertIsNotNone(first.created_at)
            self.assertIsNotNone(first.updated_at)
            second = self.application(store, self.lionel, "Alpha", title="R&D Engineer")
            self.assertEqual(second.slug, "alpha-rd-engineer-2")

        self.each_adapter(check)

    def test_save_with_fields_persists_them_and_moves_updated_at(self):
        def check(store):
            application = self.application(store, self.lionel, "Alpha")
            mark = timezone.now() - dt.timedelta(days=1)
            application.updated_at = mark
            application.follow_up_on = dt.date(2026, 10, 1)
            application.status = Status.SENT  # touched but not listed: must not be written
            store.applications.save(application, fields=["follow_up_on"])
            reloaded = store.applications.get(self.lionel, application.pk)
            self.assertEqual(reloaded.follow_up_on, dt.date(2026, 10, 1))
            self.assertEqual(reloaded.status, Status.TO_APPLY)
            self.assertGreater(reloaded.updated_at, mark)

        self.each_adapter(check)

    def test_remove_takes_the_satellites_along(self):
        def check(store):
            application = self.application(store, self.lionel, "Alpha")
            event = store.events.add(application, EventKind.NOTE, "Note")
            contact = store.contacts.add(Contact(application=application, name="X"))
            document = self.upload(store, application, "CV.docx")
            pks = (application.pk, event.pk, contact.pk, document.pk)
            store.applications.remove(application)
            with self.assertRaises(NotFound):
                store.applications.get(self.lionel, pks[0])
            with self.assertRaises(NotFound):
                store.events.get(self.lionel, pks[1])
            with self.assertRaises(NotFound):
                store.contacts.get(self.lionel, pks[2])
            with self.assertRaises(NotFound):
                store.documents.get(self.lionel, pks[3])
            self.assertEqual(store.documents.count(self.lionel), 0)

        self.each_adapter(check)

    def test_by_status_orders_by_score_then_name_with_unscored_last(self):
        def check(store):
            gamma = self.application(store, self.lionel, "Gamma", score=90)
            beta = self.application(store, self.lionel, "Beta", score=None)
            alpha = self.application(store, self.lionel, "Alpha", score=90)
            delta = self.application(store, self.lionel, "Delta", score=50, status=Status.SENT)
            self.application(store, self.marie, "Alpha", score=99)
            rows = store.applications.by_status(self.lionel, [Status.TO_APPLY])
            self.assertEqual([row.pk for row in rows], [alpha.pk, gamma.pk, beta.pk])
            rows = store.applications.by_status(self.lionel, [Status.TO_APPLY, Status.SENT])
            self.assertEqual([row.pk for row in rows], [alpha.pk, gamma.pk, delta.pk, beta.pk])

        self.each_adapter(check)

    def test_status_counts(self):
        def check(store):
            self.assertEqual(store.applications.status_counts(self.lionel), {})
            self.application(store, self.lionel, "Alpha")
            self.application(store, self.lionel, "Beta")
            self.application(store, self.lionel, "Gamma", status=Status.SENT)
            self.application(store, self.marie, "Alpha", status=Status.SENT)
            self.assertEqual(
                store.applications.status_counts(self.lionel),
                {Status.TO_APPLY: 2, Status.SENT: 1},
            )

        self.each_adapter(check)

    def test_needing_attention_matches_due_follow_ups_and_stale_rows(self):
        on = self.ON

        def check(store):
            due = self.application(
                store, self.lionel, "Alpha", status=Status.SENT,
                applied_on=on - dt.timedelta(days=3), follow_up_on=on - dt.timedelta(days=1),
            )
            stale = self.application(
                store, self.lionel, "Beta", status=Status.SENT,
                applied_on=on - dt.timedelta(days=20),
            )
            interview_due = self.application(
                store, self.lionel, "Gamma", status=Status.INTERVIEW,
                applied_on=on - dt.timedelta(days=10), follow_up_on=on,
            )
            # Fresh, reminder later; not in flight; closed; someone else's.
            self.application(
                store, self.lionel, "Delta", status=Status.SENT,
                applied_on=on - dt.timedelta(days=3), follow_up_on=on + dt.timedelta(days=5),
            )
            self.application(
                store, self.lionel, "Epsilon", status=Status.TO_APPLY,
                follow_up_on=on - dt.timedelta(days=1),
            )
            self.application(
                store, self.lionel, "Zeta", status=Status.REJECTED,
                applied_on=on - dt.timedelta(days=30), closed_on=on,
            )
            self.application(
                store, self.marie, "Alpha", status=Status.SENT,
                applied_on=on - dt.timedelta(days=20),
            )
            rows = store.applications.needing_attention(self.lionel, stale_days=14, on=on)
            self.assertEqual([row.pk for row in rows], [due.pk, interview_due.pk, stale.pk])
            self.assertEqual(
                store.applications.attention_count(self.lionel, stale_days=14, on=on), 3
            )
            self.assertEqual(
                store.applications.attention_count(self.lionel, stale_days=30, on=on), 2
            )

        self.each_adapter(check)

    def test_scores_skips_unscored_and_written_off_offers(self):
        def check(store):
            self.application(store, self.lionel, "Alpha", score=80)
            self.application(store, self.lionel, "Beta", score=None)
            self.application(store, self.lionel, "Gamma", score=60, status=Status.BACKLOG)
            self.application(store, self.lionel, "Delta", score=70, status=Status.SENT)
            self.application(store, self.lionel, "Epsilon", score=99, status=Status.DISCARDED)
            self.application(store, self.marie, "Alpha", score=10)
            scores = store.applications.scores(self.lionel, [Status.TO_APPLY, Status.BACKLOG])
            self.assertEqual(sorted(scores), [60, 80])
            self.assertEqual(store.applications.scores(self.lionel, [Status.DISCARDED]), [])

        self.each_adapter(check)

    def test_companies_are_found_case_insensitively_one_list_per_owner(self):
        def check(store):
            created = store.companies.get_or_create(self.lionel, "Alpha", Sector.PRIVATE)
            self.assertIsNotNone(created.pk)
            self.assertEqual(created.slug, "alpha")
            same = store.companies.get_or_create(self.lionel, "  alpha ", Sector.PUBLIC)
            self.assertEqual(same.pk, created.pk)
            self.assertEqual(same.sector, Sector.PUBLIC)
            untouched = store.companies.get_or_create(self.lionel, "ALPHA")
            self.assertEqual(untouched.pk, created.pk)
            self.assertEqual(untouched.sector, Sector.PUBLIC)
            theirs = store.companies.get_or_create(self.marie, "alpha")
            self.assertNotEqual(theirs.pk, created.pk)
            self.assertEqual(theirs.owner_id, self.marie.pk)
            self.assertEqual(store.companies.get_or_create(self.lionel, "Beta").sector, Sector.UNKNOWN)

        self.each_adapter(check)

    def test_documents_inherit_the_owner_demote_by_kind_and_count(self):
        def check(store):
            application = self.application(store, self.lionel, "Alpha")
            first = self.upload(store, application, "CV1.docx", primary=True)
            second = self.upload(store, application, "CV2.docx", primary=True)
            letter = self.upload(
                store, application, "LM.docx", kind=DocumentKind.COVER_LETTER, primary=True
            )
            self.assertEqual(first.owner_id, self.lionel.pk)

            store.documents.demote_primary(application, DocumentKind.CV)
            self.assertFalse(store.documents.get(self.lionel, first.pk).is_primary)
            self.assertFalse(store.documents.get(self.lionel, second.pk).is_primary)
            self.assertTrue(store.documents.get(self.lionel, letter.pk).is_primary)

            library = store.documents.add(
                Document(owner=self.lionel, kind=DocumentKind.CV, label="Base",
                         file=SimpleUploadedFile("base.docx", b"x"))
            )
            self.assertIsNone(library.application_id)
            self.assertEqual(store.documents.count(self.lionel), 4)
            self.assertEqual(store.documents.count(self.marie), 0)
            with self.assertRaises(ValueError):
                store.documents.add(
                    Document(owner=self.marie, application=application, kind=DocumentKind.CV,
                             label="Intrus", file=SimpleUploadedFile("x.docx", b"x"))
                )
            pk = library.pk
            store.documents.remove(library)
            with self.assertRaises(NotFound):
                store.documents.get(self.lionel, pk)
            self.assertEqual(store.documents.count(self.lionel), 3)

        self.each_adapter(check)

    def test_contacts_and_events_are_removed_in_place(self):
        def check(store):
            application = self.application(store, self.lionel, "Alpha")
            contact = store.contacts.add(Contact(application=application, name="Marie Dupont"))
            event = store.events.add(
                application, EventKind.CALL, "Appel", "20 minutes.", dt.date(2026, 9, 1)
            )
            self.assertEqual(event.happened_on, dt.date(2026, 9, 1))
            self.assertEqual(event.detail, "20 minutes.")
            self.assertEqual(store.events.add(application, EventKind.NOTE, "x").happened_on,
                             timezone.localdate())
            contact_pk, event_pk = contact.pk, event.pk
            store.contacts.remove(contact)
            store.events.remove(event)
            with self.assertRaises(NotFound):
                store.contacts.get(self.lionel, contact_pk)
            with self.assertRaises(NotFound):
                store.events.get(self.lionel, event_pk)

        self.each_adapter(check)

    def test_skill_gaps_save_only_the_given_fields(self):
        def check(store):
            gap = SkillGap(owner=self.lionel, name="Terraform")
            store.skill_gaps.save(gap)
            self.assertIsNotNone(gap.pk)
            gap.status = GapStatus.DOING
            gap.demand_count = 99  # touched but not listed: must not be written
            store.skill_gaps.save(gap, fields=["status"])
            reloaded = store.skill_gaps.get(self.lionel, gap.pk)
            self.assertEqual(reloaded.status, GapStatus.DOING)
            self.assertEqual(reloaded.demand_count, 0)

        self.each_adapter(check)

    def test_companies_with_case_variants_resolve_to_the_first_one(self):
        """The uniqueness constraint is on the exact name: "Acme" and "ACME"
        can coexist (admin, import) and must not make the lookup raise."""

        def check(store):
            if isinstance(store, DjangoPersistence):
                first = Company.objects.create(owner=self.lionel, name="Acme")
                Company.objects.create(owner=self.lionel, name="ACME")
            else:
                table = store.companies._table
                first = table.insert(Company(owner=self.lionel, name="Acme", slug="acme"))
                table.insert(Company(owner=self.lionel, name="ACME", slug="acme-2"))
            self.assertEqual(store.companies.get_or_create(self.lionel, "acme").pk, first.pk)

        self.each_adapter(check)

    def test_an_application_cannot_borrow_another_owner_s_company(self):
        def check(store):
            theirs = store.companies.get_or_create(self.marie, "Alpha")
            with self.assertRaises(ValueError):
                store.applications.add(Application(owner=self.lionel, company=theirs, title="Poste"))

        self.each_adapter(check)


class DomainRuleTests(SimpleTestCase):
    """The rules on bare instances: no database, no adapter."""

    TODAY = dt.date(2026, 9, 2)

    def test_plan_transition_reports_what_it_touched(self):
        application = Application(status=Status.TO_APPLY)
        transition = domain.plan_transition(
            application, Status.SENT, today=self.TODAY, follow_up_days=10
        )
        assert transition is not None
        self.assertEqual(transition.fields, ("status", "applied_on", "follow_up_on", "closed_on"))
        self.assertEqual(transition.event_kind, EventKind.APPLIED)
        self.assertEqual(transition.event_title, "À postuler → Candidature envoyée")
        self.assertEqual(application.applied_on, self.TODAY)
        self.assertEqual(application.follow_up_on, self.TODAY + dt.timedelta(days=10))
        self.assertIsNone(application.closed_on)

        closing = domain.plan_transition(
            application, Status.REJECTED, today=self.TODAY, follow_up_days=10
        )
        assert closing is not None
        self.assertEqual(closing.fields, ("status", "closed_on", "follow_up_on"))
        self.assertEqual(closing.event_kind, EventKind.REJECTION)
        self.assertIsNone(application.follow_up_on)
        self.assertEqual(application.closed_on, self.TODAY)

    def test_the_follow_up_delay_is_only_resolved_when_a_reminder_is_needed(self):
        never = mock.Mock(side_effect=AssertionError("résolu pour rien"))
        application = Application(
            status=Status.SENT, applied_on=self.TODAY, follow_up_on=self.TODAY
        )
        self.assertIsNotNone(
            domain.plan_transition(
                application, Status.INTERVIEW, today=self.TODAY, follow_up_days=never
            )
        )
        never.assert_not_called()
        resolved = mock.Mock(return_value=3)
        application = Application(status=Status.TO_APPLY)
        domain.plan_transition(application, Status.SENT, today=self.TODAY, follow_up_days=resolved)
        self.assertEqual(application.follow_up_on, self.TODAY + dt.timedelta(days=3))

    def test_a_same_or_unknown_status_is_not_a_transition(self):
        application = Application(status=Status.TO_APPLY)
        for value in (Status.TO_APPLY, "nope", ""):
            with self.subTest(value=value):
                self.assertIsNone(
                    domain.plan_transition(application, value, today=self.TODAY, follow_up_days=10)
                )
        self.assertEqual(application.status, Status.TO_APPLY)

    def test_stale_and_attention_predicates(self):
        sent = Application(status=Status.SENT, applied_on=self.TODAY - dt.timedelta(days=14))
        self.assertTrue(domain.is_stale(sent, stale_days=14, today=self.TODAY))
        self.assertFalse(domain.is_stale(sent, stale_days=15, today=self.TODAY))
        never = mock.Mock(side_effect=AssertionError("résolu pour rien"))
        self.assertFalse(
            domain.is_stale(Application(status=Status.TO_APPLY), stale_days=never, today=self.TODAY)
        )
        due = Application(status=Status.INTERVIEW, follow_up_on=self.TODAY)
        self.assertTrue(domain.needs_attention(due, stale_days=never, today=self.TODAY))
        later = Application(status=Status.INTERVIEW, follow_up_on=self.TODAY + dt.timedelta(days=1))
        self.assertFalse(domain.needs_attention(later, stale_days=14, today=self.TODAY))
        parked = Application(status=Status.TO_APPLY, follow_up_on=self.TODAY)
        self.assertFalse(domain.needs_attention(parked, stale_days=14, today=self.TODAY))


class ServiceRuleTests(SimpleTestCase):
    """The use cases on the in-memory adapter: no database, no files."""

    TODAY = dt.date(2026, 9, 2)

    def setUp(self):
        User = get_user_model()
        self.store = MemoryPersistence()
        self.files = MemoryStorageAdapter()
        self.user = User(pk=1, username="lionel")
        self.other = User(pk=2, username="marie")
        self.company = Company(pk=1, owner=self.user, name="Acme")

    def application(self, **fields):
        defaults = {
            "owner": self.user, "company": self.company, "title": "Poste",
            "status": Status.TO_APPLY, "score": 72,
        }
        defaults.update(fields)
        return self.store.applications.add(Application(**defaults))

    def change(self, application, status, **kwargs):
        kwargs.setdefault("follow_up_days", 10)
        return services.change_status(
            application, status, persistence=self.store, today=self.TODAY, **kwargs
        )

    def events(self, application):
        return [e for e in self.store.events.rows() if e.application_id == application.pk]

    def test_moving_to_sent_stamps_the_date_and_schedules_a_follow_up(self):
        application = self.application()
        self.assertTrue(self.change(application, Status.SENT, note="Via LinkedIn"))
        self.assertEqual(application.applied_on, self.TODAY)
        self.assertEqual(application.follow_up_on, self.TODAY + dt.timedelta(days=10))
        (event,) = self.events(application)
        self.assertEqual(event.kind, EventKind.APPLIED)
        self.assertEqual(event.title, "À postuler → Candidature envoyée")
        self.assertEqual(event.detail, "Via LinkedIn")
        self.assertEqual(event.happened_on, self.TODAY)

    def test_the_follow_up_delay_is_the_caller_s(self):
        application = self.application()
        self.change(application, Status.SENT, follow_up_days=3)
        self.assertEqual(application.follow_up_on, self.TODAY + dt.timedelta(days=3))

    def test_closing_clears_the_follow_up_and_stamps_the_close_date(self):
        application = self.application()
        self.change(application, Status.SENT)
        self.change(application, Status.REJECTED)
        self.assertIsNone(application.follow_up_on)
        self.assertEqual(application.closed_on, self.TODAY)
        self.assertEqual([e.kind for e in self.events(application)],
                         [EventKind.APPLIED, EventKind.REJECTION])

    def test_reopening_a_closed_application_clears_the_close_date(self):
        application = self.application()
        self.change(application, Status.REJECTED)
        self.change(application, Status.SENT)
        self.assertIsNone(application.closed_on)
        self.assertEqual(application.applied_on, self.TODAY)

    def test_a_no_op_or_unknown_transition_logs_nothing(self):
        application = self.application()
        self.assertFalse(self.change(application, Status.TO_APPLY))
        self.assertFalse(self.change(application, "nope"))
        self.assertEqual(self.events(application), [])

    def test_an_existing_applied_date_is_preserved(self):
        earlier = self.TODAY - dt.timedelta(days=12)
        application = self.application(applied_on=earlier)
        self.change(application, Status.SENT)
        self.assertEqual(application.applied_on, earlier)

    def test_advance_walks_the_pipeline_and_stops_at_the_end(self):
        application = self.application(status=Status.BACKLOG)
        walked = []
        while (
            nxt := services.advance(
                application, follow_up_days=10, persistence=self.store, today=self.TODAY
            )
        ) is not None:
            walked.append(nxt)
        self.assertEqual(
            walked,
            [Status.TO_APPLY, Status.SENT, Status.SCREENING, Status.INTERVIEW,
             Status.TECHNICAL, Status.OFFER, Status.ACCEPTED],
        )
        self.assertEqual(application.status, Status.ACCEPTED)
        self.assertEqual(application.applied_on, self.TODAY)
        self.assertEqual(application.closed_on, self.TODAY)
        self.assertIsNone(application.follow_up_on)
        self.assertEqual(len(self.events(application)), 7)

    def test_schedule_and_mark_follow_up(self):
        application = self.application(status=Status.SENT, applied_on=self.TODAY)
        services.schedule_follow_up(application, dt.date(2026, 9, 20), persistence=self.store)
        self.assertEqual(application.follow_up_on, dt.date(2026, 9, 20))

        next_on = services.mark_followed_up(
            application, follow_up_days=4, persistence=self.store, today=self.TODAY
        )
        self.assertEqual(next_on, self.TODAY + dt.timedelta(days=4))
        self.assertEqual(application.follow_up_on, next_on)
        (event,) = self.events(application)
        self.assertEqual((event.kind, event.title, event.happened_on),
                         (EventKind.FOLLOW_UP, "Relance envoyée", self.TODAY))

        services.schedule_follow_up(application, None, persistence=self.store)
        self.assertIsNone(application.follow_up_on)

    def test_record_and_add_event(self):
        application = self.application()
        created = services.record_application(application, persistence=self.store)
        self.assertEqual((created.kind, created.title), (EventKind.NOTE, "Candidature créée"))
        typed = ActivityEvent(
            kind=EventKind.CALL, title="Appel", detail="20 min", happened_on=dt.date(2026, 9, 1)
        )
        stored = services.add_event(application, typed, persistence=self.store)
        self.assertEqual(stored.application_id, application.pk)
        self.assertEqual((stored.kind, stored.title, stored.detail, stored.happened_on),
                         (EventKind.CALL, "Appel", "20 min", dt.date(2026, 9, 1)))
        self.assertEqual(len(self.events(application)), 2)

    def test_attach_document_demotes_the_previous_primary_and_logs(self):
        application = self.application()

        def upload(name):
            return services.attach_document(
                self.user, application, Document(kind=DocumentKind.CV, is_primary=True),
                label=name, upload=SimpleUploadedFile(name, b"x"),
                persistence=self.store, storage=self.files,
            )

        first = upload("CV1.docx")
        second = upload("CV2.docx")
        self.assertFalse(first.is_primary)
        self.assertTrue(second.is_primary)
        self.assertEqual(first.owner_id, self.user.pk)
        self.assertIs(first.application, application)
        self.assertEqual(first.file.name, f"documents/1/{application.slug}/CV1.docx")
        self.assertTrue(self.files.file_exists(first.file.name or ""))
        self.assertEqual(first.size_bytes, 1)
        self.assertEqual(
            [e.title for e in self.events(application)],
            ["Document ajouté : CV1.docx", "Document ajouté : CV2.docx"],
        )

        library = services.attach_document(
            self.user, None, Document(kind=DocumentKind.CV),
            label="Base", upload=SimpleUploadedFile("base.docx", b"x"),
            persistence=self.store, storage=self.files,
        )
        self.assertIsNone(library.application)
        self.assertEqual(library.owner_id, self.user.pk)
        self.assertEqual(library.file.name, "documents/1/bibliotheque/base.docx")
        self.assertEqual(len(self.events(application)), 2)
        self.assertEqual(self.store.documents.count(self.user), 3)

    def test_a_row_that_cannot_be_written_leaves_no_file_behind(self):
        application = self.application()
        with mock.patch.object(self.store.documents, "add", side_effect=RuntimeError("panne")):
            with self.assertRaises(RuntimeError):
                services.attach_document(
                    self.user, application, Document(kind=DocumentKind.CV), label="CV",
                    upload=SimpleUploadedFile("cv.docx", b"x"),
                    persistence=self.store, storage=self.files,
                )
        self.assertFalse(self.files.file_exists(f"documents/1/{application.slug}/cv.docx"))
        self.assertEqual(self.store.documents.count(self.user), 0)

    def test_deletions_return_the_application_and_refuse_a_stranger(self):
        application = self.application()
        event = self.store.events.add(application, EventKind.NOTE, "x")
        contact = self.store.contacts.add(Contact(application=application, name="X"))
        document = services.attach_document(
            self.user, application, Document(kind=DocumentKind.CV),
            label="CV", upload=SimpleUploadedFile("cv.docx", b"x"),
            persistence=self.store, storage=self.files,
        )
        for remove, pk in (
            (services.delete_event, event.pk),
            (services.delete_contact, contact.pk),
            (services.delete_document, document.pk),
        ):
            with self.subTest(remove=remove.__name__):
                with self.assertRaises(NotFound):
                    remove(self.other, pk, persistence=self.store)
        self.assertTrue(self.files.file_exists(document.file.name or ""))
        self.assertIs(services.delete_event(self.user, event.pk, persistence=self.store), application)
        self.assertIs(
            services.delete_contact(self.user, contact.pk, persistence=self.store), application
        )
        removed = services.delete_document(
            self.user, document.pk, persistence=self.store, storage=self.files
        )
        self.assertIs(removed.application, application)
        self.assertEqual(removed.label, "CV")
        self.assertEqual(self.store.documents.count(self.user), 0)
        self.assertFalse(self.files.file_exists(document.file.name or ""))
        self.assertIs(
            services.delete_application(self.user, application.pk, persistence=self.store),
            application,
        )
        with self.assertRaises(NotFound):
            self.store.applications.get(self.user, application.pk)

    def test_set_gap_status(self):
        gap = SkillGap(owner=self.user, name="Terraform")
        self.store.skill_gaps.save(gap)
        self.assertEqual(
            services.set_gap_status(self.user, gap.pk, GapStatus.DONE, persistence=self.store).status,
            GapStatus.DONE,
        )
        with self.assertRaises(ValueError):
            services.set_gap_status(self.user, gap.pk, "bogus", persistence=self.store)
        with self.assertRaises(NotFound):
            services.set_gap_status(self.other, gap.pk, GapStatus.DONE, persistence=self.store)

    def populate(self):
        self.application(score=80)
        self.application(status=Status.BACKLOG, score=None)
        self.application(status=Status.SENT, score=60, applied_on=self.TODAY - dt.timedelta(days=20))
        self.application(
            status=Status.INTERVIEW, score=90,
            applied_on=self.TODAY - dt.timedelta(days=5), follow_up_on=self.TODAY,
        )
        self.application(status=Status.REJECTED, score=30, closed_on=self.TODAY)
        self.application(status=Status.DISCARDED, score=99)
        self.store.applications.add(
            Application(owner=self.other, company=Company(pk=2, owner=self.other, name="Autre"),
                        title="Pas à moi", status=Status.TO_APPLY, score=100)
        )

    def test_dashboard_stats(self):
        self.populate()
        stats = services.dashboard_stats(self.user, persistence=self.store, today=self.TODAY)
        self.assertEqual(
            stats,
            {
                "tracked": 5, "to_apply": 1, "backlog": 1, "in_flight": 2, "interviewing": 1,
                "offers": 0, "rejected": 1, "discarded": 1, "sent_total": 3, "answered": 2,
                "response_rate": 67, "average_score": 77, "today": self.TODAY,
            },
        )
        empty = services.dashboard_stats(self.other, persistence=MemoryPersistence())
        self.assertEqual((empty["tracked"], empty["response_rate"], empty["average_score"]),
                         (0, None, None))

    def test_board(self):
        to_apply = self.application(score=80)
        backlog = self.application(status=Status.BACKLOG, score=None)
        sent = self.application(status=Status.SENT, score=60)
        self.application(status=Status.REJECTED)
        board = services.board(self.user, persistence=self.store)
        self.assertEqual([column["status"] for column in board], PIPELINE_STATUSES)
        columns = {column["status"]: column for column in board}
        self.assertEqual(columns[Status.TO_APPLY]["applications"], [to_apply])
        self.assertEqual(columns[Status.BACKLOG]["applications"], [backlog])
        self.assertEqual(columns[Status.SENT]["applications"], [sent])
        self.assertEqual(columns[Status.OFFER]["applications"], [])
        self.assertEqual((columns[Status.TO_APPLY]["count"], columns[Status.OFFER]["count"]), (1, 0))
        self.assertEqual((columns[Status.TO_APPLY]["label"], columns[Status.TO_APPLY]["tone"]),
                         ("À postuler", "amber"))

    def test_nav_counters(self):
        self.populate()
        services.attach_document(
            self.user, None, Document(kind=DocumentKind.CV), label="Base",
            upload=SimpleUploadedFile("b.docx", b"x"), persistence=self.store, storage=self.files,
        )
        self.assertEqual(
            services.nav_counters(self.user, stale_days=14, today=self.TODAY, persistence=self.store),
            {"open": 4, "tracked": 5, "attention": 2, "documents": 1},
        )
        self.assertEqual(
            services.nav_counters(self.user, stale_days=30, today=self.TODAY, persistence=self.store)["attention"],
            1,
        )

    def test_attention(self):
        self.populate()
        rows = services.attention(self.user, stale_days=14, today=self.TODAY, persistence=self.store)
        self.assertEqual([row.status for row in rows], [Status.INTERVIEW, Status.SENT])


@override_settings(MEDIA_ROOT=MEDIA)
class TransactionTests(TestCase):
    def test_a_failed_timeline_entry_rolls_the_status_change_back(self):
        application = make_application()
        store = DjangoPersistence()
        with mock.patch.object(store.events, "add", side_effect=RuntimeError("panne")):
            with self.assertRaises(RuntimeError):
                services.change_status(
                    application, Status.SENT, follow_up_days=10, persistence=store
                )
        application.refresh_from_db()
        self.assertEqual(application.status, Status.TO_APPLY)
        self.assertIsNone(application.applied_on)
        self.assertIsNone(application.follow_up_on)
        self.assertEqual(application.events.count(), 0)


class PersistenceResolverTests(SimpleTestCase):
    def test_the_setting_names_the_adapter(self):
        self.assertIsInstance(persistence(), DjangoPersistence)
        with override_settings(PERSISTENCE_ADAPTER="tracker.adapters.memory.MemoryPersistence"):
            self.assertIsInstance(persistence(), MemoryPersistence)
            self.assertIs(persistence(), persistence())
        self.assertIsInstance(persistence(), DjangoPersistence)


@override_settings(MEDIA_ROOT=MEDIA)
class NullOrderingTests(TestCase):
    """Nullable sort keys go last on every engine — the rule that keeps
    SQLite and PostgreSQL agreeing on what the pages show."""

    def setUp(self):
        self.user = make_user("Lionel")
        today = timezone.localdate()

        def application(name, **fields):
            fields.setdefault("distance_km", None)
            return make_application(owner=self.user, company_name=name, **fields)

        self.alpha = application("Alpha", score=90)
        self.beta = application("Beta", score=None, distance_km=10)
        self.gamma = application("Gamma", score=90, distance_km=5)
        self.delta = application(
            "Delta", status=Status.SENT, score=70,
            applied_on=today - dt.timedelta(days=5), follow_up_on=today + dt.timedelta(days=3),
        )
        self.epsilon = application("Epsilon", status=Status.SENT, score=60)
        self.zeta = application("Zeta", status=Status.REJECTED, score=50, closed_on=today)
        self.eta = application("Eta", status=Status.REJECTED, score=40)

    def test_the_table_sorts(self):
        rows = queries.filtered_applications(self.user, {"sort": "-score"})
        self.assertEqual(rows, [self.alpha, self.gamma, self.delta, self.epsilon, self.beta])

        rows = queries.filtered_applications(self.user, {"sort": "pipeline"})
        self.assertEqual(rows, [self.alpha, self.gamma, self.beta, self.delta, self.epsilon])

        rows = queries.filtered_applications(self.user, {"sort": "-applied_on"})
        self.assertEqual(rows[0], self.delta)
        self.assertEqual(rows[-1], self.beta)

        rows = queries.filtered_applications(self.user, {"sort": "follow_up_on"})
        self.assertEqual(rows[0], self.delta)
        self.assertEqual(rows[-1], self.beta)

    def test_the_dashboard_lists(self):
        lists = queries.dashboard_lists(self.user, today=timezone.localdate())
        self.assertEqual(lists["to_apply"], [self.gamma, self.alpha, self.beta])
        self.assertEqual(lists["in_flight"], [self.delta, self.epsilon])
        self.assertEqual(lists["backlog"], [])

    def test_closed_files(self):
        self.assertEqual(queries.closed(self.user), [self.zeta, self.eta])


@override_settings(MEDIA_ROOT=MEDIA)
class QueryBudgetTests(OwnedTestCase):
    """Upper bounds calibrated on today's counts, so an N+1 shows up.

    Transaction bookkeeping (savepoints, which a ``TestCase`` turns every
    ``atomic()`` into) is not a data query and is left out of the count.
    """

    BOOKKEEPING = ("SAVEPOINT", "RELEASE SAVEPOINT", "ROLLBACK TO SAVEPOINT")

    def setUp(self):
        super().setUp()
        self.application = make_application(company_name="Acme", status=Status.TO_APPLY)
        for index in range(3):
            make_document(
                application=self.application, kind=DocumentKind.CV, label=f"CV{index}.docx",
                language=Language.FR, is_primary=index == 0,
                file=SimpleUploadedFile(f"CV{index}.docx", b"x" * 10),
            )
        for name in ("Marie Dupont", "Jean Martin"):
            Contact.objects.create(application=self.application, name=name, role="Recruteur")
        today = timezone.localdate()
        for offset, kind in enumerate(
            [EventKind.NOTE, EventKind.CALL, EventKind.EMAIL, EventKind.INTERVIEW]
        ):
            ActivityEvent.objects.create(
                application=self.application, kind=kind, title=f"Événement {offset}",
                happened_on=today + dt.timedelta(days=offset - 2),
            )

    def count_queries(self, call):
        with CaptureQueriesContext(connection) as context:
            response = call()
        self.assertEqual(response.status_code, 200)
        return sum(
            1 for query in context.captured_queries
            if not query["sql"].startswith(self.BOOKKEEPING)
        )

    def test_application_detail(self):
        url = reverse("tracker:application_detail", args=[self.application.pk])
        self.assertLessEqual(self.count_queries(lambda: self.client.get(url)), 16)

    def test_dashboard(self):
        url = reverse("tracker:dashboard")
        self.assertLessEqual(self.count_queries(lambda: self.client.get(url)), 20)

    def test_set_status(self):
        url = reverse("tracker:set_status", args=[self.application.pk])
        post = lambda: self.client.post(  # noqa: E731
            url, {"status": Status.SENT, "source": "detail"}, HTTP_HX_REQUEST="true"
        )
        self.assertLessEqual(self.count_queries(post), 12)


class AnonymizeTests(SimpleTestCase):
    """The anonymiser: direct identifiers out, the rest of the CV intact."""

    def redact(self, text, **identity):
        return privacy.anonymize(text, known=privacy.known_identity(**identity))

    def test_contact_line(self):
        result = self.redact(
            "lionel.dupont@example.be • +32 470 12 34 56 • linkedin.com/in/lioneldupont • "
            "https://github.com/lionel"
        )
        self.assertEqual(result.text, "[EMAIL] • [TELEPHONE] • [URL] • [URL]")
        self.assertEqual(result.redactions, {"EMAIL": 1, "TELEPHONE": 1, "URL": 2})

    def test_phone_notations(self):
        for phone in ("0470 12 34 56", "0470/12.34.56", "0470123456", "+32 470 12 34 56",
                      "+32470123456", "+32 (0)2 123 45 67", "02 123 45 67", "0032 2 987 65 43",
                      "010 22 33 44", "+33 6 12 34 56 78"):
            with self.subTest(phone=phone):
                self.assertEqual(self.redact(f"Tél. {phone} (soir)").text, "Tél. [TELEPHONE] (soir)")

    def test_what_is_not_a_phone_survives(self):
        text = ("2019-2021 : Acme. 06/2019 - 08/2021 mission Odoo (version 16.0). Python 3.12. "
                "01/02/2020. 2010-2015 à l'ULB. Note 0 à 5. Budget 0470 €. 08:30. "
                "02.03.2021 - 02.05.2021. Terraform 1.5.7, 2 enfants, 10 ans d'expérience.")
        self.assertEqual(self.redact(text).text, text)

    def test_identifiers(self):
        result = self.redact(
            "IBAN BE68 5390 0754 7034 · NISS 85.07.30-033.28 · Né le 30/07/1985 à Nivelles · Âge : 41 ans"
        )
        self.assertEqual(
            result.text, "IBAN [IBAN] · NISS [NISS] · Né le [DATE-DE-NAISSANCE] · Âge : [AGE]"
        )

    def test_birth_date_and_place_notations(self):
        cases = {
            "Date de naissance : 30 juillet 1985": "Date de naissance : [DATE-DE-NAISSANCE]",
            "Née le 1er mars 1990 à La Louvière.": "Née le [DATE-DE-NAISSANCE].",
            "Born on 1985-07-30 in Mons": "Born on [DATE-DE-NAISSANCE]",
            "Né à Braine-l'Alleud, 2 enfants": "Né à [LIEU], 2 enfants",
            "Born in Mons, 1985-07-30.": "Born in [LIEU].",
            "geboren op 30.07.1985 te Gent": "geboren op [DATE-DE-NAISSANCE]",
        }
        for text, expected in cases.items():
            with self.subTest(text=text):
                self.assertEqual(self.redact(text).text, expected)

    def test_addresses(self):
        cases = {
            "Chaussée de Namur 12 bte 3, 1400 Nivelles": "[ADRESSE], [ADRESSE]",
            "Rue de la Station 3": "[ADRESSE]",
            "Avenue Louise 54 – 1050 Bruxelles – Belgique": "[ADRESSE] – [ADRESSE] – Belgique",
            "Kerkstraat 12, 2000 Antwerpen": "[ADRESSE], [ADRESSE]",
            # Right after a street, « · 2018 Python » reads as a postcode and
            # a town: the separator a two-column contact table produces is
            # exactly this one, so the address wins over the false positive.
            "Grote Markt 5 · 2018 Python": "[ADRESSE] · [ADRESSE]",
            "Rue Neuve 12 · 1400 Nivelles · Belgique": "[ADRESSE] · [ADRESSE] · Belgique",
            "Chemin de la Cure 12A": "[ADRESSE]",
            "Adresse : B-1000 Bruxelles": "Adresse : [ADRESSE]",
            "Formation à l'ULB, 1050 Bruxelles, Belgique.": "Formation à l'ULB, [ADRESSE], Belgique.",
            "Rue Neuve 12\n1400 Nivelles\n2019 Python": "[ADRESSE]\n[ADRESSE]\n2019 Python",
        }
        for text, expected in cases.items():
            with self.subTest(text=text):
                self.assertEqual(self.redact(text).text, expected)

    def test_what_is_not_an_address_survives(self):
        text = ("AWS (EC2, Route 53, IAM). Place 2 au hackathon 2019. Projet chemin de fer 2020. "
                "Bruxelles, 2019 Python. Marketplace 12 ventes. 2000 Antwerpen (sans rue) reste.")
        self.assertEqual(self.redact(text).text, text)

    def test_ordinary_french_sentences_are_not_addresses(self):
        """The keyword list is made of everyday nouns; only the shape of a
        real address — a capitalised name then a number that ends it —
        should trigger the rule."""
        for text in (
            "Cours de Java 2 à l'ULB",
            "chemin critique de la version 2 du projet",
            "Ma place dans une équipe de 4 personnes",
            "3 h par semaine sur ce chantier",
            "Boulevard des idées reçues sur 12 ans de carrière",
            "Sur la route depuis 3 ans",
            "Allée simple ou 2 allers-retours",
            "500 K€ de budget",
        ):
            with self.subTest(text=text):
                self.assertEqual(self.redact(text).text, text)

    def test_notations_a_french_cv_template_produces(self):
        cases = {
            "Né(e) le 30/07/1985": "Né(e) le [DATE-DE-NAISSANCE]",
            "Naissance : 30 juillet 1985": "Naissance : [DATE-DE-NAISSANCE]",
            "Né(e) à Nivelles": "Né(e) à [LIEU]",
            "(+32) 470 12 34 56": "[TELEPHONE]",
            "+32 (0)2 123 45 67": "[TELEPHONE]",
        }
        for text, expected in cases.items():
            with self.subTest(text=text):
                self.assertEqual(self.redact(text).text, expected)

    def test_the_profile_s_identity_is_masked_whole_word_case_and_accent_insensitive(self):
        result = self.redact(
            "LIONEL DUPONT — Lionel Dupont, dit Lionél. Habite Nivelles (NIVELLES). "
            "lionel.dupont@example.be. Le mot « Lionelle » et « dupontel » restent.",
            display_name="Lionel Dupont", username="lionel", email="lionel.dupont@example.be",
            location="Nivelles",
        )
        self.assertEqual(
            result.text,
            "[NOM] — [NOM], dit [NOM]. Habite [LIEU] ([LIEU]). [EMAIL]. "
            "Le mot « Lionelle » et « dupontel » restent.",
        )
        self.assertEqual(result.redactions["NOM"], 3)
        self.assertEqual(result.redactions["LIEU"], 2)

    def test_known_identity_skips_short_tokens_and_plain_usernames(self):
        known = privacy.known_identity(
            display_name="Jo De Smet", username="local", email="jo@x.be", location=""
        )
        self.assertEqual(known, {"Jo De Smet": "NOM", "Smet": "NOM", "jo@x.be": "EMAIL"})
        self.assertEqual(self.redact("Le marché local de Jo", display_name="Jo", username="local").text,
                         "Le marché local de Jo")

    def test_a_known_phone_backs_up_the_pattern(self):
        """Le motif couvre les formes usuelles ; le numéro que le profil
        détient rattrape celle qu'il manque — ici l'indicatif américain avec
        son indicatif régional entre parenthèses."""
        odd = "+1 (415) 555-0134"
        self.assertIn(odd, self.redact(f"Tel {odd}").text)
        self.assertEqual(self.redact(f"Tel {odd}", phone=odd).text, "Tel [TELEPHONE]")
        # Un profil sans téléphone n'ajoute rien au dictionnaire.
        self.assertEqual(privacy.known_identity(phone=""), {})
        self.assertEqual(
            privacy.known_identity(phone="+32 470 12 34 56"), {"+32 470 12 34 56": "TELEPHONE"}
        )

    def test_summary_in_french(self):
        self.assertEqual(privacy.anonymize("rien").summary(), "rien à masquer")
        self.assertEqual(self.redact("a@b.be").summary(), "1 e-mail masqué")
        self.assertEqual(
            self.redact("a@b.be c@d.be 0470 12 34 56").summary(),
            "2 e-mails, 1 numéro de téléphone masqués",
        )

    def test_empty_text(self):
        result = privacy.anonymize("")
        self.assertEqual((result.text, result.redactions, result.total), ("", {}, 0))


class DocumentTextTests(SimpleTestCase):
    """The extractor behind ``extract_and_anonymize_text``."""

    def setUp(self):
        # pypdf logs a warning for the deliberately corrupt file below.
        pypdf_logger = logging.getLogger("pypdf")
        pypdf_logger.disabled = True
        self.addCleanup(setattr, pypdf_logger, "disabled", False)

    docx = staticmethod(make_docx)
    pdf = staticmethod(make_pdf)

    def test_docx_paragraphs_then_tables(self):
        data = self.docx("Lionel Dupont", "", "Ingénieur DevOps", table=("Python", "Django"))
        self.assertEqual(
            document_text.extract_text("cv.docx", data), "Lionel Dupont\nIngénieur DevOps\nPython · Django"
        )

    def test_docx_headers_footers_and_nested_tables_are_read(self):
        """Word and Canva templates park the contact block in a header and a
        column in a text box; missing them makes a real CV look like a scan."""
        import docx

        document = docx.Document()
        document.sections[0].header.paragraphs[0].text = "Lionel Dupont · lionel@example.be"
        document.sections[0].footer.paragraphs[0].text = "Chaussée de Namur 12, 1400 Nivelles"
        document.add_paragraph("Ingénieur DevOps senior")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Python"
        table.rows[0].cells[1].text = "Django"
        table.rows[0].cells[0].add_table(rows=1, cols=1).rows[0].cells[0].text = "Terraform"
        buffer = io.BytesIO()
        document.save(buffer)

        text = document_text.extract_text("cv.docx", buffer.getvalue())
        for expected in ("lionel@example.be", "1400 Nivelles", "Ingénieur DevOps senior",
                         "Python", "Django", "Terraform"):
            with self.subTest(expected=expected):
                self.assertIn(expected, text)
        self.assertEqual(text.count("Lionel Dupont"), 1)

    def test_text_and_markdown_are_read_as_utf8(self):
        self.assertEqual(document_text.extract_text("cv.txt", "Élodie\n".encode()), "Élodie")
        self.assertEqual(document_text.extract_text("cv.MD", b"# CV\n"), "# CV")
        self.assertEqual(document_text.extract_text("cv.txt", b"\xff\xfe"), "\ufffd\ufffd")

    def test_a_pdf_without_a_text_layer_is_empty_not_an_error(self):
        self.assertEqual(document_text.extract_text("cv.pdf", self.pdf()), "")

    def test_unsupported_and_corrupt_files(self):
        for name, data in (("cv.odt", b"x"), ("cv", b"x"), ("cv.doc", b"x"),
                           ("cv.pdf", b"not a pdf"), ("cv.docx", b"not a zip")):
            with self.subTest(name=name), self.assertRaises(document_text.UnsupportedFormat):
                document_text.extract_text(name, data)
        self.assertTrue(document_text.is_supported("x/y/CV.PDF"))
        self.assertFalse(document_text.is_supported("x/y/CV.odt"))


@override_settings(MEDIA_ROOT=MEDIA)
class StoragePortContractTests(SimpleTestCase):
    """Every provider honours ``StoragePort`` the same way: disk, memory, Azure."""

    def adapters(self):
        azure, _ = azure_adapter()
        return [("local", LocalStorageAdapter()), ("memory", MemoryStorageAdapter()), ("azure", azure)]

    def each_adapter(self, check):
        for name, files in self.adapters():
            with self.subTest(provider=name):
                self.assertIsInstance(files, StoragePort)
                check(files)

    def test_round_trip_with_sanitised_and_deduplicated_names(self):
        def check(files):
            name = files.save_file("documents/1/acme-devops/CV base é.docx", b"octets")
            self.assertEqual(name, "documents/1/acme-devops/CV_base_é.docx")
            self.assertTrue(files.file_exists(name))
            with files.open_file(name) as handle:
                self.assertEqual(handle.read(), b"octets")
            again = files.save_file("documents/1/acme-devops/CV base é.docx", io.BytesIO(b"bis"))
            self.assertRegex(again, r"^documents/1/acme-devops/CV_base_é_\w{7}\.docx$")
            self.assertEqual(files.open_file(again).read(), b"bis")
            files.delete_file(name)
            files.delete_file(name)  # already gone: not an error
            self.assertFalse(files.file_exists(name))
            with self.assertRaises(MissingFile):
                files.open_file(name)
            with self.assertRaises(SuspiciousFileOperation):
                files.save_file("documents/1/../2/cv.docx", b"x")

        self.each_adapter(check)

    def test_listing_sizing_and_dating_what_is_stored(self):
        def check(files):
            # A prefix of its own: the local adapter shares MEDIA_ROOT with
            # the rest of the class.
            names = {
                files.save_file(f"documents/70{n}/bibliotheque/cv.txt", b"octets") for n in (1, 2)
            }
            files.save_file("ailleurs/note.txt", b"x")
            files.save_file("documents-anciens/note.txt", b"x")
            under = set(files.list_files("documents"))
            self.assertTrue(names <= under)
            # A folder, not a bare string prefix.
            self.assertNotIn("documents-anciens/note.txt", under)
            self.assertNotIn("ailleurs/note.txt", under)
            self.assertTrue(names <= set(files.list_files()))
            self.assertEqual(set(files.list_files("documents/701")), {sorted(names)[0]})
            one = sorted(names)[0]
            self.assertEqual(files.file_size(one), 6)
            self.assertLess(
                (dt.datetime.now(dt.timezone.utc) - files.file_modified_at(one)).total_seconds(), 120
            )
            self.assertEqual(list(files.list_files("documents/999")), [])

        self.each_adapter(check)

    def test_extract_and_anonymize_text(self):
        known = privacy.known_identity(display_name="Lionel Dupont", email="lionel@x.be")

        def check(files):
            name = files.save_file(
                "documents/1/bibliotheque/cv.txt", "Lionel Dupont · lionel@x.be · 0470 12 34 56".encode()
            )
            plain = files.extract_and_anonymize_text(name)
            self.assertEqual(plain.text, "Lionel Dupont · [EMAIL] · [TELEPHONE]")
            self.assertEqual(files.extract_and_anonymize_text(name, known=known).text,
                             "[NOM] · [EMAIL] · [TELEPHONE]")
            docx = files.save_file("documents/1/bibliotheque/cv.docx", make_docx("Marie", "0470 12 34 56"))
            self.assertEqual(files.extract_and_anonymize_text(docx).text, "Marie\n[TELEPHONE]")
            odt = files.save_file("documents/1/bibliotheque/cv.odt", b"x")
            with self.assertRaises(document_text.UnsupportedFormat):
                files.extract_and_anonymize_text(odt)
            with self.assertRaises(MissingFile):
                files.extract_and_anonymize_text("documents/1/bibliotheque/absent.txt")

        self.each_adapter(check)

    def test_links_of_the_providers_the_application_serves(self):
        for name, files in self.adapters()[:2]:
            with self.subTest(provider=name):
                self.assertTrue(files.served_by_app)
                url = files.get_secure_url("documents/7/bibliotheque/cv.pdf")
                self.assertTrue(url.startswith("/fichiers/"))
                token = url[len("/fichiers/"):].rstrip("/")
                self.assertEqual(links.read_link(token), ("documents/7/bibliotheque/cv.pdf", 7))
                # The Django face of the same object hands out the same link.
                self.assertTrue(files.url("documents/7/bibliotheque/cv.pdf").startswith("/fichiers/"))
                with self.assertRaises(ValueError):
                    files.url("")
                # Outside the layout: a link nobody can open (see the view).
                token = files.get_secure_url("ailleurs/cv.pdf")[len("/fichiers/"):].rstrip("/")
                self.assertEqual(links.read_link(token), ("ailleurs/cv.pdf", None))

    def test_the_django_face_and_the_port_hand_out_the_same_link(self):
        def check(files):
            name = files.save_file("documents/7/bibliotheque/cv.pdf", b"x")
            self.assertEqual(files.url(name), files.get_secure_url(name))
            with self.assertRaises(ValueError):
                files.url("")
            with self.assertRaises(ValueError):
                files.url(None)

        self.each_adapter(check)

    def test_a_stored_name_never_outgrows_the_column(self):
        def check(files):
            long_name = "documents/1/" + "s" * 200 + "/" + "n" * 255 + ".docx"
            for _ in range(3):
                name = files.save_file(long_name, b"x", max_length=DOCUMENT_NAME_MAX_LENGTH)
                self.assertLessEqual(len(name), DOCUMENT_NAME_MAX_LENGTH)
                self.assertTrue(files.file_exists(name))

        self.each_adapter(check)

    def test_a_handle_can_be_read_again(self):
        def check(files):
            name = files.save_file("documents/1/bibliotheque/cv.txt", b"abcdef")
            handle = files.open_file(name)
            self.assertEqual(handle.read(2), b"ab")
            handle.open("rb")
            self.assertEqual(handle.read(), b"abcdef")
            handle.close()
            self.assertEqual(handle.open("rb").read(), b"abcdef")

        self.each_adapter(check)

    def test_a_name_no_storage_would_accept_reads_as_a_missing_file(self):
        def check(files):
            for name in ("../../etc/passwd", "documents/1/../../etc/passwd", "/etc/passwd"):
                with self.subTest(name=name):
                    with self.assertRaises(MissingFile):
                        files.open_file(name)
                    self.assertFalse(files.file_exists(name))
                    files.delete_file(name)  # not an error either

        self.each_adapter(check)

    def test_link_tokens_expire_and_refuse_tampering(self):
        url = MemoryStorageAdapter().get_secure_url("documents/7/bibliotheque/cv.pdf")
        token = url[len("/fichiers/"):].rstrip("/")
        with mock.patch("django.core.signing.time.time", return_value=time.time() + 61):
            self.assertEqual(links.read_link(token, max_age=120)[1], 7)
            with self.assertRaises(links.BadLink):
                links.read_link(token, max_age=60)
        for bad in (token[:-2] + "zz", "n'importe quoi", signing.dumps(["pas", "un", "dict"], salt=links.SALT)):
            with self.subTest(token=bad[:12]), self.assertRaises(links.BadLink):
                links.read_link(bad)
        self.assertEqual(links.link_ttl(), 300)
        with override_settings(STORAGES=MEMORY_STORAGES):
            self.assertEqual(links.link_ttl(), 60)


class AzureAdapterTests(SimpleTestCase):
    """The Azure adapter against a fake account: uploads, downloads, SAS links."""

    NAME = "documents/1/acme-devops/CV base.pdf"

    def test_sas_link(self):
        files, service = azure_adapter(link_ttl=300)
        self.assertFalse(files.served_by_app)
        before = dt.datetime.now(dt.timezone.utc).replace(microsecond=0)
        url = files.get_secure_url(self.NAME)
        parts = urlsplit(url)
        self.assertEqual((parts.scheme, parts.netloc), ("https", "jobhunt.blob.core.windows.net"))
        self.assertEqual(parts.path, "/documents/documents/1/acme-devops/CV%20base.pdf")
        query = {key: values[0] for key, values in parse_qs(parts.query).items()}
        self.assertEqual((query["sp"], query["sr"]), ("r", "b"))
        self.assertIn("sig", query)
        start = dt.datetime.strptime(query["st"], "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=dt.timezone.utc)
        expiry = dt.datetime.strptime(query["se"], "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=dt.timezone.utc)
        self.assertEqual(expiry - start, dt.timedelta(minutes=5, seconds=300))
        self.assertLessEqual(abs((expiry - before).total_seconds() - 300), 2)
        self.assertEqual(query["rscd"], 'attachment; filename="CV base.pdf"')
        self.assertEqual(urlsplit(files.url(self.NAME)).path, parts.path)

    def test_the_sas_window_sits_inside_the_delegation_key_and_forbids_http(self):
        service = FakeService(token=True)
        files = AzureStorageAdapter(client=service, link_ttl=300)
        url = files.get_secure_url(self.NAME)
        self.assertEqual(url.count("?"), 1)
        query = {key: values[0] for key, values in parse_qs(urlsplit(url).query).items()}
        self.assertEqual(query["spr"], "https")
        def moment(value):
            return dt.datetime.strptime(value, "%Y-%m-%dT%H:%M:%SZ")

        self.assertLessEqual(moment(query["skt"]), moment(query["st"]))
        self.assertLessEqual(moment(query["se"]), moment(query["ske"]))

    def test_a_client_whose_url_already_carries_a_query_keeps_one_question_mark(self):
        files, service = azure_adapter()

        class Queried(FakeBlob):
            @property
            def url(self):
                return super().url + "?sv=2024-05-04&sig=déjà"

        with mock.patch.object(FakeContainer, "get_blob_client",
                               lambda container, name: Queried(container, name)):
            url = files.get_secure_url(self.NAME)
        self.assertEqual(url.count("?"), 1)
        self.assertNotIn("déjà", url)

    def test_a_missing_container_is_a_provider_failure_not_a_missing_file(self):
        files, service = azure_adapter()
        service.container.missing = True
        with self.assertRaises(StorageError):
            files.check_container()
        service.container.missing = False
        files.check_container()  # present: nothing raised

    def test_a_stream_that_fails_mid_download_is_a_storage_error(self):
        files, service = azure_adapter()
        service.container.blobs[self.NAME] = b"abcdef"
        handle = files.open_file(self.NAME)
        from azure.core.exceptions import IncompleteReadError

        with mock.patch.object(FakeDownloader, "read", side_effect=IncompleteReadError("coupé")):
            with self.assertRaises(StorageError):
                handle.read()

    def test_describe_names_the_account_and_how_links_are_signed(self):
        files, _ = azure_adapter()
        self.assertEqual(
            files.describe(),
            {"Compte": "jobhunt", "Conteneur": "documents", "Signature": "clé de compte"},
        )
        identity = AzureStorageAdapter(client=FakeService(token=True))
        self.assertEqual(identity.describe()["Signature"], "identité (clé de délégation)")

    def test_a_french_basename_gets_an_rfc_6266_disposition(self):
        files, _ = azure_adapter()
        query = parse_qs(urlsplit(files.get_secure_url("documents/1/b/CV Élodie.pdf")).query)
        self.assertEqual(query["rscd"][0], "attachment; filename*=utf-8''CV%20%C3%89lodie.pdf")

    def test_upload_rewinds_the_stream_and_declares_length_and_content_type(self):
        files, service = azure_adapter()
        upload = SimpleUploadedFile("cv.pdf", b"%PDF-1.4 contenu")
        upload.read(5)  # a reader left the position elsewhere
        name = files.save_file("documents/1/acme/cv.pdf", upload)
        self.assertEqual(name, "documents/1/acme/cv.pdf")
        self.assertEqual(service.container.uploads, [(name, 16, 16, "application/pdf")])
        self.assertEqual(service.container.blobs[name], b"%PDF-1.4 contenu")

    def test_a_name_taken_between_the_check_and_the_write_gets_another_one(self):
        files, service = azure_adapter()
        # Someone else wrote that exact blob after ``exists()`` said no.
        service.container.race_on_next_upload = True
        name = files.save_file("documents/1/acme/cv.pdf", b"x", max_length=DOCUMENT_NAME_MAX_LENGTH)
        self.assertRegex(name, r"^documents/1/acme/cv_\w{7}\.pdf$")
        self.assertEqual(service.container.blobs[name], b"x")

    def test_a_retried_name_still_fits_the_column(self):
        files, service = azure_adapter()
        service.container.race_on_next_upload = True
        long_name = "documents/1/" + "s" * 200 + "/" + "n" * 255 + ".pdf"
        name = files.save_file(long_name, b"x", max_length=DOCUMENT_NAME_MAX_LENGTH)
        self.assertLessEqual(len(name), DOCUMENT_NAME_MAX_LENGTH)
        self.assertEqual(service.container.blobs[name], b"x")

    def test_a_name_that_stays_taken_is_a_storage_error(self):
        files, service = azure_adapter()

        def always_taken(*args, **kwargs):
            from azure.core.exceptions import ResourceExistsError

            raise ResourceExistsError("existe déjà")

        service.container.upload_blob = always_taken
        with self.assertRaises(StorageError):
            files.save_file("documents/1/acme/cv2.pdf", b"x")

    def test_download_handle(self):
        files, service = azure_adapter()
        service.container.blobs[self.NAME] = b"abcdef"
        handle = files.open_file(self.NAME)
        self.assertEqual(handle.size, 6)
        self.assertEqual(handle.read(2), b"ab")
        self.assertEqual(handle.read(), b"cdef")
        self.assertFalse(handle.seekable())
        handle.close()
        handle.open("rb")
        self.assertEqual(handle.read(), b"abcdef")
        with self.assertRaises(ValueError):
            handle.open("wb")
        self.assertEqual(files.size(self.NAME), 6)
        files.delete_file(self.NAME)
        self.assertNotIn(self.NAME, service.container.blobs)
        with self.assertRaises(MissingFile):
            files.size(self.NAME)

    def test_sdk_failures_become_storage_errors(self):
        files, service = azure_adapter()
        service.container.broken = True
        for operation in (files.file_exists, files.open_file, files.delete_file, files.size):
            with self.subTest(operation=operation.__name__), self.assertRaises(StorageError):
                operation(self.NAME)

    def test_user_delegation_key_is_fetched_once_and_renewed_in_time(self):
        service = FakeService(token=True)
        files = AzureStorageAdapter(client=service, link_ttl=300)
        first = files.get_secure_url(self.NAME)
        files.get_secure_url(self.NAME)
        self.assertEqual(len(service.delegation_calls), 1)
        self.assertIn("skoid", parse_qs(urlsplit(first).query))
        start, expiry = service.delegation_calls[0]
        self.assertEqual(expiry - start, dt.timedelta(hours=2, minutes=15))
        later = dt.datetime.now(dt.timezone.utc) + dt.timedelta(hours=1, minutes=50)
        with mock.patch("tracker.adapters.azure_storage._now", return_value=later):
            files.get_secure_url(self.NAME)
        self.assertEqual(len(service.delegation_calls), 2)

    def test_a_client_that_cannot_sign_is_a_configuration_error(self):
        files = AzureStorageAdapter(client=FakeService())  # SAS-only credential
        with self.assertRaises(ImproperlyConfigured):
            files.get_secure_url(self.NAME)
        with self.assertRaises(ImproperlyConfigured):
            AzureStorageAdapter()

    def test_the_real_client_is_built_from_the_settings(self):
        from azure.storage.blob import BlobServiceClient

        files = AzureStorageAdapter(
            account_url="https://jobhunt.blob.core.windows.net", account_key=AZURE_KEY, container="cv"
        )
        self.assertIsInstance(files.client, BlobServiceClient)
        self.assertEqual(files.client.account_name, "jobhunt")
        self.assertEqual(files.container.container_name, "cv")
        url = files.get_secure_url("documents/1/b/cv.pdf")
        self.assertTrue(url.startswith("https://jobhunt.blob.core.windows.net/cv/documents/1/b/cv.pdf?"))
        # A connection string is parsed by jobhunt.storage, never handed over.
        url, key = parse_connection_string(
            "DefaultEndpointsProtocol=https;AccountName=jobhunt;"
            f"AccountKey={AZURE_KEY};EndpointSuffix=core.windows.net"
        )
        connection = AzureStorageAdapter(account_url=url, account_key=key)
        self.assertIn("sig=", connection.get_secure_url("documents/1/b/cv.pdf"))


@override_settings(MEDIA_ROOT=MEDIA)
class PrivateFileViewTests(OwnedTestCase):
    """``get_secure_url`` on the providers the application serves itself."""

    def setUp(self):
        super().setUp()
        self.document = make_document(
            kind=DocumentKind.CV, label="CV base",
            file=SimpleUploadedFile("cv-base.docx", b"contenu"),
        )
        self.name = self.document.file.name or ""
        self.link = storage().get_secure_url(self.name)

    def test_the_owner_follows_the_link(self):
        response = self.client.get(self.link)
        self.assertEqual(response.status_code, 200)
        # Files outlive the test database in MEDIA: the name may carry a suffix.
        self.assertRegex(response.headers["Content-Disposition"], r'^attachment; filename="cv-base\w*\.docx"$')
        self.assertEqual(response.headers["Content-Length"], "7")
        self.assertEqual(b"".join(getattr(response, "streaming_content")), b"contenu")
        self.assertEqual(self.document.file.url[:10], "/fichiers/")

    def test_altered_or_foreign_tokens_are_404(self):
        for url in (self.link[:-3] + "zz/", "/fichiers/nimporte-quoi/",
                    reverse("tracker:private_file", args=[signing.dumps({"n": "x.txt", "u": None}, salt=links.SALT)]),
                    reverse("tracker:private_file", args=[signing.dumps({"n": self.name}, salt="autre")])):
            with self.subTest(url=url[:30]):
                self.assertEqual(self.client.get(url).status_code, 404)

    def test_an_expired_link_is_404(self):
        with mock.patch("django.core.signing.time.time", return_value=time.time() + 301):
            self.assertEqual(self.client.get(self.link).status_code, 404)
        with mock.patch("django.core.signing.time.time", return_value=time.time() + 200):
            self.assertEqual(self.client.get(self.link).status_code, 200)

    def test_someone_else_cannot_use_the_owner_s_link(self):
        self.client.force_login(make_user("Marie"))
        self.assertEqual(self.client.get(self.link).status_code, 404)

    def test_a_link_is_checked_against_the_rows_not_only_its_signature(self):
        # A perfectly signed link to a file no document of this account
        # claims: the signature is not authority.
        orphan = storage().save_file(f"documents/{self.user.pk}/bibliotheque/orphelin.docx", b"x")
        self.assertEqual(self.client.get(links.make_link(orphan)).status_code, 404)
        # And the document's own link still works.
        self.assertEqual(self.client.get(self.link).status_code, 200)

    def test_a_file_stored_before_the_account_layout_is_served_to_its_owner(self):
        """Rows from before the per-account folders (``documents/<slug>/…``)
        carry no account in their name: the row decides, not the path."""
        legacy = storage().save_file("documents/bibliotheque/ancien.docx", b"ancien")
        Document.objects.filter(pk=self.document.pk).update(file=legacy)
        link = storage().get_secure_url(legacy)
        self.assertEqual(links.read_link(link[len("/fichiers/"):].rstrip("/"))[1], None)
        self.assertEqual(self.client.get(link).status_code, 200)
        self.client.force_login(make_user("Marie"))
        self.assertEqual(self.client.get(link).status_code, 404)

    @override_settings(AUTH_MODE="accounts")
    def test_a_visitor_is_sent_to_the_login_page(self):
        self.client.logout()
        response = self.client.get(self.link)
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.headers["Location"].startswith(reverse("accounts:login")))

    def test_a_missing_file_is_404_and_an_outage_503(self):
        with mock.patch.object(LocalStorageAdapter, "open_file", side_effect=StorageError("panne")):
            with self.assertLogs("tracker.views", "WARNING"):
                self.assertEqual(self.client.get(self.link).status_code, 503)
                self.assertEqual(self.client.get(self.document.get_download_url()).status_code, 503)
        storage().delete_file(self.name)
        self.assertEqual(self.client.get(self.link).status_code, 404)
        self.assertEqual(self.client.get(self.document.get_download_url()).status_code, 404)

    def test_the_memory_provider_serves_the_same_way(self):
        with override_settings(STORAGES=MEMORY_STORAGES):
            document = make_document(
                kind=DocumentKind.CV, label="Mémoire", file=SimpleUploadedFile("m.txt", "en mémoire".encode()),
            )
            response = self.client.get(storage().get_secure_url(document.file.name or ""))
            self.assertEqual(response.status_code, 200)
            self.assertEqual(b"".join(getattr(response, "streaming_content")), "en mémoire".encode())
            self.assertEqual(self.client.get(document.get_download_url()).status_code, 200)


def azurite_is_running() -> bool:
    """Whether the Azure emulator answers on its usual port (see the README)."""
    import socket

    with socket.socket() as probe:
        probe.settimeout(0.2)
        return probe.connect_ex(("127.0.0.1", 10000)) == 0


@skipUnless(azurite_is_running(), "Azurite n'écoute pas sur 127.0.0.1:10000")
class AzuriteTests(SimpleTestCase):
    """The Azure adapter against the real SDK and the official emulator.

    The fake client in ``AzureAdapterTests`` pins what the adapter *asks*;
    this pins what the service actually *does* — the truncation trap on a
    partly read upload, the permissions a SAS really carries, and the
    encoding of a French filename. Skipped unless Azurite is running:

        docker run -d --name jobhunt-azurite -p 127.0.0.1:10000:10000 \
          mcr.microsoft.com/azure-storage/azurite azurite-blob --blobHost 0.0.0.0
    """

    def setUp(self):
        from jobhunt.storage import parse_connection_string
        from tracker.adapters.azure_storage import AzureStorageAdapter

        # One container per test: deleting a container is asynchronous, so
        # sharing one makes the next test race the previous one's cleanup.
        name = re.sub(r"[^a-z0-9]+", "-", self._testMethodName.lower()).strip("-")
        self.CONTAINER = f"jh-{name}"[:63].rstrip("-")
        url, key = parse_connection_string("UseDevelopmentStorage=true")
        self.files = AzureStorageAdapter(account_url=url, account_key=key, container=self.CONTAINER)
        from azure.core.exceptions import ResourceExistsError

        try:
            self.files.container.create_container()
        except ResourceExistsError:
            pass
        self.addCleanup(self.files.container.delete_container)

    def test_an_upload_read_half_way_is_still_stored_whole(self):
        payload = b"%PDF-1.4 " + b"contenu " * 100
        upload = SimpleUploadedFile("CV Éléonore.pdf", payload)
        upload.read(17)  # the SDK reads from the current position and never rewinds
        name = self.files.save_file(
            "documents/1/acme/CV Éléonore.pdf", upload, max_length=DOCUMENT_NAME_MAX_LENGTH
        )
        self.assertEqual(self.files.size(name), len(payload))
        self.assertEqual(self.files.open_file(name).read(), payload)

    def test_a_secure_url_reads_and_only_reads(self):
        import requests

        payload = b"contenu"
        name = self.files.save_file("documents/1/acme/CV Éléonore.pdf", payload)
        response = requests.get(self.files.get_secure_url(name), timeout=10)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content, payload)
        # RFC 6266 for a name the header cannot carry as-is.
        self.assertEqual(
            response.headers["content-disposition"],
            "attachment; filename*=utf-8''CV_%C3%89l%C3%A9onore.pdf",
        )
        written = requests.put(
            self.files.get_secure_url(name),
            data=b"pirate",
            headers={"x-ms-blob-type": "BlockBlob"},
            timeout=10,
        )
        self.assertEqual(written.status_code, 403)
        self.assertEqual(self.files.open_file(name).read(), payload)

    def test_collisions_deletions_and_what_is_not_there(self):
        first = self.files.save_file("documents/1/acme/cv.pdf", b"un")
        second = self.files.save_file("documents/1/acme/cv.pdf", b"deux")
        self.assertNotEqual(first, second)
        self.assertEqual(self.files.open_file(first).read(), b"un")
        self.assertEqual(self.files.open_file(second).read(), b"deux")
        with self.assertRaises(MissingFile):
            self.files.open_file("documents/1/acme/absent.pdf")
        self.files.delete_file(first)
        self.files.delete_file(first)  # idempotent
        self.assertFalse(self.files.file_exists(first))

    def test_a_container_that_is_not_there_is_a_provider_failure(self):
        from jobhunt.storage import parse_connection_string
        from tracker.adapters.azure_storage import AzureStorageAdapter

        url, key = parse_connection_string("UseDevelopmentStorage=true")
        elsewhere = AzureStorageAdapter(account_url=url, account_key=key, container="nexistepas")
        with self.assertRaises(StorageError):
            elsewhere.check_container()

    def test_text_extraction_straight_from_a_blob(self):
        name = self.files.save_file(
            "documents/1/bibliotheque/cv.docx",
            make_docx("Lionel Dupont", "lionel@example.be · 0470 12 34 56", CV_BODY),
        )
        result = self.files.extract_and_anonymize_text(
            name, known=privacy.known_identity(display_name="Lionel Dupont")
        )
        self.assertTrue(result.text.startswith("[NOM]\n[EMAIL] · [TELEPHONE]"))
        self.assertNotIn("example.be", result.text)


@override_settings(MEDIA_ROOT=MEDIA)
class DocumentAdminTests(TestCase):
    """The admin reads ``document.file.url``: it must never be a raw path,
    and never raise (``ClearableFileInput`` only swallows ``AttributeError``)."""

    def test_the_change_form_shows_a_signed_link_not_a_media_path(self):
        admin = make_user("Boss", username="boss")
        admin.is_staff = admin.is_superuser = True
        admin.save()
        self.client.force_login(admin)
        application = make_application(owner=admin)
        document = make_document(
            application=application, kind=DocumentKind.CV, label="CV",
            file=SimpleUploadedFile("admin.docx", b"x"),
        )
        for url in (
            reverse("admin:tracker_document_change", args=[document.pk]),
            reverse("admin:tracker_document_changelist"),
            reverse("admin:tracker_application_change", args=[application.pk]),
        ):
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertEqual(response.status_code, 200)
        body = self.client.get(
            reverse("admin:tracker_document_change", args=[document.pk])
        ).content.decode()
        self.assertIn("/fichiers/", body)
        self.assertNotIn("/media/", body)


class DocumentDownloadOnAzureTests(OwnedTestCase):
    """With the bytes elsewhere, the download view sends the browser there."""

    def setUp(self):
        super().setUp()
        self.service = FakeService(account_key=AZURE_KEY)
        self.override = override_settings(STORAGES=azure_settings(self.service, link_ttl=300))
        self.override.enable()
        self.addCleanup(self.override.disable)
        self.application = make_application()

    def upload(self, name="CV.pdf", kind=DocumentKind.CV, **data):
        return self.client.post(
            reverse("tracker:add_document", args=[self.application.pk]),
            {"file": SimpleUploadedFile(name, b"%PDF-1.4 x"), "kind": kind, **data},
            HTTP_HX_REQUEST="true",
        )

    def test_upload_download_and_delete_go_through_the_blob_container(self):
        self.assertEqual(self.upload().status_code, 200)
        document = self.application.documents.get()
        name = document.file.name or ""
        self.assertEqual(name, f"documents/{self.user.pk}/{self.application.slug}/CV.pdf")
        self.assertEqual(self.service.container.blobs[name], b"%PDF-1.4 x")
        self.assertEqual(document.size_bytes, 10)

        response = self.client.get(document.get_download_url())
        self.assertEqual(response.status_code, 302)
        location = response.headers["Location"]
        self.assertTrue(location.startswith(
            f"https://jobhunt.blob.core.windows.net/documents/documents/{self.user.pk}/"
        ))
        self.assertEqual(parse_qs(urlsplit(location).query)["sp"], ["r"])
        self.assertEqual(urlsplit(document.file.url).netloc, "jobhunt.blob.core.windows.net")

        # The library page lists the row without a single call to the account.
        self.service.container.broken = True
        response = self.client.get(reverse("tracker:document_library"))
        self.assertContains(response, "CV.pdf")
        self.service.container.broken = False

        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(
                reverse("tracker:delete_document", args=[document.pk]), HTTP_HX_REQUEST="true"
            )
        self.assertEqual(self.service.container.blobs, {})

    def test_the_library_page_never_asks_the_provider_for_a_size(self):
        for index in range(3):
            self.upload(name=f"CV{index}.pdf")
        self.assertEqual(self.application.documents.count(), 3)
        # Every row carries its size, so a listing costs no round-trip — and
        # renders even while the account is unreachable.
        self.service.container.broken = True
        response = self.client.get(reverse("tracker:document_library"))
        self.assertEqual(response.status_code, 200)
        for index in range(3):
            self.assertContains(response, f"CV{index}.pdf")
        self.assertContains(response, "10 o")

    def test_a_row_without_its_size_falls_back_and_survives_an_outage(self):
        self.upload()
        document = self.application.documents.get()
        Document.objects.filter(pk=document.pk).update(size_bytes=None)
        document = Document.objects.get(pk=document.pk)
        self.assertEqual(document.size_display, "10 o")
        self.service.container.broken = True
        document = Document.objects.get(pk=document.pk)
        self.assertEqual(document.size_display, "—")

    def test_a_stranger_gets_404_before_any_link_is_minted(self):
        self.upload()
        document = self.application.documents.get()
        self.client.force_login(make_user("Marie"))
        self.assertEqual(self.client.get(document.get_download_url()).status_code, 404)

    def test_a_cv_upload_is_analysed_from_the_blob(self):
        analyzer = FakeAnalyzer()
        with mock.patch("tracker.views.cv_analyzer", return_value=analyzer):
            self.client.post(
                reverse("tracker:add_document", args=[self.application.pk]),
                {"file": SimpleUploadedFile("cv.txt", f"Jo jo@x.be\n{CV_BODY}".encode()),
                 "kind": DocumentKind.CV},
                HTTP_HX_REQUEST="true",
            )
        (_, call), = analyzer.calls
        self.assertTrue(call["text"].text.startswith("Jo [EMAIL]"))


class IngestCVTests(SimpleTestCase):
    """The CV use case on the memory adapters: no database, no disk, no SDK."""

    def setUp(self):
        self.store = MemoryPersistence()
        self.files = MemoryStorageAdapter()
        self.user = User(pk=1, username="lionel", email="lionel.dupont@example.be")
        self.analyzer = FakeAnalyzer()
        self.known = privacy.known_identity(
            display_name="Lionel Dupont", username="lionel", email=self.user.email, location="Nivelles"
        )

    def ingest(self, upload, application=None, analyzer: Any = "default", **fields):
        return services.ingest_cv(
            self.user, application, Document(kind=DocumentKind.CV, language="fr", **fields),
            upload=upload, label="Mon CV", known=self.known,
            analyzer=self.analyzer if analyzer == "default" else analyzer,
            persistence=self.store, storage=self.files,
        )

    def test_the_analyzer_gets_the_anonymised_text_and_nothing_else(self):
        payload = make_docx("Lionel Dupont", "lionel.dupont@example.be · 0470 12 34 56 · Nivelles",
                            CV_BODY)
        intake = self.ingest(SimpleUploadedFile("Mon CV.docx", payload))
        self.assertTrue(intake.analyzed)
        document = intake.document
        self.assertEqual(document.file.name, "documents/1/bibliotheque/Mon_CV.docx")
        self.assertTrue(self.files.file_exists(document.file.name or ""))
        self.assertEqual(document.size_bytes, len(payload))
        self.assertEqual(self.store.documents.count(self.user), 1)
        (owner, call), = self.analyzer.calls
        self.assertIs(owner, self.user)
        self.assertEqual(set(call), {"document_id", "label", "language", "text"})
        self.assertEqual((call["document_id"], call["label"], call["language"]), (document.pk, "Mon CV", "fr"))
        self.assertIs(call["text"], intake.anonymized)
        self.assertEqual(
            call["text"].text, f"[NOM]\n[EMAIL] · [TELEPHONE] · [LIEU]\n{CV_BODY}"
        )
        self.assertNotIn("lionel", call["text"].text.lower())
        self.assertEqual(call["text"].summary(), "1 e-mail, 1 numéro de téléphone, 1 nom, 1 lieu masqués")

    def test_an_attached_cv_is_filed_under_the_application_and_logged(self):
        company = self.store.companies.get_or_create(self.user, "Acme", Sector.PRIVATE)
        application = self.store.applications.add(
            Application(owner=self.user, company=company, title="DevOps", status=Status.TO_APPLY)
        )
        intake = self.ingest(
            SimpleUploadedFile("cv.txt", CV_BODY.encode()), application=application, is_primary=True
        )
        self.assertEqual(intake.document.file.name, f"documents/1/{application.slug}/cv.txt")
        self.assertIs(intake.document.application, application)
        self.assertEqual([e.title for e in self.store.events.rows()], ["Document ajouté : Mon CV"])

    def test_the_label_crosses_the_port_anonymised_too(self):
        """It defaults to the uploaded file name, which carries the applicant's
        name far more reliably than the CV body does."""
        upload = SimpleUploadedFile("CV Lionel Dupont - Nivelles.txt", CV_BODY.encode())
        intake = services.ingest_cv(
            self.user, None, Document(kind=DocumentKind.CV, language="fr"),
            upload=upload, label=upload.name or "", known=self.known,
            analyzer=self.analyzer, persistence=self.store, storage=self.files,
        )
        (_, call), = self.analyzer.calls
        self.assertEqual(call["label"], "CV [NOM] - [LIEU].txt")
        self.assertNotIn("Lionel", call["label"])
        # What the owner sees in their own library is untouched.
        self.assertEqual(intake.document.label, "CV Lionel Dupont - Nivelles.txt")

    def test_an_analyzer_that_raises_does_not_lose_the_upload(self):
        class Broken:
            def analyze_cv(self, owner, **kwargs):
                raise RuntimeError("le copilote a explosé")

        with self.assertLogs("tracker.services", "ERROR"):
            intake = self.ingest(SimpleUploadedFile("cv.txt", CV_BODY.encode()), analyzer=Broken())
        self.assertFalse(intake.analyzed)
        self.assertIsNotNone(intake.anonymized)
        self.assertEqual(self.store.documents.count(self.user), 1)
        self.assertTrue(self.files.file_exists(intake.document.file.name or ""))

    def test_a_storage_outage_while_reading_back_does_not_lose_the_upload(self):
        with mock.patch.object(MemoryStorageAdapter, "extract_and_anonymize_text",
                               side_effect=StorageError("panne")):
            with self.assertLogs("tracker.services", "ERROR"):
                intake = self.ingest(SimpleUploadedFile("cv.txt", CV_BODY.encode()))
        self.assertEqual((intake.analyzed, intake.anonymized), (False, None))
        self.assertEqual(self.analyzer.calls, [])
        self.assertTrue(self.files.file_exists(intake.document.file.name or ""))

    def test_unreadable_formats_are_stored_but_not_analysed(self):
        intake = self.ingest(SimpleUploadedFile("cv.odt", b"pas lisible"))
        self.assertFalse(intake.analyzed)
        self.assertIsNone(intake.anonymized)
        self.assertTrue(self.files.file_exists(intake.document.file.name or ""))
        self.assertEqual(self.analyzer.calls, [])

    def test_a_scan_without_text_is_not_analysed(self):
        intake = self.ingest(SimpleUploadedFile("scan.pdf", make_pdf()))
        self.assertFalse(intake.analyzed)
        assert intake.anonymized is not None
        self.assertEqual(intake.anonymized.text, "")
        self.assertEqual(self.analyzer.calls, [])

    def test_too_little_text_to_be_a_cv_is_stored_but_not_analysed(self):
        short = ("CV-" * services.MIN_TEXT_LENGTH)[: services.MIN_TEXT_LENGTH - 1]
        self.assertEqual(len(short), services.MIN_TEXT_LENGTH - 1)
        intake = self.ingest(SimpleUploadedFile("court.txt", short.encode()))
        self.assertFalse(intake.analyzed)
        self.assertEqual(self.analyzer.calls, [])
        self.assertTrue(self.files.file_exists(intake.document.file.name or ""))
        # One character more and it is analysed: the threshold is the only reason.
        self.assertTrue(self.ingest(SimpleUploadedFile("long.txt", (short + "x").encode())).analyzed)

    def test_without_an_analyzer_nothing_is_extracted(self):
        with mock.patch.object(MemoryStorageAdapter, "extract_and_anonymize_text",
                               side_effect=AssertionError("extrait")):
            intake = self.ingest(SimpleUploadedFile("cv.txt", CV_BODY.encode()), analyzer=None)
        self.assertEqual((intake.analyzed, intake.anonymized), (False, None))
        self.assertTrue(self.files.file_exists("documents/1/bibliotheque/cv.txt"))

    def test_a_failed_row_write_removes_the_stored_file(self):
        with mock.patch.object(self.store.documents, "add", side_effect=RuntimeError("panne")):
            with self.assertRaises(RuntimeError):
                self.ingest(SimpleUploadedFile("cv.txt", CV_BODY.encode()))
        self.assertFalse(self.files.file_exists("documents/1/bibliotheque/cv.txt"))
        self.assertEqual(self.analyzer.calls, [])


class StoragePruneTests(TestCase):
    """The sweep that collects what a rollback or a crash left behind."""

    def setUp(self):
        self.files = MemoryStorageAdapter()
        self.override = override_settings(
            STORAGES={
                "default": {
                    "BACKEND": "tracker.adapters.file_storage.MemoryStorageAdapter",
                    "OPTIONS": {"link_ttl": 60},
                },
                "staticfiles": MEMORY_STORAGES["staticfiles"],
            }
        )
        self.override.enable()
        self.addCleanup(self.override.disable)
        self.user = make_user("Lionel", username="lionel")

    def prune(self, *args):
        out = io.StringIO()
        call_command("storage_prune", *args, stdout=out)
        return out.getvalue()

    def test_a_referenced_file_is_never_touched(self):
        document = make_document(
            owner=self.user, kind=DocumentKind.CV, label="CV",
            file=SimpleUploadedFile("cv.docx", b"contenu"),
        )
        name = document.file.name or ""
        report = self.prune("--older-than", "0", "--delete")
        self.assertIn("Aucun fichier orphelin", report)
        self.assertTrue(storage().file_exists(name))

    def test_an_orphan_is_listed_then_removed_only_with_delete(self):
        orphan = storage().save_file(f"documents/{self.user.pk}/bibliotheque/perdu.docx", b"x")
        self.assertIn(orphan, self.prune("--older-than", "0"))
        self.assertTrue(storage().file_exists(orphan))  # dry run by default
        self.assertIn("1 fichier(s) supprimé", self.prune("--older-than", "0", "--delete"))
        self.assertFalse(storage().file_exists(orphan))

    def test_an_upload_in_flight_is_not_collected(self):
        """Its bytes are written and its row is not: the grace period is what
        keeps the sweep from deleting a file mid-request."""
        fresh = storage().save_file(f"documents/{self.user.pk}/bibliotheque/en-cours.docx", b"x")
        report = self.prune("--delete")
        self.assertIn("Aucun fichier orphelin", report)
        self.assertIn("1 fichier(s) récent(s) ignoré", report)
        self.assertTrue(storage().file_exists(fresh))

    def test_a_negative_grace_is_refused(self):
        from django.core.management.base import CommandError

        with self.assertRaises(CommandError):
            self.prune("--older-than", "-1")


class StorageResolverTests(SimpleTestCase):
    def test_the_configured_backend_is_the_port(self):
        self.assertIsInstance(storage(), LocalStorageAdapter)
        with override_settings(STORAGES=MEMORY_STORAGES):
            backend = storage()
            self.assertIsInstance(backend, MemoryStorageAdapter)
            self.assertIs(storage(), backend)
            assert isinstance(backend, MemoryStorageAdapter)
            self.assertEqual(backend.link_ttl, 60)
        self.assertIsInstance(storage(), LocalStorageAdapter)
        self.assertIs(storage(), storage())

    def test_a_backend_without_the_port_is_refused(self):
        plain = {"default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
                 "staticfiles": MEMORY_STORAGES["staticfiles"]}
        with override_settings(STORAGES=plain), self.assertRaises(ImproperlyConfigured):
            storage()

    def test_the_analyzer_uses_settings_or_publishes_a_workflow_event(self):
        from tracker.events import CVEventPublisher, cv_ingested

        with mock.patch.object(cv_ingested, "has_listeners", return_value=False):
            with override_settings(CV_ANALYZER=None):
                self.assertIsNone(cv_analyzer())
            with override_settings(CV_ANALYZER="tracker.tests.FakeAnalyzer"):
                analyzer = cv_analyzer()
                self.assertIsInstance(analyzer, FakeAnalyzer)
                self.assertIs(cv_analyzer(), analyzer)
        with mock.patch.object(cv_ingested, "has_listeners", return_value=True):
            with override_settings(CV_ANALYZER=None):
                self.assertIsInstance(cv_analyzer(), CVEventPublisher)

    def test_two_extensions_offering_the_same_service_is_a_configuration_error(self):
        two = (SimpleNamespace(name="a", cv_analyzer="x.A"), SimpleNamespace(name="b", cv_analyzer="x.B"))
        with mock.patch.object(plugin_registry, "get_plugins", return_value=two):
            with self.assertRaises(ImproperlyConfigured):
                plugin_registry.plugin_attribute("cv_analyzer")
            self.assertEqual(plugin_registry.plugin_attribute("absent"), "")


class StorageCheckTests(SimpleTestCase):
    def check(self, provider, options, missing=()):
        storages = {"default": {"BACKEND": "x", "OPTIONS": options}, "staticfiles": MEMORY_STORAGES["staticfiles"]}
        with override_settings(STORAGE_PROVIDER=provider, STORAGES=storages):
            with mock.patch("tracker.checks.find_spec", side_effect=lambda name: None if name in missing else object()):
                return [problem.id for problem in checks.check_storage(None)]

    def test_local_is_silent(self):
        self.assertEqual(self.check("local", {}), [])

    def test_memory_warns(self):
        self.assertEqual(self.check("memory", {}), ["tracker.W002"])

    def test_azure_needs_its_driver_and_an_identity_library_without_a_key(self):
        self.assertEqual(self.check("azure", {"account_key": "k"}), [])
        self.assertEqual(self.check("azure", {"account_key": "k"}, missing={"azure.storage.blob"}), ["tracker.E002"])
        self.assertEqual(self.check("azure", {"connection_string": "x"}, missing={"azure.identity"}), [])
        self.assertEqual(self.check("azure", {"account_url": "https://a"}, missing={"azure.identity"}), ["tracker.E003"])
        self.assertEqual(self.check("azure", {"account_url": "https://a"}), [])


class StorageStatusCommandTests(SimpleTestCase):
    def test_describes_the_provider_and_probes_it(self):
        out = io.StringIO()
        with override_settings(STORAGES=MEMORY_STORAGES, STORAGE_PROVIDER="memory"):
            call_command("storage_status", stdout=out)
        text = out.getvalue()
        self.assertIn("MemoryStorageAdapter", text)
        self.assertIn("absente, comme prévu", text)
        self.assertIn("/fichiers/", text)
        self.assertIn("Stockage : ok", text)

    def test_a_mistyped_container_is_reported(self):
        from django.core.management.base import CommandError

        service = FakeService(account_key=AZURE_KEY)
        service.container.missing = True
        with override_settings(STORAGES=azure_settings(service), STORAGE_PROVIDER="azure"):
            with self.assertRaises(CommandError):
                call_command("storage_status", stdout=io.StringIO())

    def test_probe_writes_reads_and_removes(self):
        out = io.StringIO()
        with override_settings(STORAGES=MEMORY_STORAGES, STORAGE_PROVIDER="memory"):
            call_command("storage_status", "--probe", stdout=out)
            self.assertFalse(storage().file_exists(storage_status.PROBE))
        self.assertIn("7 octets écrits, relus et supprimés", out.getvalue())

    def test_an_unreachable_provider_is_a_command_error(self):
        from django.core.management.base import CommandError

        service = FakeService(account_key=AZURE_KEY)
        service.container.broken = True
        with override_settings(STORAGES=azure_settings(service), STORAGE_PROVIDER="azure"):
            with self.assertRaises(CommandError):
                call_command("storage_status", stdout=io.StringIO())


class ArchitectureGuardTests(SimpleTestCase):
    """The layering, enforced by grep so it survives the next refactor."""

    ROOT = Path(tracker.__file__).parent

    def source(self, name: str) -> str:
        return (self.ROOT / name).read_text(encoding="utf-8")

    def test_views_and_services_never_query_the_orm_themselves(self):
        for name in ("views.py", "services.py"):
            text = self.source(name)
            for token in (".objects.", "owned_or_404(", "get_object_or_404(", ".for_user(",
                          "select_related(", "prefetch_related(", ".log(", "apply_status(",
                          "_default_manager"):
                with self.subTest(file=name, token=token):
                    self.assertNotIn(token, text)

    def test_rules_and_use_cases_never_take_the_model_shortcuts(self):
        for name in ("services.py", "domain.py", "adapters/memory.py"):
            text = self.source(name)
            for token in ("apply_status(", ".log(", "owner_preferences(", "preferences_for(",
                          ".objects."):
                with self.subTest(file=name, token=token):
                    self.assertNotIn(token, text)

    def runtime_imports(self, name: str) -> list[str]:
        """Every module the file really imports when it runs.

        Not a regex: ``^`` misses a function-local import, which runs like
        any other, and only what sits under ``if TYPE_CHECKING:`` is exempt
        because it never runs at all.
        """
        found: list[str] = []

        def visit(nodes):
            for node in nodes:
                if isinstance(node, ast.Import):
                    found.extend(alias.name for alias in node.names)
                elif isinstance(node, ast.ImportFrom) and node.level == 0:
                    module = node.module or ""
                    found.append(module)
                    found.extend(f"{module}.{alias.name}" for alias in node.names)
                elif isinstance(node, ast.If) and self.is_type_checking(node.test):
                    visit(node.orelse)  # the runtime half of the branch only
                    continue
                for child in ast.iter_child_nodes(node):
                    visit([child])

        visit(ast.parse(self.source(name)).body)
        return found

    @staticmethod
    def is_type_checking(test: ast.expr) -> bool:
        return (isinstance(test, ast.Name) and test.id == "TYPE_CHECKING") or (
            isinstance(test, ast.Attribute) and test.attr == "TYPE_CHECKING"
        )

    def test_the_guard_sees_an_import_wherever_it_hides(self):
        """The guard itself, checked: an indented import runs, and an earlier
        version of this test compared with ``^`` and passed ``re.MULTILINE``
        where ``assertNotRegex`` expects a message."""
        source = (
            "import os\n"
            "from typing import TYPE_CHECKING\n"
            "if TYPE_CHECKING:\n"
            "    from django.db import models\n"
            "def f():\n"
            "    from django.db import transaction\n"
            "    import azure.storage.blob\n"
        )
        with mock.patch.object(ArchitectureGuardTests, "source", return_value=source):
            found = self.runtime_imports("whatever.py")
        self.assertIn("django.db.transaction", found)
        self.assertIn("azure.storage.blob", found)
        self.assertNotIn("django.db.models", found)

    def test_services_stay_clear_of_django_plumbing(self):
        forbidden = {
            "services.py": ("django.db", "django.shortcuts", "jobhunt.plugins",
                            "tracker.queries", "tracker.adapters.django_orm",
                            "tracker.adapters.file_storage", "tracker.adapters.azure_storage",
                            "tracker.links", "azure"),
            "domain.py": ("django.db",),
            "adapters/memory.py": ("django.db",),
            # Pure by design: rules on strings and bytes, no framework at all.
            "privacy.py": ("django", "azure", "tracker.models"),
            "adapters/document_text.py": ("django", "azure", "tracker.models"),
            # The provider lives in its adapter, and nowhere else.
            "adapters/file_storage.py": ("azure",),
            "views.py": ("azure", "tracker.adapters.azure_storage", "tracker.adapters.file_storage"),
        }
        for name, modules in forbidden.items():
            imported = self.runtime_imports(name)
            for module in modules:
                with self.subTest(file=name, module=module):
                    offenders = [
                        found
                        for found in imported
                        if found == module or found.startswith(f"{module}.")
                    ]
                    self.assertEqual(offenders, [], f"{name} importe {module}")


def tearDownModule():
    shutil.rmtree(MEDIA, ignore_errors=True)
